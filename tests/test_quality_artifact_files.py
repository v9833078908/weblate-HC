# Copyright © HCGameLoc
# SPDX-License-Identifier: GPL-3.0-or-later
from __future__ import annotations

import hashlib
import io
import json
import zipfile

import pytest
from app.db import DomainError
from app.models import Job
from openpyxl import load_workbook
from test_quality_artifacts import artifact_inputs, build
from test_quality_only_run import finish, start

pytest_plugins = ["test_quality_only_run"]


def example():
    values = artifact_inputs()
    values["final_snapshot"]["judged_revisions"].pop("12")
    values["final_snapshot"]["issues"] = [
        {"group_id": "term", "kind": "glossary", "hard": True, "advisory": False, "unit_ids": [11]}
    ]
    return values["exported"], build(values)


def test_excel_colors_match_current_issues_without_changing_any_values():
    from app.services.quality_artifact_files import (
        serialize_quality_report,
        serialize_quality_translation,
    )

    exported, artifact = example()
    payload, mime, name = serialize_quality_translation(exported, artifact)
    assert name == "run-run-1-translation.xlsx"
    wb = load_workbook(io.BytesIO(payload))
    ws = wb.active
    assert ws["C2"].fill.fgColor.rgb == "FFFFC0CB"
    assert ws["C4"].fill.fgColor.rgb == "FFFFF2CC"
    assert ws["B2"].fill.patternType is None
    assert ws["D4"].fill.patternType is None
    assert ws["C3"].fill.patternType is None
    assert list(ws.values) == [
        ("key", "en", "de", "fr"),
        ("one", "Play", "Spielen", "Jouer"),
        (None, None, None, None),
        ("two", "Play", "Spiel", "Jouer"),
    ]
    report, _, _ = serialize_quality_report(artifact, exported)
    report_wb = load_workbook(io.BytesIO(report))
    assert report_wb.sheetnames == ["Сводка", "Замечания"]
    rows = list(report_wb["Замечания"].values)
    assert len(rows) == 3
    assert "C2" in rows[1] and "C4" in rows[2]
    assert report_wb["Замечания"].freeze_panes == "A2"


def test_bundle_hashes_and_separate_downloads_identify_the_same_files():
    from app.services.quality_artifact_files import (
        serialize_quality_bundle,
        serialize_quality_report,
        serialize_quality_translation,
    )

    exported, artifact = example()
    payload, _, _ = serialize_quality_bundle(artifact, exported)
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        manifest = json.loads(archive.read("manifest.json"))
        assert manifest["run_id"] == "run-1"
        for entry in manifest["files"]:
            assert hashlib.sha256(archive.read(entry["name"])).hexdigest() == entry["sha256"]
        for data, _, filename in [
            serialize_quality_translation(exported, artifact),
            serialize_quality_report(artifact, exported),
        ]:
            assert archive.read(filename) == data


def test_report_cells_cannot_execute_translation_or_reason_as_a_formula():
    from app.services.quality_artifact_files import serialize_quality_report

    exported, artifact = example()
    artifact["issues"][0].update(source="=1+1", target="@SUM(1)", reason="+cmd", key="-1+2")
    payload, _, _ = serialize_quality_report(artifact, exported)
    ws = load_workbook(io.BytesIO(payload))["Замечания"]
    assert all(cell.data_type != "f" for row in ws for cell in row)
    assert "=1+1" in list(ws.values)[1]


@pytest.mark.parametrize("value", ["x" * 32768, "bad\x00text"], ids=["too-long", "illegal-control"])
def test_excel_limits_are_explicit_instead_of_silent_truncation(value):
    from app.services.quality_artifact_files import serialize_quality_report

    exported, artifact = example()
    artifact["issues"][0]["reason"] = value
    with pytest.raises(DomainError):
        serialize_quality_report(artifact, exported)


def test_sheet_name_collisions_still_highlight_by_sheet_identity():
    from app.services.quality_artifact_files import serialize_quality_translation

    exported, artifact = example()
    for name in ["Strings", "Strings-3", "Strings"]:
        exported["sheets"].append(
            {
                "id": "extra-" + str(len(exported["sheets"])),
                "name": name,
                "file": "extra.xlsx",
                "rows": [["ok"]],
            }
        )
    payload, _, _ = serialize_quality_translation(exported, artifact)
    wb = load_workbook(io.BytesIO(payload))
    assert len({n.casefold() for n in wb.sheetnames}) == 4
    assert wb.worksheets[0]["C2"].fill.fgColor.rgb == "FFFFC0CB"
    assert all(ws["A1"].fill.patternType is None for ws in wb.worksheets[1:])


def test_api_report_files_and_json_are_frozen_and_owned(quality_setup, monkeypatch):
    store, settings, client, remote, body = quality_setup
    run_id = start(client, body).json()["id"]
    data = finish(store, settings, remote, run_id)
    artifact = store.get("quality_artifact", data["quality_artifact_id"], "owner")
    artifact.pop("revision")
    artifact["issues"] = [
        {
            "category": "unverified",
            "unverified": True,
            "blocking": False,
            "unit_id": i,
            "reason": "missing",
        }
        for i in range(500)
    ]
    store.save("quality_artifact", data["quality_artifact_id"], "owner", artifact)
    monkeypatch.setattr(
        remote, "quality_snapshot", lambda *a: pytest.fail("Download called Weblate")
    )
    base = f"/api/producer/runs/{run_id}"
    ids = []
    for page in range(1, 11):
        response = client.get(
            base + f"/quality-report/?format=json&page={page}&category=unverified"
        )
        assert response.status_code == 200, response.text
        result = response.json()
        assert result["total"] == 500
        assert result["provenance"] == artifact["provenance"]
        ids.extend(i["unit_id"] for i in result["issues"])
    assert ids == list(range(500))
    assert (
        client.get(base + "/quality-report/?format=json&category=blocking").json()["issues"] == []
    )
    for query in ["page=0", "page_size=101", "page=-1", "category=other"]:
        assert client.get(base + "/quality-report/?format=json&" + query).status_code == 422
    for path in ["quality-report/", "result-bundle/"]:
        # Restore a complete saved report for XLSX serialization.
        artifact["issues"] = []
        store.save("quality_artifact", data["quality_artifact_id"], "owner", artifact)
        response = client.get(base + "/" + path)
        assert response.status_code == 200
        assert response.headers["cache-control"] == "no-store"
        assert response.headers["x-content-type-options"] == "nosniff"
    with store.session() as session:
        session.get(Job, run_id).owner = "different-owner"
    assert client.get(base + "/quality-report/?format=json").status_code == 404
    assert client.get(base + "/result-bundle/").status_code == 404
    assert client.get("/api/producer/runs/parent/quality-report/").status_code == 404


def test_clean_report_is_not_a_failed_or_missing_report():
    from app.services.quality_artifact_files import serialize_quality_report

    values = artifact_inputs()
    payload, _, _ = serialize_quality_report(build(values), values["exported"])
    wb = load_workbook(io.BytesIO(payload))
    assert wb["Замечания"].max_row == 1
    assert "Остаточных замечаний нет" in str(list(wb["Сводка"].values))


def test_full_copy_cycle_exports_repaired_rows_with_pink_and_yellow_residuals(
    quality_setup, monkeypatch
):
    store, settings, client, remote, body = quality_setup
    parent = client.get("/api/producer/runs/parent/export/?format=json").content
    original_snapshot = remote.quality_snapshot

    def snapshot(*args):
        value = original_snapshot(*args)
        value["judged_revisions"].pop("2", None)
        value["issues"].append(
            {
                "group_id": "required-term",
                "kind": "glossary",
                "hard": True,
                "advisory": False,
                "unit_ids": [1],
            }
        )
        return value

    monkeypatch.setattr(remote, "quality_snapshot", snapshot)
    run_id = start(client, body).json()["id"]
    data = finish(store, settings, remote, run_id)
    assert data["quality"]["status"] == "completed_with_exceptions"
    assert remote.applies and remote.repeat_applied
    base = f"/api/producer/runs/{run_id}"
    report = client.get(base + "/quality-report/?format=json").json()
    assert report["summary"]["blocking_cells"] == 1
    assert report["summary"]["unverified_cells"] == 1
    assert report["summary"]["automatically_corrected"] == 2
    wb = load_workbook(io.BytesIO(client.get(base + "/export/?format=xlsx").content))
    assert wb.active["C2"].value == wb.active["C3"].value == "corrected"
    assert wb.active["C2"].fill.fgColor.rgb == "FFFFC0CB"
    assert wb.active["C3"].fill.fgColor.rgb == "FFFFF2CC"
    count = len(remote.starts)
    for _ in range(3):
        response = client.get(base + "/result-bundle/")
        assert response.status_code == 200
        with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
            assert len(archive.namelist()) == 3
    assert len(remote.starts) == count
    assert client.get("/api/producer/runs/parent/export/?format=json").content == parent
