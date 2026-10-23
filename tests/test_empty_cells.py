# Copyright © HCGameLoc
# SPDX-License-Identifier: GPL-3.0-or-later
from __future__ import annotations

import pytest
from app.db import Store
from app.schemas import UploadPatch
from app.services.drafts import create_draft, patch_draft
from app.services.exports import build_export
from app.services.kits import analyze_file
from openpyxl import Workbook
from openpyxl.styles import PatternFill


def cleared_workbook(tmp_path, *, target_header=True, only_empty=False):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Strings"
    if not only_empty:
        sheet.append(["Key", "en", "de"] if target_header else ["Key", "en"])
        sheet.append(["first", "  Hello  ", "Fertig"] if target_header else ["first", "Hello"])
        sheet.append([" \t", "\u00a0", "\n"])
        sheet.append(["second", "Goodbye", ""] if target_header else ["second", "Goodbye"])
    # Excel retains these dimensions/style after the actual value is cleared.
    cell = sheet["Z40"]
    cell.value = "removed old content"
    cell.fill = PatternFill(fill_type="solid", fgColor="FFFF00")
    cell.value = None
    path = tmp_path / "cleared.xlsx"
    workbook.save(path)
    workbook.close()
    return path


def test_analysis_counts_content_not_historical_excel_dimensions(tmp_path):
    path = cleared_workbook(tmp_path)
    original = path.read_bytes()

    analysis, tables, _ = analyze_file(path, path.name)

    sheet = analysis["sheets"][0]
    assert sheet["rows"] == 3  # Header + two real rows, not Excel's 40-row range.
    assert sheet["columns"] == ["Key", "en", "de"]
    assert tables["Strings"][1] == ["first", "  Hello  ", "Fertig"]
    assert tables["Strings"][3] == ["second", "Goodbye", ""]
    assert path.read_bytes() == original


def test_entirely_cleared_sheet_has_no_content_or_columns(tmp_path):
    path = cleared_workbook(tmp_path, only_empty=True)

    analysis, _, _ = analyze_file(path, path.name)

    assert analysis["sheets"][0]["rows"] == 0
    assert analysis["sheets"][0]["columns"] == []


@pytest.mark.parametrize("suffix,delimiter", [("csv", ","), ("tsv", "\t")])
def test_empty_records_do_not_inflate_analysis_counts(tmp_path, suffix, delimiter):
    path = tmp_path / ("kit." + suffix)
    path.write_text(
        delimiter.join(["Key", "en", "de"])
        + "\n"
        + delimiter.join(["a", "Hello", ""])
        + "\n"
        + delimiter * 2
        + "\n"
        + delimiter.join(["b", "Goodbye", ""])
        + "\n"
        + delimiter * 2
        + "\n",
        encoding="utf-8",
    )

    analysis, tables, _ = analyze_file(path, path.name)

    assert analysis["sheets"][0]["rows"] == 3
    assert tables["kit"][3] == ["b", "Goodbye", ""]


@pytest.mark.parametrize("target_header", [True, False])
def test_import_and_export_keep_coordinates_and_real_translation_cells(tmp_path, target_header):
    path = cleared_workbook(tmp_path, target_header=target_header)
    original = path.read_bytes()
    store = Store(tmp_path / "data")
    draft = create_draft(store, "owner", "local", [(path.name, original)], "string", "quick")
    saved = patch_draft(
        store,
        "owner",
        draft["token"],
        UploadPatch(
            revision=draft["revision"],
            ui_mode="quick",
            current_step="confirm",
            answers={
                "source_language": "en",
                "included_target_languages": ["de"],
                "selected_sheets": [{"file": path.name, "sheet": "Strings"}],
            },
        ),
    )

    assert saved["errors"] == []
    manifest = saved["manifest"]
    assert manifest["totals"] == {"input_rows": 2, "imported_rows": 2, "quarantined_rows": 0}
    component = manifest["components"][0]
    assert [unit["original_row"] for unit in component["units"]] == [2, 4]
    assert component["profile"]["languages"][-1]["column"] == 3
    assert component["units"][1]["values"]["de"] == ""
    exported = build_export(manifest, {})["sheets"][0]["rows"]
    assert exported[0] == ["Key", "en", "de"]
    assert exported[3] == ["second", "Goodbye", ""]
    if target_header:
        assert exported[1] == ["first", "  Hello  ", "Fertig"]
    assert path.read_bytes() == original


def test_zero_false_and_formula_are_not_treated_as_empty(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Strings"
    sheet.append(["Key", "en", "de"])
    sheet.append(["zero", 0, False])
    sheet.append(["formula", "Hello", '=""'])
    path = tmp_path / "values.xlsx"
    workbook.save(path)
    workbook.close()

    analysis, tables, _ = analyze_file(path, path.name)

    assert analysis["sheets"][0]["rows"] == 3
    assert tables["Strings"][1] == ["zero", "0", "False"]
    assert tables["Strings"][2][2] == '=""'
    assert analysis["sheets"][0]["formula_cells"][0]["cell"] == "C3"
    assert not analysis["sheets"][0]["formula_cells"][0]["auto_ignore"]


def test_ghost_cells_do_not_disable_raw_reader_resource_limits(tmp_path, monkeypatch):
    from app.db import DomainError
    from loc_kit_ingest import reader

    path = cleared_workbook(tmp_path)
    monkeypatch.setattr(reader, "_MAX_ROWS", 10)

    with pytest.raises(DomainError, match="input limit exceeded"):
        analyze_file(path, path.name)
