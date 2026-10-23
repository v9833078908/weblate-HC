# Copyright © HCGameLoc
# SPDX-License-Identifier: GPL-3.0-or-later
from __future__ import annotations

import pytest


def test_complete_export_keeps_problem_only_sheet_and_selected_export_merges_groups(tmp_path):
    import copy

    from app.services.exports import build_export, component_exports
    from test_multisource_manifest import prepared

    manifest = prepared(tmp_path)
    translations = {
        c["id"]: {"pl": {u["key"]: "Ready" for u in c["units"]}} for c in manifest["components"]
    }
    parts = component_exports(manifest, translations)
    assert set(parts) == {"kit-1"}
    assert parts["kit-1"]["keys"] == 3
    assert len(parts["kit-1"]["sheets"]) == 1
    assert parts["kit-1"]["per_language"]["pl"]["translated"] == 3
    single = copy.deepcopy(manifest)
    single["components"] = single["components"][:1]
    problems = copy.deepcopy(single["sheets"][0])
    problems.update(id="kit-2", sheet="Problems")
    single["sheets"].append(problems)
    assert len(build_export(single, {})["sheets"]) == 2


def test_multisource_export_merges_original_sheet(tmp_path):
    import io

    from app.services.exports import build_export, serialize_export
    from openpyxl import load_workbook
    from test_multisource_manifest import prepared

    manifest = prepared(tmp_path)
    result = build_export(
        manifest,
        {
            "kit-1-src-ru": {"pl": {"b": "Świat"}},
            "kit-1-src-en": {"ru": {"c": "Играть"}, "pl": {"c": "Graj"}},
        },
    )
    assert len(result["sheets"]) == 1
    rows = result["sheets"][0]["rows"]
    assert rows[2] == ["b", "", "Мир", "World", "Świat"]
    assert rows[3] == ["c", "", "Играть", "Play", "Graj"]
    assert result["keys"] == 3
    assert len(rows) == 6
    payload, _, _ = serialize_export(result, "local", "xlsx")
    workbook = load_workbook(io.BytesIO(payload))
    assert len(workbook.sheetnames) == 1


@pytest.mark.parametrize(
    "code,header,text",
    [
        ("sr_Latn", "sr-latn", "Zdravo"),
        ("la", "Latin(la)", "Salve"),
        ("zh_Hant_HK", "zh-hant-hk", "你好"),
    ],
)
@pytest.mark.parametrize("as_source", [False, True])
def test_catalog_locale_reuses_original_column_and_preserves_translation(
    tmp_path, code, header, text, as_source
):
    from app.db import Store
    from app.schemas import UploadPatch
    from app.services.drafts import create_draft, patch_draft
    from app.services.exports import build_export

    store = Store(tmp_path)
    rows = [["Key", "ru", header], ["one", "Привет", text]]
    kit = "\n".join(",".join(row) for row in rows) + "\n"
    draft = create_draft(store, "owner", "local", [("kit.csv", kit.encode())], "string", "quick")
    source, target = (code, "ru") if as_source else ("ru", code)
    saved = patch_draft(
        store,
        "owner",
        draft["token"],
        UploadPatch(
            revision=draft["revision"],
            ui_mode="quick",
            current_step="confirm",
            answers={
                "source_language": source,
                "included_target_languages": [target],
                "selected_sheets": [{"file": "kit.csv", "sheet": "kit"}],
            },
        ),
    )
    assert saved["errors"] == []
    manifest = saved["manifest"]
    assert manifest["source"] == source
    assert manifest["targets"] == [target]
    component = manifest["components"][0]
    assert component["profile"]["source_lang"] == source
    assert component["units"][0]["values"] == {"ru": "Привет", code: text}
    assert build_export(manifest, {})["sheets"][0]["rows"] == rows
    assert component["profile"]["comments"] == []


def test_new_regional_targets_add_columns_without_replacing_base_translations(tmp_path):
    from app.db import Store
    from app.schemas import UploadPatch
    from app.services.drafts import create_draft, patch_draft
    from app.services.exports import build_export

    store = Store(tmp_path)
    kit = "Key,ru,pt,es,de\nold,Привет,Olá,Hola,Hallo\n"
    draft = create_draft(store, "owner", "local", [("kit.csv", kit.encode())], "string", "quick")
    targets = ["pt_BR", "es_MX", "zh_Hans"]
    saved = patch_draft(
        store,
        "owner",
        draft["token"],
        UploadPatch(
            revision=draft["revision"],
            ui_mode="quick",
            current_step="confirm",
            answers={
                "source_language": "ru",
                "included_target_languages": targets,
                "selected_sheets": [{"file": "kit.csv", "sheet": "kit"}],
            },
        ),
    )
    assert saved["errors"] == []
    assert saved["manifest"]["targets"] == targets
    export = build_export(saved["manifest"], {})
    assert export["sheets"][0]["rows"] == [
        ["Key", "ru", "pt", "es", "de", *targets],
        ["old", "Привет", "Olá", "Hola", "Hallo", "", "", ""],
    ]


@pytest.mark.parametrize("extra", [False, True])
def test_rare_source_and_empty_target_keep_headers_without_internal_codes(extra):
    from app.services.kits import infer_selected_component

    rows = [["Key", "la", "sr-latn"], ["one", "Salve", ""]]
    if extra:
        rows[0].append("en_XW0")
        rows[1].append("Existing")
    original = [list(row) for row in rows]
    document = infer_selected_component("kit", rows, "kit-1", "la", ["sr_Latn"])
    assert document["source_lang"] == "la"
    assert document["languages"][0]["code"] == "la"
    assert document["languages"][0]["header"] == "la"
    assert rows == original
    if extra:
        assert any(language["code"] == "en_XW0" for language in document["languages"])


def test_unrecognized_sheets_report_problems_before_selection_and_can_be_excluded(tmp_path):
    from app.db import Store
    from app.schemas import UploadPatch
    from app.services.drafts import create_draft, patch_draft
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.active.title = "UI"
    workbook.active.append(["Key", "ru", "en"])
    workbook.active.append(["play", "Играть", "Play"])
    questions = workbook.create_sheet("Questions_Answers")
    questions.append(["Tab", "Line", "Key", "Text", "Question/Comment", "Answer"])
    general = workbook.create_sheet("General")
    general.append(["Russian", "English"])
    hidden = workbook.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden.append(["Text", "Notes"])
    path = tmp_path / "kit.xlsx"
    workbook.save(path)
    store = Store(tmp_path / "data")
    draft = create_draft(
        store, "producer", "local", [("kit.xlsx", path.read_bytes())], "string", "quick"
    )
    sheets = {s["name"]: s for s in draft["analysis"]["loc_kit"]["files"][0]["sheets"]}
    assert set(sheets) == {"UI", "Questions_Answers", "General"}
    assert sheets["UI"].get("issues", []) == []
    for name in ["Questions_Answers", "General"]:
        assert sheets[name]["issues"][0]["code"] == "loc_kit.no_language_headers"
        assert sheets[name]["issues"][0]["scope"] == "sheet"
    saved = patch_draft(
        store,
        "producer",
        draft["token"],
        UploadPatch(
            revision=draft["revision"],
            ui_mode="quick",
            current_step="confirm",
            answers={
                "selected_sheets": [{"file": "kit.xlsx", "sheet": "UI"}],
                "source_language": "ru",
                "included_target_languages": ["en"],
            },
        ),
    )
    assert not saved["errors"]
    assert [c["sheet"] for c in saved["manifest"]["components"]] == ["UI"]


def test_empty_existing_target_column_is_filled_in_place(tmp_path):
    from app.db import Store
    from app.schemas import UploadPatch
    from app.services.drafts import create_draft, patch_draft
    from app.services.exports import build_export

    store = Store(tmp_path)
    draft = create_draft(
        store,
        "producer",
        "local",
        [("kit.csv", b"Key,en,Russian(ru)\none,Hello,\n")],
        "string",
        "quick",
    )
    draft = patch_draft(
        store,
        "producer",
        draft["token"],
        UploadPatch(
            revision=draft["revision"],
            ui_mode="quick",
            current_step="confirm",
            answers={
                "selected_sheets": [{"file": "kit.csv", "sheet": "kit"}],
                "source_language": "en",
                "included_target_languages": ["ru"],
            },
        ),
    )
    assert not draft["errors"]
    exported = build_export(draft["manifest"], {"kit-1": {"ru": {"one": "Привет"}}})
    assert exported["sheets"][0]["rows"] == [
        ["Key", "en", "Russian(ru)"],
        ["one", "Hello", "Привет"],
    ]


def test_hidden_sheets_excluded_before_analysis(tmp_path):
    from loc_kit_ingest.reader import inspect_sheets
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.active.title = "Visible"
    workbook.active.append(["Key", "English(en)", "Russian(ru)"])
    workbook.active.append(["hello", "Hello", "Привет"])
    for name, state in [("Secret", "hidden"), ("Other", "veryHidden")]:
        sheet = workbook.create_sheet(name)
        sheet.sheet_state = state
        sheet.append(["Key", "English(en)"])
        sheet.append(["private", "DO_NOT_PROCESS"])
    path = tmp_path / "kit.xlsx"
    workbook.save(path)
    result = inspect_sheets(path)
    assert list(result.sheets) == ["Visible"]
    assert "DO_NOT_PROCESS" not in repr(result)
    from app.db import DomainError, Store
    from app.schemas import UploadPatch
    from app.services.drafts import create_draft, patch_draft
    from app.services.exports import build_export

    store = Store(tmp_path / "data")
    draft = create_draft(
        store, "producer", "test-local", [("kit.xlsx", path.read_bytes())], "string", "quick"
    )
    answers = {
        "selected_files": ["kit.xlsx"],
        "selected_sheets": [{"file": "kit.xlsx", "sheet": "Visible"}],
        "source_language": "en",
        "included_target_languages": ["ru"],
    }
    saved = patch_draft(
        store,
        "producer",
        draft["token"],
        UploadPatch(
            revision=draft["revision"], ui_mode="quick", current_step="confirm", answers=answers
        ),
    )
    assert not saved["errors"]
    assert "DO_NOT_PROCESS" not in repr(saved["manifest"])
    assert [sh["name"] for sh in build_export(saved["manifest"], {})["sheets"]] == ["Visible"]
    import pytest

    for hidden in ["Secret", "Other"]:
        answers["selected_sheets"] = [{"file": "kit.xlsx", "sheet": hidden}]
        with pytest.raises(DomainError) as error:
            patch_draft(
                store,
                "producer",
                draft["token"],
                UploadPatch(
                    revision=saved["revision"],
                    ui_mode="quick",
                    current_step="confirm",
                    answers=answers,
                ),
            )
        assert error.value.code == "sheet_not_visible" and error.value.status == 400


def test_upload_confirmation_and_export_preserve_targets(tmp_path):
    from app.main import create_app
    from app.settings import Settings
    from fastapi.testclient import TestClient

    app = create_app(Settings(data_dir=tmp_path, public_origin="http://testserver"))
    app.state.store.create_user("producer", "a-long-test-password")
    with TestClient(app) as client:
        csrf = client.post(
            "/auth/login", json={"username": "producer", "password": "a-long-test-password"}
        ).json()["csrf_token"]
        headers = {"X-CSRFToken": csrf, "Origin": "http://testserver"}
        slug = client.post(
            "/api/producer/projects/", json={"name": "тест"}, headers=headers
        ).json()["slug"]
        upload = client.post(
            f"/api/producer/projects/{slug}/uploads/",
            headers=headers,
            files=[
                (
                    "files",
                    (
                        "kit.csv",
                        b"Key,English(en),Russian(ru)\nhello,Hello,Already done\nbye,Goodbye,\n",
                        "text/csv",
                    ),
                )
            ],
        )
        assert upload.status_code == 201
        draft = upload.json()
        sheet = draft["analysis"]["loc_kit"]["files"][0]["sheets"][0]["name"]
        answers = {
            "selected_files": ["kit.csv"],
            "selected_sheets": [{"file": "kit.csv", "sheet": sheet}],
            "source_language": "en",
            "included_target_languages": ["ru"],
            "column_roles": [],
            "duplicate_policy": "separate-components",
            "profile_answers": {},
        }
        body = {
            "revision": draft["revision"],
            "answers": answers,
            "ui_mode": "quick",
            "current_step": "confirm",
        }
        confirmed = client.patch(
            f"/api/producer/uploads/{draft['token']}/", json=body, headers=headers
        )
        assert confirmed.status_code == 200
        assert confirmed.json()["analysis"]["confirmation"]["totals"]["imported_rows"] == 2
        assert (
            client.patch(
                f"/api/producer/uploads/{draft['token']}/", json=body, headers=headers
            ).status_code
            == 409
        )
        assert confirmed.json()["errors"] == []
        manifest = app.state.store.get("draft", draft["token"], "producer")["manifest"]
        assert manifest["components"][0]["units"][0]["values"]["ru"] == "Already done"


def test_formula_cannot_become_translation_text(tmp_path):
    from app.db import Store
    from app.schemas import UploadPatch
    from app.services.drafts import create_draft, patch_draft
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.active.title = "Kit"
    workbook.active.append(["Key", "en", "ru"])
    workbook.active.append(["one", "=1+1", ""])
    path = tmp_path / "formula.xlsx"
    workbook.save(path)
    store = Store(tmp_path / "data")
    draft = create_draft(
        store, "producer", "test-local", [("kit.xlsx", path.read_bytes())], "string", "quick"
    )
    saved = patch_draft(
        store,
        "producer",
        draft["token"],
        UploadPatch(
            revision=draft["revision"],
            ui_mode="quick",
            current_step="confirm",
            answers={
                "selected_sheets": [{"file": "kit.xlsx", "sheet": "Kit"}],
                "source_language": "en",
                "included_target_languages": ["ru"],
            },
        ),
    )
    assert saved["errors"] and saved["manifest"] is None


def test_incomplete_rows_and_duplicate_keys_go_to_quarantine(tmp_path):
    from app.db import Store
    from app.schemas import UploadPatch
    from app.services.drafts import create_draft, patch_draft

    store = Store(tmp_path / "data")
    payload = b"Key,en,ru\ngood,Hello,Keep\nduplicate,First,\nduplicate,Second,\nempty,,Existing\n"
    draft = create_draft(store, "producer", "test-local", [("kit.csv", payload)], "string", "quick")
    saved = patch_draft(
        store,
        "producer",
        draft["token"],
        UploadPatch(
            revision=draft["revision"],
            ui_mode="quick",
            current_step="confirm",
            answers={
                "selected_sheets": [{"file": "kit.csv", "sheet": "kit"}],
                "source_language": "en",
                "included_target_languages": ["ru"],
            },
        ),
    )
    assert not saved["errors"]
    manifest = saved["manifest"]
    assert manifest["totals"] == {"input_rows": 4, "imported_rows": 2, "quarantined_rows": 2}
    assert manifest["components"][0]["units"][0]["key"] == "good"
    assert {q["reason_code"] for q in manifest["quarantine"]} == {"duplicate_key"}
