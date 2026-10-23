# Copyright © HCGameLoc
# SPDX-License-Identifier: GPL-3.0-or-later
from __future__ import annotations

import copy

import pytest
from test_quality_artifacts import artifact_inputs, build


@pytest.mark.parametrize("seats", [{}, {"1": "parsed", "2": "parsed"}])
def test_old_major_with_new_response_revision_is_yellow_without_current_verdict_proof(seats):
    # Real API: run severity is historical; response revision is computed from the live unit.
    values = artifact_inputs()
    values["final_snapshot"]["judged_revisions"].pop("12")
    values["unit_results"] = [
        {"unit_id": 12, "severity": "major", "input_revision": "r2", "live_seat_results": seats}
    ]
    result = build(values)
    assert result["highlights"] == [
        {"sheet_id": "c", "row": 4, "column": 3, "category": "unverified"}
    ]
    assert [i["kind"] for i in result["issues"]] == ["unverified"]


def test_available_reason_and_check_id_survive_for_row_and_project_issues():
    values = artifact_inputs()
    issue = {
        "group_id": "coverage",
        "kind": "check",
        "hard": True,
        "advisory": False,
        "check_id": "repeat-coverage",
        "reason": "translation_propagation_disabled",
        "unit_ids": [11],
    }
    values["final_snapshot"]["issues"] = [
        issue,
        {**issue, "group_id": "project-coverage", "unit_ids": []},
    ]
    result = build(values)
    assert result["issues"][0]["check_id"] == "repeat-coverage"
    assert result["issues"][0]["reason"] == "translation_propagation_disabled"
    assert result["issues"][0]["action"] == "check_component_repeat_settings"
    assert result["summary"]["project_issues"][0]["reason"] == "translation_propagation_disabled"
    assert result["summary"]["project_issues"][0]["check_id"] == "repeat-coverage"


def test_report_uses_actual_output_tab_when_original_names_collide():
    values = artifact_inputs()
    second = copy.deepcopy(values["manifest"]["components"][0])
    second.update(id="c2", file="second.xlsx")
    values["manifest"]["components"].append(second)
    sheet = copy.deepcopy(values["exported"]["sheets"][0])
    sheet.update(id="c2", file="second.xlsx")
    values["exported"]["sheets"].append(sheet)
    values["bindings"].extend(
        [{**b, "component_id": "c2", "sheet_id": "c2"} for b in list(values["bindings"])]
    )
    values["final_snapshot"]["judged_revisions"].pop("12")
    rows = build(values)["issues"]
    assert [(i["sheet"], i["original_sheet"]) for i in rows] == [
        ("Strings", "Strings"),
        ("Strings-2", "Strings"),
    ]
    import io

    from app.services.quality_artifact_files import (
        serialize_quality_report,
        serialize_quality_translation,
    )
    from openpyxl import load_workbook

    artifact = build(values)
    file, _, _ = serialize_quality_translation(values["exported"], artifact)
    workbook = load_workbook(io.BytesIO(file))
    assert workbook.worksheets[1].title == rows[1]["sheet"] == "Strings-2"
    file, _, _ = serialize_quality_report(artifact, values["exported"])
    report = load_workbook(io.BytesIO(file))
    assert list(report["Замечания"].values)[2][1] == "Strings-2"


def test_quality_export_excludes_service_columns_when_finding_added_target_language():
    from app.services.exports import build_export
    from app.services.quality_artifacts import bind_quality_cells

    values = artifact_inputs()
    manifest = values["manifest"]
    manifest.update(manifest_version=2, quarantine=[])
    component = manifest["components"][0]
    component.update(logical_sheet_id="logical", header_row=0)
    profile = {
        "header_row": 1,
        "languages": [{"code": "en", "column": 2}],
        "key": {"column": 1},
        "ignored_columns": [{"column": 3}],
    }
    manifest["sheets"] = [
        {
            "id": "logical",
            "file": "kit.xlsx",
            "sheet": "Strings",
            "profile": profile,
            "rows": [
                ["key", "en", "de"],
                ["one", "Play", "service"],
                ["", "", ""],
                ["two", "Play", "service"],
            ],
        }
    ]
    exported = build_export(manifest, {"c": {"de": {"one": "Spielen", "two": "Spiel"}}}, True)
    assert exported["sheets"][0]["rows"][1] == ["one", "Play", "service", "Spielen"]
    remote = {
        "issues": [],
        "components": {
            "remote": {
                "complete": True,
                "source": "en",
                "units": [
                    {
                        "key": "one",
                        "source": "Play",
                        "target_ids": {"de": 11},
                        "targets": {"de": "Spielen"},
                    },
                    {
                        "key": "two",
                        "source": "Play",
                        "target_ids": {"de": 12},
                        "targets": {"de": "Spiel"},
                    },
                ],
            }
        },
    }
    bindings = bind_quality_cells(
        manifest,
        exported,
        remote,
        {"kit.xlsx\u001fStrings": "remote"},
        values["final_snapshot"],
        "de",
    )
    assert [b["column"] for b in bindings] == [4, 4]
