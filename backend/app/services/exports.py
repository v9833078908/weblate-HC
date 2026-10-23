# Copyright © HCGameLoc
# SPDX-License-Identifier: GPL-3.0-or-later
from __future__ import annotations

import copy
import csv
import io
import json
import re
import zipfile

from loc_kit_ingest.infer import language_code
from openpyxl import Workbook
from sqlalchemy import select

from app.db import DomainError, revision
from app.models import Job
from app.services.kits import selected_header_language
from app.services.languages import normalized_code
from app.services.manifest import component_mapping_key, component_source


def _build_legacy_export(manifest, translations, prefer_remote=False):
    sheets = []
    per_language = {
        code: {"strings": 0, "translated": 0, "empty": 0, "blocked": 0}
        for code in manifest["targets"]
    }
    keys = 0
    for component in manifest["components"]:
        rows = copy.deepcopy(component["rows"])
        languages = {
            language["code"]: language["column"] - 1
            for language in component["profile"]["languages"]
        }
        for unit in component["units"]:
            keys += 1
            for code in manifest["targets"]:
                original = unit["values"].get(code, "")
                remote = translations.get(component["id"], {}).get(code, {}).get(unit["key"], "")
                target = (remote or original) if prefer_remote else (original or remote)
                row = rows[unit["row"] - 1]
                column = languages[code]
                row.extend([""] * max(0, column + 1 - len(row)))
                row[column] = target
                counts = per_language[code]
                counts["strings"] += 1
                counts["translated"] += int(bool(target))
                counts["empty"] += int(not target)
        sheets.append(
            {
                "id": component["id"],
                "file": component["file"],
                "name": component["sheet"],
                "rows": rows,
            }
        )
    return {
        "sheets": sheets,
        "keys": keys,
        "empty": sum(v["empty"] for v in per_language.values()),
        "removed": 0,
        "per_language": per_language,
        "quarantine": manifest["quarantine"],
        "formats": ["xlsx", "csv", "json"],
    }


def build_export(manifest, translations, prefer_remote=False):
    if manifest.get("manifest_version") != 2:
        return _build_legacy_export(manifest, translations, prefer_remote)
    sheets = {}
    columns, reserved_columns = {}, {}
    for sheet in manifest.get("sheets", []):
        rows = copy.deepcopy(sheet["rows"])
        sheets[sheet["id"]] = {
            "id": sheet["id"],
            "file": sheet["file"],
            "name": sheet["sheet"],
            "rows": rows,
        }
        columns[sheet["id"]] = {
            lang["code"]: lang["column"] - 1 for lang in sheet["profile"]["languages"]
        }
        reserved_columns[sheet["id"]] = reserved_export_columns(sheet["profile"])
    counts = {
        code: {"strings": 0, "translated": 0, "empty": 0, "blocked": 0}
        for code in manifest["targets"]
    }
    writes, keys = {}, set()
    for component in manifest["components"]:
        logical = component["logical_sheet_id"]
        rows, lang_cols = sheets[logical]["rows"], columns[logical]
        for code in manifest["targets"]:
            if code not in lang_cols:
                existing = [
                    i
                    for i, h in enumerate(rows[component["header_row"]])
                    if selected_header_language(h, {normalized_code(code): code}) == code
                    and i not in lang_cols.values()
                    and i not in reserved_columns[logical]
                ]
                if len(existing) > 1:
                    raise DomainError("export_columns", "Неоднозначные языковые колонки", 409)
                col = existing[0] if existing else max(map(len, rows), default=0)
                lang_cols[code] = col
                for row in rows:
                    row.extend([""] * (col + 1 - len(row)))
                if not existing:
                    rows[component["header_row"]][col] = code
        actual = component_source(manifest, component)
        for unit in component["units"]:
            row_number = unit.get("original_row", unit["row"])
            keys.add((logical, unit["key"]))
            for code in manifest["targets"]:
                original = unit["values"].get(code, "")
                remote = translations.get(component["id"], {}).get(code, {}).get(unit["key"], "")
                target = (
                    original
                    if code == actual
                    else ((remote or original) if prefer_remote else (original or remote))
                )
                coordinate = (logical, row_number, lang_cols[code])
                if coordinate in writes and writes[coordinate] != target:
                    raise DomainError(
                        "export_cell_conflict",
                        "Несогласованные результаты для одной ячейки лок-кита",
                        409,
                    )
                if coordinate in writes:
                    continue
                writes[coordinate] = target
                if code != actual:
                    rows[row_number - 1].extend(
                        [""] * max(0, lang_cols[code] + 1 - len(rows[row_number - 1]))
                    )
                    rows[row_number - 1][lang_cols[code]] = target
                count = counts[code]
                count["strings"] += 1
                count["translated"] += int(bool(target))
                count["empty"] += int(not target)
    return {
        "sheets": list(sheets.values()),
        "keys": len(keys),
        "empty": sum(c["empty"] for c in counts.values()),
        "removed": 0,
        "per_language": counts,
        "quarantine": manifest["quarantine"],
        "formats": ["xlsx", "csv", "json"],
    }


def component_exports(manifest, translations, prefer_remote=False):
    """User selections address logical sheets, never internal source groups."""
    selections = (
        manifest.get("sheets", [])
        if manifest.get("manifest_version") == 2
        else manifest["components"]
    )
    parts = {}
    for sheet in selections:
        groups = [
            c for c in manifest["components"] if c.get("logical_sheet_id", c["id"]) == sheet["id"]
        ]
        scoped = {
            **manifest,
            "components": groups,
            "quarantine": [
                q
                for q in manifest["quarantine"]
                if q["file"] == sheet["file"] and q["sheet"] == sheet["sheet"]
            ],
        }
        if manifest.get("manifest_version") == 2:
            scoped["sheets"] = [sheet]
        parts[sheet["id"]] = build_export(scoped, translations, prefer_remote)
    return parts


def project_snapshot(store, owner, slug):
    """Combine completed language results without changing the saved run exports."""
    from app.services.project_links import get_link
    from app.services.runs import source_identity

    link = get_link(store, owner, slug)
    if link and link.get("snapshot", {}).get("components"):
        # A project download reflects the last explicit refresh; run exports stay immutable.
        if link.get("read_only"):
            return reconstructed_export(link["snapshot"])
        try:
            pointer = store.get("export_pointer", "export-pointer:" + slug, owner)
            with store.session() as session:
                job = session.get(Job, pointer["run_id"])
                manifest = (
                    copy.deepcopy(job.data["snapshot"]["manifest"])
                    if job and job.owner == owner and job.status == "completed"
                    else None
                )
            if manifest and all(
                component_mapping_key(manifest, c) in link.get("component_map", {})
                for c in manifest["components"]
            ):
                return version_export(manifest, link["snapshot"], link["component_map"])
        except DomainError as error:
            if error.status != 404:
                raise
        return reconstructed_export(link["snapshot"])

    try:
        pointer = store.get("export_pointer", "export-pointer:" + slug, owner)
    except DomainError as error:
        if error.status != 404:
            raise
        from app.services.project_links import get_link

        link = get_link(store, owner, slug)
        if not link:
            raise
        return reconstructed_export(link["snapshot"])
    value = store.get("export", pointer["export_id"], owner)
    value.pop("revision")
    with store.session() as session:
        latest = session.get(Job, pointer["run_id"])
        if not latest or latest.owner != owner or latest.status != "completed":
            return value
        latest_data = latest.data
        if latest_data["snapshot"].get("change_plan"):
            return value
        identity = source_identity(latest_data["snapshot"]["manifest"])
        history = sorted(
            (
                job.data
                for job in session.scalars(
                    select(Job).where(Job.owner == owner, Job.status == "completed")
                )
                if job.data["project"] == slug
                and job.data["remote_project"] == latest_data["remote_project"]
                and (job.data.get("finished") or job.data["created"])
                <= (latest_data.get("finished") or latest_data["created"])
                and job.data.get("export_id")
                and source_identity(job.data["snapshot"]["manifest"]) == identity
            ),
            key=lambda data: (
                data.get("finished") or data["created"],
                data["export_id"] == pointer["export_id"],
                data["export_id"],
            ),
        )
    if len(history) < 2:
        return value
    manifest = copy.deepcopy(latest_data["snapshot"]["manifest"])
    manifest["targets"] = list(
        dict.fromkeys(code for data in history for code in data["snapshot"]["manifest"]["targets"])
    )
    components = {component["id"]: component for component in manifest["components"]}
    for component in components.values():
        for unit in component["units"]:
            unit["values"] = {
                component_source(manifest, component): unit["values"][
                    component_source(manifest, component)
                ]
            }
    for data in history:
        previous_manifest = data["snapshot"]["manifest"]
        exported = store.get("export", data["export_id"], owner)
        sheets = {sheet["id"]: sheet for sheet in exported["sheets"]}
        for previous in previous_manifest["components"]:
            component = components[previous["id"]]
            profile = component["profile"]
            languages = {lang["code"]: lang for lang in profile["languages"]}
            header = component["rows"][profile["header_row"] - 1]
            units = {unit["key"]: unit for unit in component["units"]}
            old_columns = {
                lang["code"]: lang["column"] - 1 for lang in previous["profile"]["languages"]
            }
            for code in previous_manifest["targets"]:
                if code not in languages:
                    reserved = {
                        field["column"]
                        for fields in profile.values()
                        for field in (fields if isinstance(fields, list) else [fields])
                        if isinstance(field, dict) and "column" in field
                    }
                    column = next(
                        (
                            i
                            for i, cell in enumerate(header, 1)
                            if i not in reserved and language_code(cell) == code
                        ),
                        None,
                    )
                    if column is None:
                        column = max(map(len, component["rows"]), default=0) + 1
                        for row in component["rows"]:
                            row.extend([""] * (column - len(row)))
                        header[column - 1] = code
                    lang = {"code": code, "column": column, "header": header[column - 1]}
                    profile["languages"].append(lang)
                    languages[code] = lang
                for old_unit in previous["units"]:
                    row = sheets[previous.get("logical_sheet_id", previous["id"])]["rows"][
                        old_unit["row"] - 1
                    ]
                    units[old_unit["key"]]["values"][code] = row[old_columns[code]]
    combined = build_export(manifest, {})
    combined["components"] = component_exports(manifest, {})
    return combined


def reconstructed_export(snapshot):
    parts = {}
    for cid, component in snapshot["components"].items():
        if component.get("glossary"):
            continue
        codes = sorted(
            {code for unit in component["units"] for code in unit["targets"]}
            | set(snapshot["languages"])
        )
        rows = [["key", component["source"], *codes, "note"]]
        counts = {
            code: {
                "strings": len(component["units"]),
                "translated": sum(bool(u["targets"].get(code)) for u in component["units"]),
                "empty": sum(not u["targets"].get(code) for u in component["units"]),
                "blocked": 0,
            }
            for code in codes
        }
        rows.extend(
            [
                [
                    u["key"],
                    u["source"],
                    *[u["targets"].get(code, "") for code in codes],
                    u.get("note", ""),
                ]
                for u in component["units"]
            ]
        )
        parts[cid] = {
            "sheets": [
                {"id": cid, "file": "reconstructed.xlsx", "name": component["name"], "rows": rows}
            ],
            "keys": len(component["units"]),
            "empty": sum(c["empty"] for c in counts.values()),
            "removed": 0,
            "per_language": counts,
            "quarantine": [],
            "formats": ["xlsx", "csv", "json"],
        }
    codes = sorted(snapshot["languages"])
    return {
        "components": parts,
        "sheets": [s for part in parts.values() for s in part["sheets"]],
        "keys": sum(p["keys"] for p in parts.values()),
        "empty": sum(p["empty"] for p in parts.values()),
        "removed": 0,
        "per_language": {
            code: {
                field: sum(p["per_language"].get(code, {}).get(field, 0) for p in parts.values())
                for field in ("strings", "translated", "empty", "blocked")
            }
            for code in codes
        },
        "quarantine": [],
        "formats": ["xlsx", "csv", "json"],
        "reconstructed": True,
    }


def version_export(manifest, snapshot, mapping):
    """Enrich a new version with all observed locales; old exports stay immutable."""

    manifest = copy.deepcopy(manifest)
    manifest["targets"] = sorted(set(manifest["targets"]) | set(snapshot["languages"]))
    translations = {}
    for component in manifest["components"]:
        remote = snapshot["components"][mapping[component_mapping_key(manifest, component)]]
        for unit in component["units"]:
            unit["values"] = {
                component_source(manifest, component): unit["values"][
                    component_source(manifest, component)
                ]
            }
        profile = component["profile"]
        languages = {lang["code"] for lang in profile["languages"]}
        header = component["rows"][profile["header_row"] - 1]
        for code in manifest["targets"]:
            if code not in languages:
                column = max(map(len, component["rows"])) + 1
                for row in component["rows"]:
                    row.extend([""] * (column - len(row)))
                header[column - 1] = code
                profile["languages"].append({"code": code, "column": column, "header": code})
        translations[component["id"]] = {
            code: {u["key"]: u["targets"].get(code, "") for u in remote["units"]}
            for code in manifest["targets"]
        }
    value = build_export(manifest, translations, True)
    value["components"] = component_exports(manifest, translations, True)
    return value


def snapshot(store, owner, slug, content):
    value = project_snapshot(store, owner, slug)
    available = {sheet["id"] for sheet in value["sheets"]}
    if content and not set(content) <= available:
        raise DomainError("export_selection", "Unknown content")
    if content:
        # Use saved per-component snapshots so the summary always matches the selected files.
        selected = [value["components"][key] for key in content]
        value = {
            "sheets": [sh for part in selected for sh in part["sheets"]],
            "keys": sum(part["keys"] for part in selected),
            "empty": sum(part["empty"] for part in selected),
            "removed": 0,
            "formats": ["xlsx", "csv", "json"],
            "per_language": {
                code: {
                    field: sum(part["per_language"][code][field] for part in selected)
                    for field in ["strings", "translated", "empty", "blocked"]
                }
                for code in value["per_language"]
            },
            "quarantine": [q for part in selected for q in part["quarantine"]],
        }
    else:
        value.pop("components", None)
    return value


def export_summary(store, owner, slug, content):
    value = snapshot(store, owner, slug, content)
    return {key: value[key] for key in ["keys", "empty", "removed", "per_language", "formats"]} | {
        "revision": revision(value)
    }


def export_file(store, owner, slug, content, format, expected):
    value = snapshot(store, owner, slug, content)
    if revision(value) != expected:
        raise DomainError("stale-revision", "Reload the export summary", 409)
    return serialize_export(value, slug, format)


def serialize_export(value, slug, format):
    if format == "json":
        return (
            json.dumps(value, ensure_ascii=False).encode(),
            "application/json",
            slug + "-loc-kit.json",
        )
    if format == "xlsx":
        workbook, _sheets = export_workbook(value)
        output = io.BytesIO()
        workbook.save(output)
        return (
            output.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            slug + "-loc-kit.xlsx",
        )
    if format == "csv":
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
            for sheet in value["sheets"]:
                text = io.StringIO()
                writer = csv.writer(text)
                # CSV has no cell types: neutralise executable spreadsheet expressions.
                writer.writerows(
                    [
                        [
                            "'" + cell
                            if isinstance(cell, str) and cell.startswith(("=", "+", "-", "@"))
                            else cell
                            for cell in row
                        ]
                        for row in sheet["rows"]
                    ]
                )
                archive.writestr(sheet["id"] + ".csv", text.getvalue().encode("utf-8-sig"))
        return output.getvalue(), "application/zip", slug + "-loc-kit-csv.zip"
    raise DomainError("export_format", "Choose xlsx, csv or json")


def reserved_export_columns(profile):
    """Zero-based service columns that can never become translation output."""
    return {
        field["column"] - 1
        for name in ("key", "comments", "references", "ignored_columns", "flags", "explanation")
        for field in (
            profile.get(name, []) if isinstance(profile.get(name, []), list) else [profile[name]]
        )
        if "column" in field
    }


def export_sheet_titles(value):
    """Use one identity-to-title map for the workbook and its quality report."""
    titles, used = {}, set()
    for sheet in value["sheets"]:
        base = re.sub(r"[\\/*?:\[\]]", "_", sheet["name"])[:27] or "Sheet"
        title, suffix = base, 2
        while title.casefold() in used:
            tail = f"-{suffix}"
            title = base[: 31 - len(tail)] + tail
            suffix += 1
        used.add(title.casefold())
        titles[sheet["id"]] = title
    return titles


def export_workbook(value):
    """Build literal-value sheets and retain identity despite renamed Excel tabs."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    titles, sheets = export_sheet_titles(value), {}
    for sheet in value["sheets"]:
        ws = workbook.create_sheet(titles[sheet["id"]])
        sheets[sheet["id"]] = ws
        for row in sheet["rows"]:
            ws.append(row)
            for cell in ws[ws.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
    return workbook, sheets
