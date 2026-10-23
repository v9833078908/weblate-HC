# Copyright © HCGameLoc
# SPDX-License-Identifier: GPL-3.0-or-later
from __future__ import annotations

import hashlib
import io
import json
import re
import zipfile
from datetime import datetime

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.db import DomainError
from app.services.exports import export_workbook

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
COLORS = {"blocking": "FFFFC0CB", "unverified": "FFFFF2CC"}
LABELS = {
    "blocking": "Критическое замечание",
    "unverified": "Проверка не завершена",
    "advisory": "Рекомендация",
    "glossary": "Глоссарий",
    "repeat-drift": "Разные переводы повторов",
    "check": "Техническая проверка",
    "judge": "Замечание судьи",
    "parsed": "Ответ разобран (не означает одобрение)",
    "missing": "Нет актуального ответа",
    "approved": "Одобрено для текущей версии",
    "translation_propagation_disabled": "Контроль повторов выключен для компонента",
    "check_component_repeat_settings": "Проверить настройки контроля повторов компонента",
    "unparsed": "Ответ не разобран",
    "manual_review": "Проверить и исправить вручную",
    "complete_verification": "Завершить проверку",
    "review_recommendation": "Рассмотреть рекомендацию",
    "current_revision_not_verified": "Нет подтверждённого одобрения обоих судей для текущей версии",
    "judge_major": "Существенное замечание судьи к текущей версии",
    "judge_critical": "Критическое замечание судьи к текущей версии",
    "completed": "Завершено",
    "completed_with_exceptions": "Завершено с замечаниями",
    "required": "Обязательное",
    "not_verified": "Не проверено",
}


def label(value):
    return LABELS.get(value, value)


def _name(artifact, suffix):
    run_id = artifact["run_id"]
    if not re.fullmatch(r"[A-Za-z0-9_-]{1,100}", run_id):
        raise DomainError("quality_filename", "Invalid result identity", 409)
    return f"run-{run_id}-{suffix}"


def _check_values(rows):
    for row in rows:
        for value in row:
            if isinstance(value, str) and (
                len(value) > 32767 or ILLEGAL_CHARACTERS_RE.search(value)
            ):
                raise DomainError(
                    "quality_excel_limit", "Report text exceeds Excel cell limits", 409
                )


def _save(workbook):
    # Stable file bytes: no download-time timestamps in a frozen result bundle.
    workbook.properties.created = workbook.properties.modified = datetime(2000, 1, 1)
    initial, output = io.BytesIO(), io.BytesIO()
    workbook.save(initial)
    with (
        zipfile.ZipFile(initial) as source,
        zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, (2000, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            data = source.read(name)
            if name == "docProps/core.xml":
                # openpyxl updates modified during save, even when set beforehand.
                data = re.sub(
                    rb"(<dcterms:modified[^>]*>).*?(</dcterms:modified>)",
                    rb"\g<1>2000-01-01T00:00:00Z\2",
                    data,
                )
            target.writestr(info, data)
    return output.getvalue()


def serialize_quality_translation(exported: dict, artifact: dict) -> tuple[bytes, str, str]:
    for sheet in exported["sheets"]:
        _check_values(sheet["rows"])
    workbook, sheets = export_workbook(exported)
    for mark in artifact["highlights"]:
        cell = sheets[mark["sheet_id"]].cell(mark["row"], mark["column"])
        cell.fill = PatternFill("solid", fgColor=COLORS[mark["category"]])
        cell.font = Font(color="FF111111")
    return _save(workbook), XLSX, _name(artifact, "translation.xlsx")


def serialize_quality_report(artifact: dict, exported: dict) -> tuple[bytes, str, str]:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Сводка"
    rows = [
        ["Показатель", "Значение"],
        ["Задание", artifact["run_id"]],
        ["Язык", artifact["summary"]["language"]],
        ["Версия результата", artifact["provenance"]["scope_revision"]],
        ["Статус", label(artifact["summary"]["status"])],
    ]
    names = {
        "processed": "Обработано строк",
        "automatically_corrected": "Автоматически исправлено строк",
        "blocking_cells": "Ячеек с критическими замечаниями",
        "unverified_cells": "Непроверенных ячеек",
        "highlighted_cells": "Всего выделено ячеек",
        "issue_count": "Остаточных замечаний",
        "repeat_groups": "Групп повторов с замечаниями",
        "judge_verified_units": "Строк одобрено обоими судьями",
        "quarantined_transport": "Изолированных ошибок сервиса",
    }
    rows.extend([title, artifact["summary"].get(key, 0)] for key, title in names.items())
    for seat, states in artifact["summary"].get("judge_seats", {}).items():
        rows.append(
            [
                f"Судья {seat}",
                ", ".join(f"{label(state)}: {count}" for state, count in sorted(states.items())),
            ]
        )
    rows.extend(
        [
            ["Розовый", "Критические обязательные замечания"],
            ["Жёлтый", "Непроверенные ячейки; при критической ошибке приоритет у розового"],
            ["Ответ разобран", "Не означает одобрение перевода"],
            [
                "Результат",
                "Остаточных замечаний нет"
                if not artifact["issues"]
                and not artifact["summary"].get("project_issues")
                and not artifact["summary"].get("quarantined_transport")
                else "См. замечания и состояние проверки",
            ],
        ]
    )
    for issue in artifact["summary"].get("project_issues", []):
        rows.append(
            [
                "Проблема проекта",
                issue["group_id"] + ": " + label(issue.get("reason", issue["kind"])),
            ]
        )
    _check_values(rows)
    for row in rows:
        summary.append(row)
    details = workbook.create_sheet("Замечания")
    details.append(
        [
            "Файл",
            "Лист",
            "Ячейка",
            "Ключ",
            "Unit ID",
            "Язык",
            "Исходник",
            "Перевод",
            "Категория",
            "Важность",
            "Причина",
            "Судья 1",
            "Судья 2",
            "Действие",
            "Проверка не завершена",
            "Группа",
            "Проверка",
            "Исходный лист",
        ]
    )
    for issue in artifact["issues"]:
        row = [
            issue["file"],
            issue["sheet"],
            get_column_letter(issue["column"]) + str(issue["row"]),
            issue["key"],
            issue["unit_id"],
            issue["language"],
            issue["source"],
            issue["target"],
            label(issue["category"]),
            label(issue["severity"]),
            label(issue["reason"]),
            label(issue["judge_1"]),
            label(issue["judge_2"]),
            label(issue["action"]),
            "Да" if issue["unverified"] else "Нет",
            issue["group_id"],
            issue.get("check_id", ""),
            issue.get("original_sheet", issue["sheet"]),
        ]
        _check_values([row])
        details.append(row)
    for ws in workbook:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if cell.row == 1:
                    cell.font = Font(bold=True)
        for col in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(col)].width = 32 if col in {7, 8, 11} else 24
    summary.column_dimensions["A"].width = 46
    summary.column_dimensions["B"].width = 80
    return _save(workbook), XLSX, _name(artifact, "quality-report.xlsx")


def serialize_quality_bundle(artifact: dict, exported: dict) -> tuple[bytes, str, str]:
    files = [
        serialize_quality_translation(exported, artifact),
        serialize_quality_report(artifact, exported),
    ]
    manifest = {
        "schema_version": 1,
        "run_id": artifact["run_id"],
        "export_id": artifact["export_id"],
        "provenance": artifact["provenance"],
        "files": [
            {"name": name, "sha256": hashlib.sha256(data).hexdigest()} for data, _, name in files
        ],
    }
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        for data, _, name in files:
            archive.writestr(name, data)
        archive.writestr("manifest.json", json.dumps(manifest, ensure_ascii=False).encode())
    return output.getvalue(), "application/zip", _name(artifact, "results.zip")
