# Copyright © HCGameLoc
#
# SPDX-License-Identifier: GPL-3.0-or-later

"""
Weblate format contract tests for loc_kit_ingest output.

Generated PO/TBX files are loaded through Weblate's real format classes and,
for the glossary, through an actual glossary upload so the resulting database
units can be inspected. This verifies the generated files satisfy the Weblate
component/glossary contract end to end.

Run via: ./rundev.sh test weblate/trans/tests/test_loc_kit_ingest_contract.py
"""

from __future__ import annotations

import csv
import io
import json
import re
import shutil
import sys
import tempfile
import threading
import uuid
import zipfile
from dataclasses import replace
from datetime import timedelta
from pathlib import Path
from unittest.mock import patch

from django.conf import settings
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.files.base import ContentFile
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import DatabaseError, connection, transaction
from django.test import (
    RequestFactory,
    SimpleTestCase,
    TransactionTestCase,
)
from django.test.utils import modify_settings, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook
from translate.storage.pypo import pofile

from weblate.auth.data import SELECTION_ALL
from weblate.auth.models import Group, Permission, Role, User
from weblate.formats.models import FILE_FORMATS
from weblate.glossary.models import (
    build_glossary_prompt_entry,
    get_glossary_term_modes,
)
from weblate.glossary.tasks import flag_glossary_terminology, sync_terminology
from weblate.lang.models import Language
from weblate.trans import loc_kit
from weblate.trans.loc_kit import PREVIEW_WARNING_LIMIT
from weblate.trans.models import (
    Category,
    Component,
    PendingUnitChange,
    Project,
    Translation,
)
from weblate.trans.models.loc_kit import (
    LOC_KIT_DRAFT_STORAGE,
    LOC_KIT_STRING_UPDATE_APPLY_TIME_LIMIT,
    LocKitImportDraft,
)
from weblate.trans.tasks import (
    LOC_KIT_DISPATCH_MAX_ATTEMPTS,
    _flush_loc_kit_pending_changes,
    _publish_loc_kit_dispatch,
    apply_loc_kit_string_update_draft,
    drain_loc_kit_dispatches,
    perform_load,
    prepare_loc_kit_string_update,
)
from weblate.trans.tests.test_views import ViewTestCase
from weblate.trans.tests.utils import (
    RepoTestMixin,
    create_another_user,
    create_test_user,
)
from weblate.utils.lock import WeblateLockTimeoutError
from weblate.utils.state import (
    STATE_APPROVED,
    STATE_EMPTY,
    STATE_READONLY,
    STATE_TRANSLATED,
)
from weblate.utils.tests import http_mock
from weblate.utils.views import create_component_from_kit
from weblate.vcs.git import LocalRepository

# loc_kit_ingest is a standalone package at the repository root.
_REPO_ROOT = Path(__file__).resolve().parents[3]
if str(_REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(_REPO_ROOT))

# ruff: ignore[module-import-not-at-top-of-file]
from loc_kit_ingest.model import GlossaryTerm, ParseResult, StringUnit

# ruff: ignore[module-import-not-at-top-of-file]
from loc_kit_ingest.profile import (
    ComponentProfile,
    KeyColumn,
    KeyedGrammar,
    LanguageColumn,
    MetadataColumn,
    PairRegion,
    PairsGrammar,
)

# ruff: ignore[module-import-not-at-top-of-file]
from loc_kit_ingest.writer import render_component

# --------------------------------------------------------------------------- #
# Shared component profiles
# --------------------------------------------------------------------------- #


def make_po_component(source_lang: str, target_lang: str) -> ComponentProfile:
    return ComponentProfile(
        sheet="temple",
        component="Temple",
        kind="po",
        source_lang=source_lang,
        header_row=0,
        first_data_row=2,
        languages=(
            LanguageColumn(source_lang, source_lang, 2, source_lang),
            LanguageColumn(target_lang, target_lang, 3, target_lang),
        ),
        key=KeyColumn(0, "id"),
        comments=(MetadataColumn(1, "Character", "Character"),),
        references=(MetadataColumn(4, "Id", "Id"),),
        grammar=KeyedGrammar(skip_rows=(1,), allow_blank_rows=True),
        key_language=None,
        initial_target_languages=(),
    )


def make_tbx_component(source_lang: str, target_lang: str) -> ComponentProfile:
    return ComponentProfile(
        sheet="terms",
        component="Terms",
        kind="tbx",
        source_lang=source_lang,
        header_row=0,
        first_data_row=None,
        languages=(
            LanguageColumn(source_lang, source_lang, 0, source_lang),
            LanguageColumn(target_lang, target_lang, 1, target_lang),
        ),
        key=None,
        comments=(),
        references=(),
        grammar=PairsGrammar(skip_rows=(1,), regions=(PairRegion(2, 3, 4),)),
        key_language=target_lang,
        initial_target_languages=(target_lang,),
    )


# --------------------------------------------------------------------------- #
# Format-level contract (no database needed)
# --------------------------------------------------------------------------- #


class LocKitFormatContractTest(SimpleTestCase):
    """Generated files load through Weblate's real format classes."""

    def setUp(self) -> None:
        super().setUp()
        self.tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.tmp, ignore_errors=True)

    def render_po(self) -> Path:
        component = make_po_component("ru", "en")
        result = ParseResult(
            component="Temple",
            kind="po",
            units=(
                StringUnit(
                    key="dialog.intro",
                    values={
                        "ru": " Привет, [shake]путник[/shake]!\r\n",
                        "en": "<color=#E3BA59>Hello, traveler!</color>",
                    },
                    comments=("Character: Sage",),
                    references=("101",),
                    row=4,
                ),
                StringUnit(
                    key="dialog.bye",
                    values={"ru": "Прощай.", "en": "Farewell."},
                    comments=("Character: Sage",),
                    references=("102",),
                    row=5,
                ),
            ),
            diagnostics=(),
            skipped_rows=(),
        )
        render_component(component, result, self.tmp)
        return self.tmp / "Temple"

    def render_tbx(self) -> Path:
        component = make_tbx_component("ru", "en")
        result = ParseResult(
            component="Terms",
            kind="tbx",
            units=(
                GlossaryTerm(
                    context="characters.sage",
                    values={"ru": "Мудрец", "en": "Sage"},
                    source_explanation="Источник описание",
                    target_explanations={"en": "Wise person"},
                    section="Characters",
                    term_row=4,
                    source_flags=("read-only",),
                    note_rows=(5,),
                ),
            ),
            diagnostics=(),
            skipped_rows=(),
        )
        render_component(component, result, self.tmp)
        return self.tmp / "Terms" / "tbx"

    @staticmethod
    def parse_mono(path: Path, template: Path):
        """Monolingual formats need an explicit template store."""
        cls = FILE_FORMATS["po-mono"]
        return cls(path, template_store=cls(template, is_template=True))

    def test_po_mono_template_maps_key_to_context_and_source(self) -> None:
        po_dir = self.render_po()
        store = self.parse_mono(po_dir / "ru.po", po_dir / "ru.po")

        contexts = {unit.context for unit in store.all_units if unit.context}
        self.assertEqual(contexts, {"dialog.intro", "dialog.bye"})

        unit = next(u for u in store.all_units if u.context == "dialog.intro")
        self.assertEqual(unit.source, " Привет, [shake]путник[/shake]!\r\n")

    def test_po_mono_translation_keeps_markup_and_source(self) -> None:
        po_dir = self.render_po()
        store = self.parse_mono(po_dir / "en.po", po_dir / "ru.po")

        unit = next(u for u in store.all_units if u.context == "dialog.intro")
        self.assertEqual(unit.target, "<color=#E3BA59>Hello, traveler!</color>")

    def test_po_mono_keeps_developer_comment_and_location(self) -> None:
        po_dir = self.render_po()
        store = self.parse_mono(po_dir / "ru.po", po_dir / "ru.po")

        unit = next(u for u in store.all_units if u.context == "dialog.intro")
        self.assertIn("Character: Sage", unit.notes)
        self.assertIn("101", unit.locations)

    def test_tbx_is_bilingual_with_both_explanations(self) -> None:
        tbx_dir = self.render_tbx()
        # Weblate always supplies the configured languages when parsing.
        store = FILE_FORMATS["tbx"](
            tbx_dir / "en.tbx", source_language="ru", language_code="en"
        )

        units = [u for u in store.all_units if u.context]
        self.assertEqual(len(units), 1)
        unit = units[0]
        self.assertEqual(unit.context, "characters.sage")
        self.assertEqual(unit.source, "Мудрец")
        self.assertEqual(unit.target, "Sage")
        self.assertEqual(unit.source_explanation, "Источник описание")
        self.assertEqual(unit.explanation, "Wise person")

    def test_tbx_carries_source_glossary_flags(self) -> None:
        tbx_dir = self.render_tbx()
        store = FILE_FORMATS["tbx"](
            tbx_dir / "en.tbx", source_language="ru", language_code="en"
        )

        (unit,) = [item for item in store.all_units if item.context]

        self.assertIn("read-only", unit.flags)

    def test_tbx_never_written_for_source_language(self) -> None:
        tbx_dir = self.render_tbx()
        self.assertFalse((tbx_dir / "ru.tbx").exists())
        self.assertTrue((tbx_dir / "en.tbx").is_file())

    def test_tbx_uses_profile_xml_lang_not_weblate_code(self) -> None:
        # The Weblate code zh_Hans must be emitted as the BCP-47 tag zh-Hans.
        component = replace(
            make_tbx_component("ru", "zh_Hans"),
            languages=(
                LanguageColumn("ru", "ru", 0, "ru"),
                LanguageColumn("zh_Hans", "zh-Hans", 1, "zh_Hans"),
            ),
        )
        result = ParseResult(
            component="Terms",
            kind="tbx",
            units=(
                GlossaryTerm(
                    context="characters.sage",
                    values={"ru": "Мудрец", "zh_Hans": "贤者"},
                    source_explanation="Источник",
                    target_explanations={"zh_Hans": "解释"},
                    section="Characters",
                    term_row=4,
                    note_rows=(5,),
                ),
            ),
            diagnostics=(),
            skipped_rows=(),
        )
        render_component(component, result, self.tmp)
        xml = (self.tmp / "Terms" / "tbx" / "zh_Hans.tbx").read_text(encoding="utf-8")
        self.assertIn('xml:lang="zh-Hans"', xml)
        self.assertNotIn('xml:lang="zh_Hans"', xml)


# --------------------------------------------------------------------------- #
# Glossary import contract (real upload, real database units)
# --------------------------------------------------------------------------- #


class LocKitGlossaryImportContractTest(ViewTestCase):
    """A generated TBX imports into a real glossary and drives the LLM payload."""

    CREATE_GLOSSARIES: bool = True

    def setUp(self) -> None:
        super().setUp()
        self.glossary_component = self.project.glossaries[0]
        self.glossary = self.glossary_component.translation_set.get(
            language=self.get_translation().language
        )
        self.tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.tmp, ignore_errors=True)

    def generate_tbx(self) -> Path:
        """Render a TBX matching the test project's source/target languages."""
        source_code = self.component.source_language.code
        target_code = self.glossary.language.code
        component = make_tbx_component(source_code, target_code)
        result = ParseResult(
            component="Terms",
            kind="tbx",
            units=(
                GlossaryTerm(
                    context="characters.sage",
                    values={source_code: "Sage", target_code: "Mudrc"},
                    source_explanation="A wise character",
                    target_explanations={target_code: "Moudra postava"},
                    section="Characters",
                    term_row=4,
                    note_rows=(5,),
                ),
            ),
            diagnostics=(),
            skipped_rows=(),
        )
        render_component(component, result, self.tmp)
        return self.tmp / "Terms" / "tbx" / f"{target_code}.tbx"

    def upload_tbx(self, path: Path):
        with path.open("rb") as handle:
            return self.client.post(
                reverse("upload", kwargs={"path": self.glossary.get_url_path()}),
                {"file": handle, "method": "add"},
            )

    def test_generated_tbx_imports_term_with_context_and_target(self) -> None:
        response = self.upload_tbx(self.generate_tbx())
        self.assertRedirects(response, self.glossary.get_absolute_url())

        unit = self.glossary.unit_set.get(context="characters.sage")
        self.assertEqual(unit.source, "Sage")
        self.assertEqual(unit.target, "Mudrc")

    def test_translation_upload_does_not_carry_explanations(self) -> None:
        """
        Documents a Weblate upload-path limitation, not a writer defect.

        The translation-upload view imports source/target only. Explanations
        reach the database through repository synchronisation instead, which is
        why the runbook attaches Terms.zip as component files. The generated
        file does carry both explanations - see
        LocKitFormatContractTest.test_tbx_is_bilingual_with_both_explanations.
        """
        self.upload_tbx(self.generate_tbx())

        unit = self.glossary.unit_set.get(context="characters.sage")
        self.assertEqual(unit.explanation, "")

    def test_imported_unit_produces_full_llm_glossary_entry(self) -> None:
        """A real database unit with explanations yields all payload fields."""
        self.upload_tbx(self.generate_tbx())
        unit = self.glossary.unit_set.get(context="characters.sage")

        # Apply the explanations that repository synchronisation would set.
        unit.explanation = "Moudra postava"
        unit.save(update_fields=["explanation"])
        unit.source_unit.explanation = "A wise character"
        unit.source_unit.save(update_fields=["explanation"])
        unit.refresh_from_db()

        entry = build_glossary_prompt_entry(unit)

        self.assertIsNotNone(entry)
        self.assertEqual(entry["source"], "Sage")
        self.assertEqual(entry["target"], "Mudrc")
        self.assertEqual(entry["source_explanation"], "A wise character")
        self.assertEqual(entry["target_explanation"], "Moudra postava")


# --------------------------------------------------------------------------- #
# Universal component upload: table kits through the real create view
# --------------------------------------------------------------------------- #


class LocKitUniversalUploadContractTest(ViewTestCase):
    """
    A table kit uploaded in the create UI becomes a live component.

    Covers the N-language contract: every populated language column arrives
    as a translation, regardless of how many there are.
    """

    KIT_CSV = (
        "id,,ru,en,ja,ko,de\n"
        "id-ignore,Actor,Russian ,English ,,,\n"
        "line_1,Ann,Привет,Hello,こんにちは,안녕,Hallo\n"
        "line_2,Bob,Пока,Bye,さようなら,잘 가,Tschüss\n"
    )

    def _upload(self, name: str, body: str):
        return SimpleUploadedFile(name, body.encode(), content_type="text/csv")

    def _kit_data(self, slug: str) -> dict:
        return {"project": self.project, "slug": slug, "name": slug.title()}

    def test_five_language_kit_renders_every_language(self) -> None:
        fake, info = create_component_from_kit(
            self._kit_data("dialogs"),
            self._upload("Space Kit - Dialogs.csv", self.KIT_CSV),
        )
        self.assertEqual(info["languages"], ["ru", "en", "ja", "ko", "de"])
        self.assertEqual(info["source_lang"], "ru")
        self.assertEqual(info["units"], 2)
        self.assertEqual(info["template"], "ru.po")
        repo = Path(fake.full_path)
        self.assertEqual(
            sorted(path.name for path in repo.glob("*.po")),
            ["de.po", "en.po", "ja.po", "ko.po", "ru.po"],
        )
        shutil.rmtree(fake.full_path, ignore_errors=True)

    def test_keyless_kit_uses_source_term_as_key(self) -> None:
        # No dedicated key column: the ru column is both the PO key and the
        # source language. The "Use as glossary" checkbox does not exist at
        # this call boundary, so it cannot influence this ordinary PO path.
        keyless_csv = "ru,en,ja\nПривет,Hello,こんにちは\nПока,Bye,さようなら\n"
        fake, info = create_component_from_kit(
            self._kit_data("keyless"), self._upload("Dialogs.csv", keyless_csv)
        )
        self.assertEqual(info["source_lang"], "ru")
        self.assertEqual(info["languages"], ["ru", "en", "ja"])
        self.assertEqual(info["template"], "ru.po")
        self.assertEqual(info["units"], 2)
        self.assertTrue(
            any(
                "is both the PO key and a language column" in note
                for note in info["notes"]
            )
        )

        repo = Path(fake.full_path)
        self.assertEqual(
            sorted(path.name for path in repo.glob("*.po")),
            ["en.po", "ja.po", "ru.po"],
        )

        ru_units = {u.getid(): u for u in pofile.parsefile(str(repo / "ru.po")).units}
        en_units = {u.getid(): u for u in pofile.parsefile(str(repo / "en.po")).units}
        self.assertEqual(ru_units["Привет"].target, "Привет")
        self.assertEqual(en_units["Привет"].target, "Hello")
        self.assertEqual(en_units["Пока"].target, "Bye")

        shutil.rmtree(fake.full_path, ignore_errors=True)

    def test_zip_upload_keeps_historical_behavior(self) -> None:
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("cs.po", 'msgid "hello"\nmsgstr "ahoj"\n')
        fake, info = create_component_from_kit(
            self._kit_data("zipped"),
            SimpleUploadedFile(
                "kit.zip", buffer.getvalue(), content_type="application/zip"
            ),
        )
        self.assertIsNone(info)
        self.assertTrue((Path(fake.full_path) / "cs.po").exists())
        shutil.rmtree(fake.full_path, ignore_errors=True)

    def test_duplicate_keys_block_intake(self) -> None:
        broken = "id,ru,en\na,Один,One\na,Два,Two\n"
        with self.assertRaises(ValidationError) as ctx:
            create_component_from_kit(
                self._kit_data("broken"), self._upload("Broken.csv", broken)
            )
        self.assertIn("duplicate key", "".join(ctx.exception.messages))

    def test_unsupported_suffix_is_rejected(self) -> None:
        with self.assertRaises(ValidationError) as ctx:
            create_component_from_kit(
                self._kit_data("plain"), self._upload("notes.txt", "hello")
            )
        self.assertIn(".csv", "".join(ctx.exception.messages))

    def test_create_view_builds_component_from_csv(self) -> None:
        self.user.is_superuser = True
        self.user.save()

        with modify_settings(INSTALLED_APPS={"remove": "weblate.billing"}):
            response = self.client.post(
                reverse("create-component-zip"),
                {
                    "zipfile": self._upload("Space Kit - Dialogs.csv", self.KIT_CSV),
                    "name": "Dialogs",
                    "slug": "dialogs",
                    "project": self.project.pk,
                    "source_language": self.component.source_language.pk,
                },
            )
            # The kit skips discovery: the create form arrives prefilled.
            self.assertContains(response, "Loc-kit converted")
            self.assertContains(response, "ru.po")

            form = response.context["form"]
            params = {field: form[field].value() or "" for field in form.fields}
            params.pop("inherit_new_lang", None)
            params["new_lang"] = "none"
            response = self.client.post(
                reverse("create-component-zip"), params, follow=True
            )

        component = Component.objects.get(slug="dialogs")
        self.assertEqual(component.file_format, "po-mono")
        self.assertEqual(component.template, "ru.po")
        self.assertEqual(component.source_language.code, "ru")
        codes = sorted(
            component.translation_set.values_list("language__code", flat=True)
        )
        self.assertEqual(codes, ["de", "en", "ja", "ko", "ru"])
        self.assertEqual(component.source_translation.unit_set.count(), 2)

        # Unit identity: the game key is the context, texts are source/target.
        ja_unit = component.translation_set.get(language__code="ja").unit_set.get(
            context="line_1"
        )
        self.assertEqual(ja_unit.source, "Привет")
        self.assertEqual(ja_unit.target, "こんにちは")

        # The Character column survives as a developer comment on the unit.
        source_unit = component.source_translation.unit_set.get(context="line_1")
        self.assertIn("Ann", source_unit.note)

    def test_game_language_codes_become_weblate_translations(self) -> None:
        # Game kits label Chinese, Japanese and Korean with their own codes;
        # unrecognised they would silently become developer comments.
        kit = (
            "id,ru,en,ch-s,jp,kr\n"
            "line_1,Привет,Hello,你好,こんにちは,안녕\n"
            "line_2,Пока,Bye,再见,さようなら,잘 가\n"
        )
        fake, info = create_component_from_kit(
            self._kit_data("codes"), self._upload("Space Kit - Codes.csv", kit)
        )
        self.assertEqual(info["languages"], ["ru", "en", "zh_Hans", "ja", "ko"])
        repo = Path(fake.full_path)
        self.assertEqual(
            sorted(path.name for path in repo.glob("*.po")),
            ["en.po", "ja.po", "ko.po", "ru.po", "zh_Hans.po"],
        )
        self.assertEqual(
            Language.objects.filter(code__in=["zh_Hans", "ja", "ko"]).count(), 3
        )
        shutil.rmtree(fake.full_path, ignore_errors=True)

    def test_key_without_any_text_blocks_intake(self) -> None:
        broken = "id,ru,en\nline_1,Привет,Hello\nline_2,,\n"
        with self.assertRaises(ValidationError) as ctx:
            create_component_from_kit(
                self._kit_data("empty-row"), self._upload("Empty.csv", broken)
            )
        self.assertIn("no text in any language", "".join(ctx.exception.messages))

    def test_sourceless_key_becomes_a_unit_with_an_empty_source(self) -> None:
        self.user.is_superuser = True
        self.user.save()
        kit = (
            "id,ru,en,ja\n"
            "line_1,Привет,Hello,こんにちは\n"
            "line_2,,Beta only,ベータ限定\n"
        )

        with modify_settings(INSTALLED_APPS={"remove": "weblate.billing"}):
            response = self.client.post(
                reverse("create-component-zip"),
                {
                    "zipfile": self._upload("Space Kit - Beta.csv", kit),
                    "name": "Beta",
                    "slug": "beta",
                    "project": self.project.pk,
                    "source_language": self.component.source_language.pk,
                },
            )
            self.assertContains(
                response, "1 strings were imported without a source string"
            )
            self.assertContains(response, "po.missing_source")

            form = response.context["form"]
            params = {field: form[field].value() or "" for field in form.fields}
            params.pop("inherit_new_lang", None)
            params["new_lang"] = "none"
            self.client.post(reverse("create-component-zip"), params, follow=True)

        component = Component.objects.get(slug="beta")
        self.assertEqual(component.source_translation.unit_set.count(), 2)

        source_unit = component.source_translation.unit_set.get(context="line_2")
        self.assertEqual(source_unit.source, "")
        self.assertEqual(source_unit.target, "")

        ja_unit = component.translation_set.get(language__code="ja").unit_set.get(
            context="line_2"
        )
        self.assertEqual(ja_unit.source, "")
        self.assertEqual(ja_unit.target, "ベータ限定")

    def test_source_markup_defect_warns_but_keeps_create_available(self) -> None:
        """
        A bad closing tag in the kit source is a warning, not a blocker.

        Defect 525591: the ru source carries `</color=yellow>`. The wizard
        must surface the diagnostic (code and key) and still let the create
        form through, because the import itself is not refused.
        """
        self.user.is_superuser = True
        self.user.save()
        kit = (
            "id,ru,en\n"
            "spawn_rate,Заказчики приходят в </color=yellow>{0}</color> раза реже.,Fewer customers\n"
            "line_1,Привет,Hello\n"
        )

        with modify_settings(INSTALLED_APPS={"remove": "weblate.billing"}):
            response = self.client.post(
                reverse("create-component-zip"),
                {
                    "zipfile": self._upload("Space Kit - Markup.csv", kit),
                    "name": "Markup",
                    "slug": "markup",
                    "project": self.project.pk,
                    "source_language": self.component.source_language.pk,
                },
            )
            self.assertContains(response, "source.tag_closing_has_attribute")
            self.assertContains(response, "spawn_rate")

            # The create form arrives prefilled and remains submittable.
            form = response.context["form"]
            params = {field: form[field].value() or "" for field in form.fields}
            params.pop("inherit_new_lang", None)
            params["new_lang"] = "none"
            self.client.post(reverse("create-component-zip"), params, follow=True)

        component = Component.objects.get(slug="markup")
        source_unit = component.source_translation.unit_set.get(context="spawn_rate")
        self.assertIn("</color=yellow>", source_unit.source)

    def test_create_view_applies_kit_explanations_to_source_units(self) -> None:
        self.user.is_superuser = True
        self.user.save()
        kit = (
            "id,ru,en,Comment,Explanation\n"
            "line_1,Привет,Hello,Shown on load,Casual greeting\n"
            "line_2,Пока,Bye,,Farewell line\n"
        )

        with modify_settings(INSTALLED_APPS={"remove": "weblate.billing"}):
            response = self.client.post(
                reverse("create-component-zip"),
                {
                    "zipfile": self._upload("Space Kit - Explained.csv", kit),
                    "name": "Explained",
                    "slug": "explained",
                    "project": self.project.pk,
                    "source_language": self.component.source_language.pk,
                },
            )
            self.assertContains(response, "will set 2 strings")

            form = response.context["form"]
            params = {field: form[field].value() or "" for field in form.fields}
            params.pop("inherit_new_lang", None)
            params["new_lang"] = "none"
            self.client.post(reverse("create-component-zip"), params, follow=True)

        component = Component.objects.get(slug="explained")
        line1 = component.source_translation.unit_set.get(context="line_1")
        line2 = component.source_translation.unit_set.get(context="line_2")
        self.assertEqual(line1.explanation, "Casual greeting")
        self.assertEqual(line1.note, "Shown on load")
        self.assertEqual(line2.explanation, "Farewell line")
        po_bytes = Path(component.full_path, "ru.po").read_bytes()
        self.assertNotIn(b"Casual greeting", po_bytes)
        self.assertNotIn(b"Farewell line", po_bytes)

    def test_deferred_perform_load_applies_staged_explanations(self) -> None:
        """
        A load deferred to perform_load carries its staged explanations.

        When create_translations hits a lock timeout and defers the load to
        the perform_load task, the wizard's explanation map used to die with
        the request; the task must carry it and apply it once translations
        exist. Simulated by calling the task body synchronously the same way
        the worker does.
        """
        self.user.is_superuser = True
        self.user.save()
        kit = "id,ru,en,Explanation\nline_1,Привет,Hello,Casual greeting\n"

        with modify_settings(INSTALLED_APPS={"remove": "weblate.billing"}):
            response = self.client.post(
                reverse("create-component-zip"),
                {
                    "zipfile": self._upload("Deferred Kit.csv", kit),
                    "name": "Deferred",
                    "slug": "deferred",
                    "project": self.project.pk,
                    "source_language": self.component.source_language.pk,
                },
            )
            form = response.context["form"]
            params = {field: form[field].value() or "" for field in form.fields}
            params.pop("inherit_new_lang", None)
            params["new_lang"] = "none"
            self.client.post(reverse("create-component-zip"), params, follow=True)

        component = Component.objects.get(slug="deferred")
        # Wipe explanations the eager path may already have applied, then
        # rerun the load exactly as the deferred task would: the map must
        # still reach the source units.
        component.source_translation.unit_set.update(explanation="")
        perform_load(
            component.pk,
            force=True,
            loc_kit_explanations={"line_1": "Casual greeting"},
            user_id=self.user.pk,
        )

        line1 = component.source_translation.unit_set.get(context="line_1")
        self.assertEqual(line1.explanation, "Casual greeting")

    def test_create_view_without_explanation_column_needs_no_session_state(
        self,
    ) -> None:
        """A kit with no Explanation column never touches the pending session key."""
        self.user.is_superuser = True
        self.user.save()

        with modify_settings(INSTALLED_APPS={"remove": "weblate.billing"}):
            self.client.post(
                reverse("create-component-zip"),
                {
                    "zipfile": self._upload("Space Kit - Dialogs.csv", self.KIT_CSV),
                    "name": "Dialogs",
                    "slug": "dialogs",
                    "project": self.project.pk,
                    "source_language": self.component.source_language.pk,
                },
            )

        self.assertNotIn("loc_kit_pending_explanations", self.client.session)


# --------------------------------------------------------------------------- #
# Glossary intake UI: sheet selection, preview, correction, confirmation
# --------------------------------------------------------------------------- #

GLOSSARY_CSV = (
    "domain,ru,en,note_ru,note_en\n"
    "Персонажи,Герой,Hero,Главный протагонист,Main protagonist\n"
    "Оружие,Меч,Sword,Ближний бой,Melee weapon\n"
)


# Language-only kit with no extra columns: deterministic inference must give
# a preview. GLOSSARY_CSV above is intentionally NOT parsed deterministically:
# `domain` and `note_ru`/`note_en` are in neither _NOTE_HEADERS nor
# _IGNORABLE_HEADERS (loc_kit_ingest/infer.py), so it keeps covering the
# LLM/manual path. Widening either set breaks that coverage - re-point these
# tests at a new fixture instead of relaxing them.
GLOSSARY_LANG_ONLY_CSV = "ru,en\nRussian,English\nГерой,Hero\nМеч,Sword\n"

GLOSSARY_NOTE_CSV = (
    "ru,en,fr,note\n"
    "Russian,English,French,Note\n"
    "Партия,Party,Parti,"
    '"Правящая политическая партия. Во французском le Parti, мужской род."\n'
    "Самосбор,Samosbor,Samosbor,Термин вселенной.\n"
)
GLOSSARY_FLAGS_CSV = (
    "ru,en,cs,flags\n"
    "HeroCraft,HeroCraft,HeroCraft,read-only\n"
    "Судно,Vessel,Plavidlo,forbidden\n"
    "Точное,Exact,,exact\n"
)


# Real terminology exports put a term on one row and its description on the
# next, under a section caption. Inference must map the descriptions as
# explanations, not as twice as many terms.
_DESC_RU = "Главный протагонист истории и первый играбельный персонаж. " * 3
_DESC_EN = "The main protagonist of the story and the first playable hero. " * 3
_DESC2_RU = "Клинок ближнего боя и стартовое оружие в первой главе игры. " * 3
_DESC2_EN = "A melee blade handed to the player in the first chapter. " * 3
GLOSSARY_PAIRS_CSV = (
    "ru,en\n"
    "Russian,English\n"
    "Персонажи,Characters\n"
    f"Герой,Hero\n{_DESC_RU},{_DESC_EN}\n"
    f"Меч,Sword\n{_DESC2_RU},{_DESC2_EN}\n"
)

GLOSSARY_SEMICOLON_CSV = (
    'ru;en;notes\nЛеон;Leon;"Имя собственное, мужской род."\nАки;Aki;Сестра Леона.\n'
)

GLOSSARY_ID_PARTIAL_CSV = (
    "id,ru,en,ja,zh-TC,notes\n"
    "char_leon,Леон,Leon,レオン,,главный герой\n"
    "char_aki,Аки,Aki,,阿姬,\n"
    "char_joe,Джо,Joe,,,паук\n"
)

GLOSSARY_SEMICOLON_PARTIAL_CSV = (
    "ru;en;ja;zh-TC;notes\nЛеон;Leon;レオン;;главный герой\nАки;Aki;;阿姬;\n"
)


def _glossary_profile(sheet: str, *, source_lang: str = "ru") -> dict:
    return {
        "schema_version": 2,
        "components": [
            {
                "sheet": sheet,
                "component": "ignored-by-server",
                "kind": "tbx",
                "source_lang": source_lang,
                "header_row": 1,
                "languages": [
                    {"code": "ru", "xml_lang": "ru", "column": 2, "header": "ru"},
                    {"code": "en", "xml_lang": "en", "column": 3, "header": "en"},
                ],
                "grammar": {
                    "type": "record-map",
                    "skip_rows": [],
                    "regions": [
                        {
                            "first_record_row": 2,
                            "last_record_row": 3,
                            "record_stride": 1,
                        }
                    ],
                    "term_row_offset": 0,
                    "section_field": {
                        "column": 1,
                        "header": "domain",
                        "row_offset": 0,
                    },
                    "notes": [
                        {
                            "scope": "source",
                            "column": 4,
                            "header": "note_ru",
                            "row_offset": 0,
                        },
                        {
                            "scope": "target",
                            "language": "en" if source_lang == "ru" else "ru",
                            "column": 5,
                            "header": "note_en",
                            "row_offset": 0,
                        },
                    ],
                },
                "initial_target_languages": ["en" if source_lang == "ru" else "ru"],
            }
        ],
    }


class LocKitGlossaryUploadUITest(ViewTestCase):
    """The glossary intake stages, end to end, through the real views."""

    def setUp(self) -> None:
        super().setUp()
        self.user.is_superuser = True
        self.user.save()
        # Component repositories live under a DATA_DIR path that xdist
        # workers share, so two tests creating the same slug race on disk
        # even though their databases are separate.
        self.slug = f"gloss-{uuid.uuid4().hex[:8]}"

    def _csv(self, name: str = "Glossary.csv", body: str = GLOSSARY_CSV):
        return SimpleUploadedFile(name, body.encode(), content_type="text/csv")

    def _start(self, upload=None, slug: str | None = None):
        """Upload a table with the glossary checkbox set."""
        slug = slug or self.slug
        with modify_settings(INSTALLED_APPS={"remove": "weblate.billing"}):
            return self.client.post(
                reverse("create-component-zip"),
                {
                    "zipfile": upload or self._csv(),
                    "name": slug.title(),
                    "slug": slug,
                    "project": self.project.pk,
                    "source_language": self.component.source_language.pk,
                    "is_glossary": "1",
                },
            )

    def _draft(self):
        return LocKitImportDraft.objects.get()

    def _profile_upload(self, document: dict, name: str = "fix.loc-ingest.json"):
        return SimpleUploadedFile(
            name, json.dumps(document).encode(), content_type="application/json"
        )

    def _upload_profile(self, draft, document: dict):
        return self.client.post(
            reverse("loc-kit-glossary-preview", kwargs={"token": draft.token}),
            {"action": "upload-profile", "profile": self._profile_upload(document)},
            follow=True,
        )

    def _select_sheet(self, draft, sheet: str = "Glossary"):
        return self.client.post(
            reverse("loc-kit-sheet-select", kwargs={"token": draft.token}),
            {"sheet": sheet},
        )

    def _confirm(self):
        """Drive preview confirmation and the final component form."""
        draft = self._draft()
        with modify_settings(INSTALLED_APPS={"remove": "weblate.billing"}):
            response = self.client.post(
                reverse("loc-kit-glossary-preview", kwargs={"token": draft.token}),
                {"action": "confirm"},
                follow=True,
            )
            form = response.context["form"]
            params = {field: form[field].value() or "" for field in form.fields}
            params.pop("inherit_new_lang", None)
            params["new_lang"] = "none"
            return self.client.post(
                reverse("loc-kit-glossary-confirm", kwargs={"token": draft.token}),
                params,
                follow=True,
            )

    # ----------------------------------------------------------------- #
    # Happy path
    # ----------------------------------------------------------------- #

    def test_glossary_upload_creates_a_draft_and_asks_for_a_sheet(self) -> None:
        response = self._start()

        draft = self._draft()
        self.assertRedirects(
            response,
            reverse("loc-kit-sheet-select", kwargs={"token": draft.token}),
        )
        self.assertEqual(draft.owner, self.user)
        # Single sheet: auto-skip fixed it, but GLOSSARY_CSV's domain/note_*
        # columns refuse deterministic inference, so it stays at sheet-select.
        self.assertEqual(draft.state, LocKitImportDraft.State.SHEET_SELECTED)
        # No component exists yet.
        self.assertFalse(Component.objects.filter(slug=self.slug).exists())

    @override_settings(LOC_KIT_PROFILE_ANALYSIS_ENABLED=False)
    def test_language_only_sheet_gets_deterministic_preview(self) -> None:
        """Языковая таблица даёт превью локально, без OpenRouter и без JSON."""
        self._start(
            upload=self._csv("Terms.csv", GLOSSARY_LANG_ONLY_CSV), slug=self.slug
        )
        draft = self._draft()

        draft.refresh_from_db()
        self.assertEqual(draft.state, LocKitImportDraft.State.PREVIEW_READY)
        preview = json.loads(draft.preview_json)
        self.assertEqual(preview["source_language"], "ru")
        self.assertEqual(preview["target_languages"], ["en"])
        self.assertEqual(preview["term_count"], 2)

        page = self.client.get(
            reverse("loc-kit-glossary-preview", kwargs={"token": draft.token})
        )
        self.assertContains(page, "Герой")

    @override_settings(LOC_KIT_PROFILE_ANALYSIS_ENABLED=False)
    def test_deterministic_preview_confirms_into_live_component(self) -> None:
        self._start(
            upload=self._csv("Terms.csv", GLOSSARY_LANG_ONLY_CSV), slug=self.slug
        )
        self._confirm()
        component = Component.objects.get(slug=self.slug)
        self.assertTrue(component.is_glossary)
        self.assertEqual(component.source_language.code, "ru")
        codes = set(component.translation_set.values_list("language__code", flat=True))
        self.assertIn("en", codes)

    @override_settings(LOC_KIT_PROFILE_ANALYSIS_ENABLED=False)
    def test_confirmed_glossary_flags_terminology(self) -> None:
        self._start(
            upload=self._csv("Terms.csv", GLOSSARY_LANG_ONLY_CSV), slug=self.slug
        )
        with (
            patch.object(Component, "update_branch", return_value=True),
            self.captureOnCommitCallbacks(execute=True),
        ):
            self._confirm()
        component = Component.objects.get(slug=self.slug, is_glossary=True)
        sources = component.source_translation.unit_set.all()
        self.assertTrue(sources)
        for unit in sources:
            self.assertIn("terminology", unit.all_flags)

    @override_settings(LOC_KIT_PROFILE_ANALYSIS_ENABLED=False)
    def test_source_flags_are_previewed_and_created(self) -> None:
        self._start(
            upload=self._csv("Terms.csv", GLOSSARY_FLAGS_CSV),
            slug=self.slug,
        )
        draft = self._draft()
        draft.refresh_from_db()

        preview = json.loads(draft.preview_json)
        self.assertEqual(preview["terms"][0]["source_flags"], ["read-only"])
        page = self.client.get(
            reverse("loc-kit-glossary-preview", kwargs={"token": draft.token})
        )
        self.assertContains(page, "read-only")
        self.assertContains(page, "forbidden")

        with (
            patch.object(Component, "update_branch", return_value=True),
            self.captureOnCommitCallbacks(execute=True),
        ):
            self._confirm()

        component = Component.objects.get(slug=self.slug, is_glossary=True)
        source_units = {
            unit.source: unit for unit in component.source_translation.unit_set.all()
        }
        target_units = {
            unit.source: unit
            for unit in component.translation_set.get(
                language__code="en"
            ).unit_set.all()
        }
        self.assertIn("read-only", target_units["HeroCraft"].flags)
        self.assertIn("forbidden", target_units["Судно"].flags)
        read_only_entry = build_glossary_prompt_entry(target_units["HeroCraft"])
        forbidden_entry = build_glossary_prompt_entry(target_units["Судно"])
        self.assertIn("read-only", source_units["HeroCraft"].extra_flags)
        self.assertIn("forbidden", source_units["Судно"].extra_flags)
        self.assertIn("read-only", read_only_entry["flags"])
        self.assertIn("forbidden", forbidden_entry["flags"])
        self.assertIn("exact", target_units["Точное"].flags)
        self.assertNotIn("exact", source_units["Точное"].extra_flags)
        exact_modes = get_glossary_term_modes(target_units["Точное"])
        self.assertEqual(exact_modes, {"exact"})
        cs_exact = component.translation_set.get(language__code="cs").unit_set.get(
            source="Точное"
        )
        self.assertNotIn("exact", cs_exact.extra_flags)
        source_units["HeroCraft"].update_extra_flags("terminology", self.user)
        flag_glossary_terminology(component.pk)
        source_units["HeroCraft"].refresh_from_db()
        self.assertIn("read-only", source_units["HeroCraft"].extra_flags)

    @override_settings(LOC_KIT_PROFILE_ANALYSIS_ENABLED=False)
    def test_term_description_sheet_maps_descriptions_as_explanations(self) -> None:
        """Термин и описание на соседних строках - одна запись, не две."""
        self._start(upload=self._csv("Terms.csv", GLOSSARY_PAIRS_CSV), slug=self.slug)
        draft = self._draft()

        draft.refresh_from_db()

        self.assertEqual(draft.state, LocKitImportDraft.State.PREVIEW_READY)
        preview = json.loads(draft.preview_json)
        self.assertEqual(preview["term_count"], 2)
        self.assertEqual(preview["note_count"], 4)
        self.assertEqual(preview["terms"][0]["section"], "Персонажи")
        self.assertEqual(preview["terms"][0]["source"], "Герой")

    @override_settings(LOC_KIT_PROFILE_ANALYSIS_ENABLED=False)
    def test_confirmed_glossary_populates_a_new_language(self) -> None:
        self._start(
            upload=self._csv("Terms.csv", GLOSSARY_LANG_ONLY_CSV), slug=self.slug
        )
        with (
            patch.object(Component, "update_branch", return_value=True),
            self.captureOnCommitCallbacks(execute=True),
        ):
            self._confirm()

        component = Component.objects.get(slug=self.slug, is_glossary=True)
        expected = component.source_translation.unit_set.count()
        component.add_new_language(Language.objects.get(code="fr"), None)

        added = component.translation_set.get(language__code="fr")
        self.assertEqual(added.unit_set.count(), expected)

    @override_settings(LOC_KIT_PROFILE_ANALYSIS_ENABLED=False)
    def test_note_column_reaches_the_created_glossary_llm_entry(self) -> None:
        self._start(upload=self._csv("Terms.csv", GLOSSARY_NOTE_CSV), slug=self.slug)
        draft = self._draft()
        draft.refresh_from_db()

        self.assertEqual(draft.state, LocKitImportDraft.State.PREVIEW_READY)
        preview = json.loads(draft.preview_json)
        self.assertEqual(preview["term_count"], 2)
        self.assertEqual(preview["note_count"], 2)
        self.assertIn("мужской род", preview["terms"][0]["source_explanation"])

        page = self.client.get(
            reverse("loc-kit-glossary-preview", kwargs={"token": draft.token})
        )
        self.assertContains(page, "мужской род")

        with (
            patch.object(Component, "update_branch", return_value=True),
            self.captureOnCommitCallbacks(execute=True),
        ):
            self._confirm()
        component = Component.objects.get(slug=self.slug)
        translation = component.translation_set.get(language__code="fr")
        unit = translation.unit_set.get(source="Партия")
        entry = build_glossary_prompt_entry(unit)
        self.assertIsNotNone(entry)
        self.assertEqual(
            entry["source_explanation"],
            "Правящая политическая партия. Во французском le Parti, мужской род.",
        )

    @override_settings(LOC_KIT_PROFILE_ANALYSIS_ENABLED=False)
    def test_operator_switches_the_layout_from_the_preview(self) -> None:
        """Неверно угаданная раскладка чинится кнопкой, а не JSON-профилем."""  # ruff: ignore[ambiguous-unicode-character-docstring]
        self._start(upload=self._csv("Terms.csv", GLOSSARY_PAIRS_CSV), slug=self.slug)
        draft = self._draft()

        response = self.client.post(
            reverse("loc-kit-glossary-preview", kwargs={"token": draft.token}),
            {"action": "relayout", "layout": "flat"},
        )
        self.assertEqual(response.status_code, 302)
        draft.refresh_from_db()
        # Flat reading turns every description row into its own term.
        self.assertEqual(json.loads(draft.preview_json)["term_count"], 5)

        self.client.post(
            reverse("loc-kit-glossary-preview", kwargs={"token": draft.token}),
            {"action": "relayout", "layout": "pairs"},
        )
        draft.refresh_from_db()
        self.assertEqual(json.loads(draft.preview_json)["term_count"], 2)

    @override_settings(LOC_KIT_PROFILE_ANALYSIS_ENABLED=False)
    def test_unknown_layout_is_refused(self) -> None:
        self._start(upload=self._csv("Terms.csv", GLOSSARY_PAIRS_CSV), slug=self.slug)
        response = self.client.post(
            reverse("loc-kit-glossary-preview", kwargs={"token": self._draft().token}),
            {"action": "relayout", "layout": "../../etc"},
        )
        self.assertEqual(response.status_code, 404)

    @override_settings(LOC_KIT_PROFILE_ANALYSIS_ENABLED=False)
    def test_term_with_outer_whitespace_is_trimmed_not_refused(self) -> None:
        """TBX не хранит внешний пробел: обрезаем и предупреждаем."""
        self._start(
            upload=self._csv("Terms.csv", "ru,en\nГерой ,Hero\nМеч,Sword\n"),
            slug=self.slug,
        )
        draft = self._draft()
        draft.refresh_from_db()

        self.assertEqual(draft.state, LocKitImportDraft.State.PREVIEW_READY)
        preview = json.loads(draft.preview_json)
        self.assertEqual(preview["terms"][0]["source"], "Герой")
        self.assertTrue(
            any("trimmed" in warning for warning in preview["warnings"]),
            preview["warnings"],
        )

    @override_settings(LOC_KIT_PROFILE_ANALYSIS_ENABLED=False)
    def test_stored_preview_warnings_are_bounded(self) -> None:
        """
        Warnings must not be an unbounded write path into the draft.

        One note is emitted per skipped row, and the row count comes from the
        uploaded file. Errors and sample terms are already capped; without a
        cap here the draft row and the preview page grow with the upload.
        """
        # Alternate "has a ru term" / "has only en": every second row is
        # skipped and contributes a note of its own.
        body = "".join(
            f"термин{index},term{index}\n,stray{index}\n" for index in range(60)
        )
        self._start(upload=self._csv("Sparse.csv", "ru,en\n" + body), slug=self.slug)
        draft = self._draft()
        draft.refresh_from_db()
        self.assertEqual(draft.state, LocKitImportDraft.State.PREVIEW_READY)

        warnings = json.loads(draft.preview_json)["warnings"]
        self.assertEqual(len(warnings), PREVIEW_WARNING_LIMIT + 1)
        self.assertIn("more warning", warnings[-1])

    def test_failed_draft_insert_leaves_no_orphaned_file(self) -> None:
        """
        Storage is not transactional.

        The upload lands on disk before the row is inserted, so a failed
        insert must not strand a file the row-driven cleanup can never see.
        """
        storage = LOC_KIT_DRAFT_STORAGE

        def before() -> set[str]:
            try:
                return set(storage.listdir("drafts")[1])
            except FileNotFoundError:
                return set()

        with (
            patch.object(LocKitImportDraft, "save", side_effect=DatabaseError("boom")),
            self.assertRaises(DatabaseError),
        ):
            self._start()

        self.assertFalse(LocKitImportDraft.objects.exists())
        self.assertEqual(before(), set())

    def test_disabled_analyzer_offers_manual_profile_upload(self) -> None:
        self._start()
        draft = self._draft()
        response = self._select_sheet(draft, "Glossary")

        draft.refresh_from_db()
        self.assertEqual(draft.sheet, "Glossary")
        self.assertEqual(draft.state, LocKitImportDraft.State.SHEET_SELECTED)
        # Analysis is off by default, so no profile was produced.
        self.assertEqual(draft.profile_json, "")
        page = self.client.get(response["Location"])
        self.assertContains(page, "Upload corrected profile")

    def test_stage_templates_never_nest_a_form(self) -> None:
        """
        Crispy must not emit a form tag of its own.

        Both stage templates supply their own <form> and submit button. When
        the helper wraps another one the source stays balanced, so only a
        real parser shows the damage: it closes the outer form at crispy's
        </form> and every later control - including submit - lands outside
        any form. The page renders normally and no click can post it.
        """
        self._start()
        draft = self._draft()

        pages = {
            "sheet": self.client.get(
                reverse("loc-kit-sheet-select", kwargs={"token": draft.token})
            ),
            "preview": self.client.get(
                self._select_sheet(draft, "Glossary")["Location"]
            ),
        }
        for stage, page in pages.items():
            depth = 0
            for tag in re.finditer(r"</?form\b", page.content.decode()):
                depth += -1 if tag.group().startswith("</") else 1
                self.assertLessEqual(
                    depth, 1, f"{stage}: a form is nested inside another form"
                )

    @override_settings(LOC_KIT_PROFILE_ANALYSIS_ENABLED=False)
    def test_sheet_selection_is_not_rate_limited_without_analysis(self) -> None:
        """
        Picking a worksheet must stay free.

        A multi-sheet workbook needs several POSTs before the operator even
        reaches the sheet they want, and with the analyzer off none of them
        can reach a provider. Spending the analysis budget here locked the
        operator out of their own upload on the default configuration.
        """
        self._start()
        draft = self._draft()
        # Superusers bypass check_rate_limit, so the regression is only
        # visible as an ordinary user. The permission gate has its own test.
        self.user.is_superuser = False
        self.user.save()

        # Comfortably more than RATELIMIT_LOC_KIT_ANALYSIS_ATTEMPTS (3).
        with patch(
            "weblate.trans.views.create.get_creatable_projects",
            return_value=Project.objects.filter(pk=self.project.pk),
        ):
            preview = reverse("loc-kit-glossary-preview", kwargs={"token": draft.token})
            for attempt in range(6):
                response = self._select_sheet(draft, "Glossary")
                # A throttled POST also answers 302, but back to itself.
                # Only the destination distinguishes accepted from refused.
                self.assertEqual(response["Location"], preview, f"attempt {attempt}")

        draft.refresh_from_db()
        self.assertEqual(draft.sheet, "Glossary")
        self.assertEqual(draft.state, LocKitImportDraft.State.SHEET_SELECTED)

    @override_settings(
        LOC_KIT_PROFILE_ANALYSIS_ENABLED=True,
        LOC_KIT_PROFILE_OPENROUTER_KEY="sk-test-secret-do-not-leak",
        LOC_KIT_PROFILE_OPENROUTER_MODEL="openai/gpt-4o",
    )
    @http_mock.activate
    def test_analysis_attempts_are_capped_per_session(self) -> None:
        """
        The provider budget is bounded and the lockout stays recoverable.

        Superusers bypass check_rate_limit entirely, so the cap can only be
        observed as an ordinary user.
        """
        http_mock.register(
            "POST",
            "https://openrouter.ai/api/v1/chat/completions",
            json={"choices": [{"message": {"content": "not json"}}]},
        )
        self._start()
        draft = self._draft()
        # Demote only once the draft exists: the upload gate itself needs a
        # user the wizard would accept. The gate is covered separately, so
        # hold it open here and let the cap be the only thing under test.
        self.user.is_superuser = False
        self.user.save()

        attempts = settings.RATELIMIT_LOC_KIT_ANALYSIS_ATTEMPTS
        with patch(
            "weblate.trans.views.create.get_creatable_projects",
            return_value=Project.objects.filter(pk=self.project.pk),
        ):
            for _ in range(attempts):
                self._select_sheet(draft)
            self.assertEqual(len(http_mock.calls), attempts)

            # The next attempt is refused before any request leaves the
            # server, and the message keeps the manual profile route open.
            response = self._select_sheet(draft)
            self.assertEqual(len(http_mock.calls), attempts)
            page = self.client.get(response["Location"])
        self.assertContains(page, "Upload a profile to continue")

    @http_mock.activate
    def test_manual_profile_produces_a_preview_without_any_outbound_request(
        self,
    ) -> None:
        self._start()
        draft = self._draft()
        self._select_sheet(draft)

        response = self._upload_profile(draft, _glossary_profile("Glossary"))
        # A correction is revalidated locally; the analyzer is never called.
        self.assertEqual(len(http_mock.calls), 0)

        draft.refresh_from_db()
        self.assertEqual(draft.state, LocKitImportDraft.State.PREVIEW_READY)
        preview = json.loads(draft.preview_json)
        self.assertEqual(preview["source_language"], "ru")
        self.assertEqual(preview["target_languages"], ["en"])
        self.assertEqual(preview["term_count"], 2)
        self.assertContains(response, "Герой")

    def test_confirmed_glossary_becomes_a_live_tbx_component(self) -> None:
        self._start()
        draft = self._draft()
        self._select_sheet(draft)
        self._upload_profile(draft, _glossary_profile("Glossary"))
        self._confirm()

        component = Component.objects.get(slug=self.slug)
        self.assertEqual(component.file_format, "tbx")
        self.assertEqual(component.filemask, "tbx/*.tbx")
        self.assertEqual(component.template, "")
        self.assertTrue(component.is_glossary)
        self.assertEqual(component.source_language.code, "ru")

        # A glossary is created for every project language; the two the kit
        # actually carries must be among them.
        codes = set(component.translation_set.values_list("language__code", flat=True))
        self.assertLessEqual({"en", "ru"}, codes)

        unit = component.translation_set.get(language__code="en").unit_set.get(
            context='["Персонажи","Герой"]'
        )
        self.assertEqual(unit.source, "Герой")
        self.assertEqual(unit.target, "Hero")

        # The draft and its file are gone once the component exists.
        self.assertFalse(LocKitImportDraft.objects.exists())

    def test_created_glossary_matches_a_component_with_the_same_source(self) -> None:
        self._start()
        draft = self._draft()
        self._select_sheet(draft)
        self._upload_profile(draft, _glossary_profile("Glossary"))
        self._confirm()

        glossary = Component.objects.get(slug=self.slug)
        self.assertTrue(glossary.is_glossary)
        # Glossary matching keys off the source language.
        self.assertEqual(
            glossary.source_language.code,
            "ru",
            "a glossary only helps components sharing its source language",
        )

    # ----------------------------------------------------------------- #
    # Multi-sheet
    # ----------------------------------------------------------------- #

    def test_multiple_worksheets_require_an_explicit_choice(self) -> None:
        buffer = io.BytesIO()
        workbook = Workbook()
        first = workbook.active
        first.title = "Glossary"
        for row in csv.reader(io.StringIO(GLOSSARY_CSV)):
            if row:
                first.append(row)
        second = workbook.create_sheet("Other")
        second.append(["domain", "ru", "en", "note_ru", "note_en"])
        second.append(["X", "Щит", "Shield", "", ""])
        workbook.save(buffer)

        self._start(
            upload=SimpleUploadedFile(
                "Book.xlsx",
                buffer.getvalue(),
                content_type=(
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                ),
            )
        )
        draft = self._draft()
        page = self.client.get(
            reverse("loc-kit-sheet-select", kwargs={"token": draft.token})
        )
        self.assertContains(page, "Glossary")
        self.assertContains(page, "Other")

        # A profile naming the unselected sheet is refused.
        self._select_sheet(draft, "Glossary")
        self._upload_profile(draft, _glossary_profile("Other"))
        draft.refresh_from_db()
        self.assertNotEqual(draft.state, LocKitImportDraft.State.PREVIEW_READY)

    # ----------------------------------------------------------------- #
    # Rejections: nothing is ever created
    # ----------------------------------------------------------------- #

    def test_draft_endpoints_enforce_the_same_gate_as_the_wizard(self) -> None:
        """
        A draft must never be a laxer way in than the ordinary wizard.

        The wizard gates component creation on get_creatable_projects(), which
        is managed_projects intersected with valid billing. A draft endpoint
        that only checked "project.edit" would let a user whose billing has
        lapsed create a component the wizard refuses.
        """
        self._start()
        draft = self._draft()

        # Drop the user out of the creatable set the wizard itself uses.
        self.user.is_superuser = False
        self.user.save()

        with patch(
            "weblate.trans.views.create.get_creatable_projects",
            return_value=Project.objects.none(),
        ):
            for name in (
                "loc-kit-sheet-select",
                "loc-kit-glossary-preview",
                "loc-kit-glossary-confirm",
            ):
                response = self.client.get(reverse(name, kwargs={"token": draft.token}))
                self.assertEqual(response.status_code, 404, name)

        self.assertFalse(Component.objects.filter(slug=self.slug).exists())

    def test_confirm_never_writes_into_another_components_repository(self) -> None:
        """
        The staged repo must land on the component's own path.

        LocalRepository.from_files removes an existing target before cloning.
        Building full_path from the draft's category while the form saves a
        different one aimed the write at another component's VCS directory:
        slug uniqueness is per (project, category), so a slug that collides
        at the draft's category level is perfectly valid at another one.
        """
        victim = self.component
        victim_path = victim.full_path

        # A second category makes the victim's slug reusable, so the final
        # form accepts it while the draft still carries category=None.
        category = Category.objects.create(
            project=self.project, name="Elsewhere", slug="elsewhere"
        )

        self._start()
        draft = self._draft()
        self.assertIsNone(draft.category)
        self._select_sheet(draft)
        self._upload_profile(draft, _glossary_profile("Glossary"))

        targets: list[str] = []
        real_from_files = LocalRepository.from_files

        def spy(target, files):
            targets.append(target)
            return real_from_files(target, files)

        with modify_settings(INSTALLED_APPS={"remove": "weblate.billing"}):
            response = self.client.post(
                reverse("loc-kit-glossary-preview", kwargs={"token": draft.token}),
                {"action": "confirm"},
                follow=True,
            )
            form = response.context["form"]
            params = {field: form[field].value() or "" for field in form.fields}
            params.pop("inherit_new_lang", None)
            params["new_lang"] = "none"
            # Aim the write at the victim: its slug, but another category.
            params["slug"] = victim.slug
            params["category"] = category.pk
            with patch("weblate.trans.views.create.LocalRepository.from_files", spy):
                self.client.post(
                    reverse("loc-kit-glossary-confirm", kwargs={"token": draft.token}),
                    params,
                    follow=True,
                )

        self.assertNotIn(
            victim_path, targets, "the staged repo was written over another component"
        )

    def test_invalid_profile_creates_no_component(self) -> None:
        self._start()
        draft = self._draft()
        self._select_sheet(draft)
        broken = _glossary_profile("Glossary")
        broken["components"][0]["grammar"]["notes"][0]["header"] = "wrong_header"
        self._upload_profile(draft, broken)

        draft.refresh_from_db()
        self.assertEqual(draft.profile_json, "")
        self.assertFalse(Component.objects.filter(slug=self.slug).exists())

    def test_source_only_profile_creates_no_component(self) -> None:
        self._start()
        draft = self._draft()
        self._select_sheet(draft)
        source_only = _glossary_profile("Glossary")
        source_only["components"][0]["initial_target_languages"] = []
        self._upload_profile(draft, source_only)

        draft.refresh_from_db()
        self.assertEqual(draft.profile_json, "")

    def test_confirm_before_a_preview_exists_is_refused(self) -> None:
        self._start()
        draft = self._draft()
        self._select_sheet(draft)
        response = self.client.post(
            reverse("loc-kit-glossary-preview", kwargs={"token": draft.token}),
            {"action": "confirm"},
        )
        self.assertEqual(response.status_code, 404)
        self.assertFalse(Component.objects.filter(slug=self.slug).exists())

    def test_cancel_deletes_the_draft_and_creates_nothing(self) -> None:
        self._start()
        draft = self._draft()
        self._select_sheet(draft)
        self._upload_profile(draft, _glossary_profile("Glossary"))

        self.client.post(
            reverse("loc-kit-glossary-preview", kwargs={"token": draft.token}),
            {"action": "cancel"},
            follow=True,
        )
        self.assertFalse(LocKitImportDraft.objects.exists())
        self.assertFalse(Component.objects.filter(slug=self.slug).exists())

    def test_another_user_cannot_touch_the_draft(self) -> None:
        self._start()
        draft = self._draft()

        other = User.objects.create_user("intruder", "intruder@example.com", "x")
        other.is_superuser = True
        other.save()
        self.client.force_login(other)

        for name in (
            "loc-kit-sheet-select",
            "loc-kit-glossary-preview",
            "loc-kit-glossary-confirm",
        ):
            response = self.client.get(reverse(name, kwargs={"token": draft.token}))
            self.assertEqual(response.status_code, 404, name)

    def test_expired_draft_is_unavailable(self) -> None:
        self._start()
        draft = self._draft()
        draft.expires_at = timezone.now() - timedelta(minutes=1)
        draft.save(update_fields=["expires_at"])

        response = self.client.get(
            reverse("loc-kit-sheet-select", kwargs={"token": draft.token})
        )
        self.assertEqual(response.status_code, 404)

    def test_unknown_token_is_unavailable(self) -> None:
        response = self.client.get(
            reverse("loc-kit-sheet-select", kwargs={"token": uuid.uuid4()})
        )
        self.assertEqual(response.status_code, 404)

    # ----------------------------------------------------------------- #
    # The final form cannot repoint the component
    # ----------------------------------------------------------------- #

    def test_tampered_final_post_is_overridden_by_the_draft(self) -> None:
        self._start()
        draft = self._draft()
        self._select_sheet(draft)
        self._upload_profile(draft, _glossary_profile("Glossary"))

        with modify_settings(INSTALLED_APPS={"remove": "weblate.billing"}):
            response = self.client.post(
                reverse("loc-kit-glossary-preview", kwargs={"token": draft.token}),
                {"action": "confirm"},
                follow=True,
            )
            form = response.context["form"]
            params = {field: form[field].value() or "" for field in form.fields}
            params.pop("inherit_new_lang", None)
            params["new_lang"] = "none"
            # Try to repoint every conversion-derived field.
            params["file_format"] = "po-mono"
            params["filemask"] = "evil/*.po"
            params["template"] = "evil.po"
            params["is_glossary"] = ""
            params["source_language"] = Language.objects.get(code="en").pk
            self.client.post(
                reverse("loc-kit-glossary-confirm", kwargs={"token": draft.token}),
                params,
                follow=True,
            )

        component = Component.objects.get(slug=self.slug)
        self.assertEqual(component.file_format, "tbx")
        self.assertEqual(component.filemask, "tbx/*.tbx")
        self.assertEqual(component.template, "")
        self.assertTrue(component.is_glossary)
        self.assertEqual(component.source_language.code, "ru")

    @override_settings(LOC_KIT_PROFILE_ANALYSIS_ENABLED=False)
    def test_semicolon_kit_gets_a_deterministic_preview(self) -> None:
        self._start(
            upload=self._csv("Temple.csv", GLOSSARY_SEMICOLON_CSV), slug=self.slug
        )
        draft = self._draft()
        draft.refresh_from_db()

        self.assertEqual(draft.state, LocKitImportDraft.State.PREVIEW_READY)
        preview = json.loads(draft.preview_json)
        self.assertEqual(preview["source_language"], "ru")
        self.assertEqual(preview["target_languages"], ["en"])
        self.assertEqual(preview["term_count"], 2)
        self.assertIn("мужской род", preview["terms"][0]["source_explanation"])

    @override_settings(LOC_KIT_PROFILE_ANALYSIS_ENABLED=False)
    def test_semicolon_partial_kit_maps_the_vendor_code(self) -> None:
        self._start(
            upload=self._csv("Terms.csv", GLOSSARY_SEMICOLON_PARTIAL_CSV),
            slug=self.slug,
        )
        draft = self._draft()
        draft.refresh_from_db()

        self.assertEqual(draft.state, LocKitImportDraft.State.PREVIEW_READY)
        preview = json.loads(draft.preview_json)
        self.assertIn("zh_Hant", preview["target_languages"])

    @override_settings(LOC_KIT_PROFILE_ANALYSIS_ENABLED=False)
    def test_id_partial_kit_creates_a_component_with_untranslated_terms(self) -> None:
        self._start(
            upload=self._csv("Terms.csv", GLOSSARY_ID_PARTIAL_CSV), slug=self.slug
        )
        draft = self._draft()
        draft.refresh_from_db()

        self.assertEqual(draft.state, LocKitImportDraft.State.PREVIEW_READY)
        preview = json.loads(draft.preview_json)
        self.assertEqual(preview["source_language"], "ru")
        self.assertEqual(preview["term_count"], 3)
        self.assertIn("zh_Hant", preview["target_languages"])
        self.assertTrue(
            any("untranslated" in warning for warning in preview["warnings"]),
            preview["warnings"],
        )

        self._confirm()
        component = Component.objects.get(slug=self.slug, project=self.project)
        self.assertTrue(component.is_glossary)
        self.assertEqual(component.source_language.code, "ru")

        translation = component.translation_set.get(language__code="zh_Hant")
        filled = translation.unit_set.get(source="Аки")
        self.assertEqual(filled.target, "阿姬")
        blank = translation.unit_set.get(source="Джо")
        self.assertEqual(blank.target, "")
        self.assertEqual(blank.state, STATE_EMPTY)


class LocKitUpdateDraftModelTest(ViewTestCase):
    """A glossary-update draft is bound to one existing glossary component."""

    CREATE_GLOSSARIES: bool = True

    def test_draft_binds_to_existing_glossary(self) -> None:
        glossary = self.project.glossaries[0]
        draft = LocKitImportDraft.objects.create(
            owner=self.user,
            session_key="x" * 10,
            project=self.project,
            slug=glossary.slug,
            name=glossary.name,
            source_filename="Terms.csv",
            target_component=glossary,
        )
        self.assertEqual(draft.target_component, glossary)

    def test_deleting_target_glossary_deletes_its_draft(self) -> None:
        glossary = self.project.glossaries[0]
        LocKitImportDraft.objects.create(
            owner=self.user,
            session_key="x" * 10,
            project=self.project,
            slug=glossary.slug,
            name=glossary.name,
            source_filename="Terms.csv",
            target_component=glossary,
        )
        glossary.delete()
        self.assertFalse(LocKitImportDraft.objects.exists())


class LocKitGlossaryUpdateGateTest(ViewTestCase):
    """Update drafts are gated on the target glossary, not on creation."""

    CREATE_GLOSSARIES: bool = True

    def setUp(self) -> None:
        super().setUp()
        self.glossary = self.project.glossaries[0]
        # A component-level upload.perform resolves through its child
        # translations, and a freshly created glossary has none. Seed one
        # target language so the restricted-role tests see a real glossary.
        self.user.is_superuser = True
        self.user.save()
        self.glossary.add_new_language(Language.objects.get(code="cs"), None)
        self.user.is_superuser = False
        self.user.save()
        self.user.clear_permissions_cache()

    def _draft(self, component=None) -> LocKitImportDraft:
        component = component or self.glossary
        session = self.client.session
        if not session.session_key:
            session.create()
        draft = LocKitImportDraft(
            owner=self.user,
            session_key=session.session_key,
            project=self.project,
            slug=component.slug,
            name=component.name,
            source_filename="Terms.csv",
            target_component=component,
        )
        draft.uploaded.save("Terms.csv", ContentFile(b"ru,en\n"), save=False)
        draft.save()
        return draft

    def _restrict_to_upload(self) -> None:
        """Keep upload.perform, drop translation.add/glossary.add/unit.add."""
        self.user.groups.clear()
        role = Role.objects.create(name=f"Update gate {uuid.uuid4().hex[:6]}")
        role.permissions.add(
            Permission.objects.get(codename="glossary.upload"),
            # check_upload rewrites glossary unit.edit to glossary.edit
            Permission.objects.get(codename="glossary.edit"),
            Permission.objects.get(codename="unit.edit"),
        )
        group = Group.objects.create(name=role.name, language_selection=SELECTION_ALL)
        group.roles.add(role)
        group.projects.add(self.project)
        self.user.groups.add(group)
        self.user.clear_permissions_cache()

    def test_update_draft_requires_upload_permission(self) -> None:
        draft = self._draft()
        self.user.groups.clear()
        self.user.clear_permissions_cache()
        for name in ("loc-kit-sheet-select", "loc-kit-glossary-preview"):
            response = self.client.get(reverse(name, kwargs={"token": draft.token}))
            self.assertEqual(response.status_code, 404, name)

    def test_locked_glossary_hides_its_update_draft(self) -> None:
        draft = self._draft()
        self.user.is_superuser = True
        self.user.save()
        self.glossary.locked = True
        self.glossary.save()
        response = self.client.get(
            reverse("loc-kit-glossary-preview", kwargs={"token": draft.token})
        )
        self.assertEqual(response.status_code, 404)

    def test_draft_pointing_at_a_regular_component_is_rejected(self) -> None:
        draft = self._draft(component=self.component)
        self.user.is_superuser = True
        self.user.save()
        response = self.client.get(
            reverse("loc-kit-glossary-preview", kwargs={"token": draft.token})
        )
        self.assertEqual(response.status_code, 404)

    def test_upload_permission_alone_opens_the_preview(self) -> None:
        """
        translation.add is only needed when a language must be created.

        A user without glossary.add/unit.add still opens the preview; the
        apply path refuses those languages per language instead of giving a
        permission bypass (covered by the UpdateApply tests).
        """
        draft = self._draft()
        self._restrict_to_upload()
        self.assertTrue(self.user.has_perm("upload.perform", self.glossary))
        self.assertFalse(self.user.has_perm("translation.add", self.project))
        self.assertFalse(
            self.user.has_perm("unit.add", self.glossary.source_translation)
        )
        response = self.client.get(
            reverse("loc-kit-glossary-preview", kwargs={"token": draft.token})
        )
        self.assertEqual(response.status_code, 200)

    def test_confirm_rejects_an_update_draft(self) -> None:
        """
        Confirm requires the creatable-projects permission.

        An update draft is only gated by upload.perform on the existing
        glossary and must never reach this endpoint - that would let
        upload.perform on one glossary create an unrelated component.
        """
        draft = self._draft()
        # Reach the state the view actually gates on, so a passing test
        # proves the new guard fires rather than the pre-existing
        # PREVIEW_READY check that would 404 an untouched draft anyway.
        draft.state = LocKitImportDraft.State.PREVIEW_READY
        draft.preview_json = json.dumps({"source_language": "en"})
        draft.save(update_fields=["state", "preview_json"])
        self._restrict_to_upload()
        before = Component.objects.count()

        response = self.client.get(
            reverse("loc-kit-glossary-confirm", kwargs={"token": draft.token})
        )
        self.assertEqual(response.status_code, 404)
        response = self.client.post(
            reverse("loc-kit-glossary-confirm", kwargs={"token": draft.token}), {}
        )
        self.assertEqual(response.status_code, 404)
        self.assertEqual(Component.objects.count(), before)


class LocKitGlossaryUpdateStartTest(ViewTestCase):
    """An operator stages an append table from an existing glossary."""

    CREATE_GLOSSARIES: bool = True

    def setUp(self) -> None:
        super().setUp()
        self.glossary = self.project.glossaries[0]
        self.user.is_superuser = True
        self.user.save()

    def _url(self):
        return reverse(
            "loc-kit-glossary-update", kwargs={"path": self.glossary.get_url_path()}
        )

    def _table(self, body: str = GLOSSARY_LANG_ONLY_CSV):
        return SimpleUploadedFile("Terms.csv", body.encode(), content_type="text/csv")

    @override_settings(LOC_KIT_PROFILE_ANALYSIS_ENABLED=False)
    def test_update_upload_stages_a_bound_draft(self) -> None:
        response = self.client.post(self._url(), {"table": self._table()})

        draft = LocKitImportDraft.objects.get()
        self.assertEqual(draft.target_component, self.glossary)
        self.assertEqual(draft.project, self.project)
        # Single sheet: auto-skip mapped it straight to a preview.
        self.assertEqual(draft.state, LocKitImportDraft.State.PREVIEW_READY)
        self.assertRedirects(
            response,
            reverse("loc-kit-glossary-preview", kwargs={"token": draft.token}),
        )
        # The component is not touched by staging.
        self.glossary.refresh_from_db()
        self.assertEqual(self.glossary.source_translation.unit_set.count(), 0)

    def test_update_start_requires_upload_permission(self) -> None:
        # Superusers bypass every permission; the gate is only observable
        # on an ordinary account.
        self.user.is_superuser = False
        self.user.save()
        self.user.groups.clear()
        self.user.clear_permissions_cache()

        response = self.client.get(self._url())
        self.assertEqual(response.status_code, 404)
        response = self.client.post(self._url(), {"table": self._table()})
        self.assertEqual(response.status_code, 404)
        self.assertFalse(LocKitImportDraft.objects.exists())

    def test_update_stage_template_never_nests_a_form(self) -> None:
        """
        Same contract as the creation stages, for the update page.

        The template owns the single <form>; crispy must not emit one of
        its own. The rendered source stays balanced either way, so only
        walking tag depth exposes a nested form that browsers reject.
        """
        page = self.client.get(self._url())
        depth = 0
        for tag in re.finditer(r"</?form\b", page.content.decode()):
            depth += -1 if tag.group().startswith("</") else 1
            self.assertLessEqual(depth, 1, "a form is nested inside another form")


# --------------------------------------------------------------------------- #
# Append-only application service
# --------------------------------------------------------------------------- #

# ruff: ignore[module-import-not-at-top-of-file, import-private-name]
from loc_kit_ingest.parser import _context_key


def _append_preview(
    *terms: dict, source_language: str = "ru"
) -> loc_kit.GlossaryPreview:
    """
    Build a validated preview straight from term rows, no UI round trip.

    Each term is ``{"section": str, "values": {code: str}, "notes":
    {"source": str, "target": {code: str}}, "flags": tuple[str, ...]}``.
    The section column drives the same (section, term) context the writer
    derives, so previews can match units seeded through add_unit. Optional
    note and flag columns exist only when a term carries them, matching the
    record-map shape the flow accepts.
    """
    codes = [
        code
        for code in dict.fromkeys(value for term in terms for value in term["values"])
        if code != source_language
    ]
    has_source_notes = any(term.get("notes", {}).get("source") for term in terms)
    noted_targets = [
        code
        for code in codes
        if any(term.get("notes", {}).get("target", {}).get(code) for term in terms)
    ]
    has_source_flags = any(term.get("flags") for term in terms)

    # column 1 is the section field; languages follow.
    header = ["domain", source_language, *codes]
    note_fields = []
    column = len(header) + 1
    if has_source_notes:
        header.append("note_ru")
        note_fields.append(
            {
                "scope": "source",
                "column": column,
                "header": "note_ru",
                "row_offset": 0,
            }
        )
        column += 1
    for code in noted_targets:
        name = f"note_{code}"
        header.append(name)
        note_fields.append(
            {
                "scope": "target",
                "language": code,
                "column": column,
                "header": name,
                "row_offset": 0,
            }
        )
        column += 1
    source_flags_field = None
    if has_source_flags:
        header.append("flags")
        source_flags_field = {
            "column": column,
            "header": "flags",
            "row_offset": 0,
        }
        column += 1

    languages = [
        {
            "code": source_language,
            "xml_lang": source_language,
            "column": 2,
            "header": source_language,
        },
        *[
            {"code": code, "xml_lang": code, "column": i + 3, "header": code}
            for i, code in enumerate(codes)
        ],
    ]
    document = {
        "schema_version": 2,
        "components": [
            {
                "sheet": "Terms",
                "component": "glossary",
                "kind": "tbx",
                "source_lang": source_language,
                "header_row": 1,
                "languages": languages,
                "grammar": {
                    "type": "record-map",
                    "skip_rows": [],
                    "regions": [
                        {
                            "first_record_row": 2,
                            "last_record_row": 1 + len(terms),
                            "record_stride": 1,
                        }
                    ],
                    "term_row_offset": 0,
                    "allow_empty_targets": True,
                    "notes": note_fields,
                    "section_field": {
                        "column": 1,
                        "header": "domain",
                        "row_offset": 0,
                    },
                },
                "initial_target_languages": codes,
            }
        ],
    }
    if source_flags_field is not None:
        document["components"][0]["grammar"]["source_flags"] = source_flags_field
    rows = [header]
    for term in terms:
        row = [term.get("section", "Термины")]
        row.extend(term["values"].get(code, "") for code in [source_language, *codes])
        if has_source_notes:
            row.append(term.get("notes", {}).get("source", ""))
        row.extend(
            term.get("notes", {}).get("target", {}).get(code, "")
            for code in noted_targets
        )
        if has_source_flags:
            row.append(", ".join(term.get("flags", ())))
        rows.append(row)
    return loc_kit.validate_glossary_profile(
        profile_document=document,
        rows=rows,
        sheet_name="Terms",
        component_name="glossary",
    )


class LocKitGlossaryAppendServiceTest(ViewTestCase):
    """append_glossary_terms only ever appends, never overwrites."""

    CREATE_GLOSSARIES: bool = True

    def setUp(self) -> None:
        super().setUp()
        self.user.is_superuser = True
        self.user.save()
        self.factory = RequestFactory()
        self.glossary = self.project.glossaries[0]
        # The stock setup already mirrors every project language (cs, de)
        # into a fresh glossary; only languages outside the project can be
        # added later.
        self.cs = self.glossary.translation_set.get(language__code="cs")

    def _request(self):
        request = self.factory.post("/")
        request.user = self.user
        return request

    def _seed(self, section, en, cs="", *, explanation="Old meaning.", state=None):
        """Create one glossary term the way the stock UI would."""
        context = _context_key(section, en)
        # Adding through the source translation creates the source unit and
        # an empty target pair in every existing language.
        self.glossary.source_translation.add_unit(
            None, context, en, explanation=explanation, author=self.user
        )
        target_unit = self.cs.unit_set.get(context=context, source=en)
        if cs:
            target_unit.target = cs
            target_unit.state = STATE_TRANSLATED
            target_unit.save(update_fields=["target", "state"], same_content=True)
        if state is not None:
            target_unit.state = state
            target_unit.save(update_fields=["state"], same_content=True)
        return target_unit

    def test_existing_empty_target_is_never_filled(self) -> None:
        seeded = self._seed("Characters", "Hero")
        preview = _append_preview(
            {"section": "Characters", "values": {"en": "Hero", "cs": "Hrdina"}},
            source_language="en",
        )

        result = loc_kit.append_glossary_terms(self._request(), self.glossary, preview)

        seeded.refresh_from_db()
        self.assertEqual(seeded.target, "")
        # The explanation lives on the source unit; the target unit carries
        # its own (empty) one, and neither is rewritten by the append.
        self.assertEqual(seeded.explanation, "")
        self.assertEqual(seeded.source_unit.explanation, "Old meaning.")
        self.assertEqual(result.added_terms, 0)
        self.assertEqual(result.languages["cs"].existing, 1)
        self.assertEqual(result.languages["cs"].added, 0)

    def test_translated_and_approved_terms_untouched(self) -> None:
        translated = self._seed("One", "First", "Prvni")
        approved = self._seed("Two", "Second", "Druhy", state=STATE_APPROVED)
        preview = _append_preview(
            {"section": "One", "values": {"en": "First", "cs": "Jedna"}},
            {"section": "Two", "values": {"en": "Second", "cs": "Zwei"}},
            source_language="en",
        )

        result = loc_kit.append_glossary_terms(self._request(), self.glossary, preview)

        translated.refresh_from_db()
        approved.refresh_from_db()
        self.assertEqual(translated.target, "Prvni")
        self.assertEqual(approved.target, "Druhy")
        self.assertEqual(result.added_terms, 0)
        self.assertEqual(result.languages["cs"].existing, 2)

    def test_new_term_creates_every_nonempty_target(self) -> None:
        pl = self.glossary.add_new_language(Language.objects.get(code="pl"), None)
        preview = _append_preview(
            {
                "values": {"en": "Shield", "cs": "Stit", "pl": "Tarcza"},
                "notes": {"source": "Defense.", "target": {"cs": "Obrana."}},
            },
            source_language="en",
        )

        result = loc_kit.append_glossary_terms(self._request(), self.glossary, preview)

        self.assertEqual(result.added_terms, 1)
        cs_unit = self.cs.unit_set.get(source="Shield")
        pl_unit = pl.unit_set.get(source="Shield")
        self.assertEqual(cs_unit.target, "Stit")
        self.assertEqual(pl_unit.target, "Tarcza")
        # Meanings: target explanation on the target unit, source
        # explanation on its source unit.
        self.assertEqual(cs_unit.explanation, "Obrana.")
        source_unit = cs_unit.source_unit
        self.assertEqual(source_unit.explanation, "Defense.")
        self.assertIn("terminology", source_unit.extra_flags)
        self.assertEqual(result.languages["cs"].added, 1)
        self.assertEqual(result.languages["pl"].added, 1)

    def test_new_terms_merge_source_flags_with_terminology(self) -> None:
        preview = _append_preview(
            {"values": {"en": "HeroCraft", "cs": "HeroCraft"}, "flags": ("read-only",)},
            {"values": {"en": "Vessel", "cs": "Plavidlo"}, "flags": ("forbidden",)},
            source_language="en",
        )

        self.assertEqual(preview.terms[0].source_flags, ("read-only",))
        loc_kit.append_glossary_terms(self._request(), self.glossary, preview)

        read_only = self.glossary.source_translation.unit_set.get(source="HeroCraft")
        forbidden = self.glossary.source_translation.unit_set.get(source="Vessel")
        self.assertEqual(
            set(read_only.extra_flags.split(", ")),
            {"read-only", "terminology"},
        )
        self.assertEqual(
            set(forbidden.extra_flags.split(", ")),
            {"forbidden", "terminology"},
        )

    def test_new_exact_flag_is_stored_only_on_nonempty_target(self) -> None:
        preview = _append_preview(
            {
                "values": {"en": "Canonical", "cs": "Kanonický", "pl": ""},
                "flags": ("exact",),
            },
            source_language="en",
        )

        result = loc_kit.append_glossary_terms(self._request(), self.glossary, preview)

        source = self.glossary.source_translation.unit_set.get(source="Canonical")
        target = self.cs.unit_set.get(source="Canonical")
        self.assertIn("terminology", source.extra_flags)
        self.assertNotIn("exact", source.extra_flags)
        self.assertIn("exact", target.extra_flags)
        self.assertEqual(result.languages["pl"].blank, 1)

    def test_incoming_source_flags_never_rewrite_existing_terms(self) -> None:
        seeded = self._seed("Characters", "Hero", "Hrdina")
        seeded.source_unit.update_extra_flags("terminology, read-only", self.user)
        preview = _append_preview(
            {
                "section": "Characters",
                "values": {"en": "Hero", "cs": "Hrdina"},
                "flags": ("forbidden",),
            },
            source_language="en",
        )

        loc_kit.append_glossary_terms(self._request(), self.glossary, preview)

        seeded.source_unit.refresh_from_db()
        self.assertIn("read-only", seeded.source_unit.extra_flags)
        self.assertNotIn("forbidden", seeded.source_unit.extra_flags)

    def test_blank_target_cell_is_skipped_not_created(self) -> None:
        """Blank pl for a new term: Czech lands, Polish stays absent."""
        preview = _append_preview(
            {"values": {"en": "Bow", "cs": "Luk", "pl": ""}},
            source_language="en",
        )

        result = loc_kit.append_glossary_terms(self._request(), self.glossary, preview)

        self.assertEqual(result.added_terms, 1)
        self.assertTrue(self.cs.unit_set.filter(source="Bow").exists())
        self.assertFalse(
            self.glossary.translation_set.filter(language__code="pl").exists()
        )
        self.assertEqual(result.languages["cs"].added, 1)
        self.assertEqual(result.languages["pl"].blank, 1)
        self.assertEqual(result.languages["pl"].added, 0)

    def test_absent_column_does_not_block_other_languages(self) -> None:
        """A glossary language with no column is absent, not an error."""
        pl = self.glossary.add_new_language(Language.objects.get(code="pl"), None)
        preview = _append_preview(
            {"values": {"en": "Potion", "cs": "Lektvar"}}, source_language="en"
        )

        result = loc_kit.append_glossary_terms(self._request(), self.glossary, preview)

        self.assertEqual(result.added_terms, 1)
        self.assertEqual(result.languages["cs"].added, 1)
        self.assertTrue(result.languages["pl"].absent)
        # Terminology sync still mirrors the new source term into every
        # glossary language, but the absent column never writes a target.
        self.assertEqual(pl.unit_set.get(source="Potion").target, "")

    def test_nonempty_column_creates_a_missing_language(self) -> None:
        preview = _append_preview(
            {"values": {"en": "Mage", "ja": "Mahou"}}, source_language="en"
        )

        result = loc_kit.append_glossary_terms(self._request(), self.glossary, preview)

        ja = self.glossary.translation_set.get(language__code="ja")
        self.assertEqual(ja.unit_set.get(source="Mage").target, "Mahou")
        self.assertEqual(result.languages["ja"].added, 1)
        self.assertEqual(result.added_terms, 1)

    def test_restricted_user_gets_unavailable_language_only(self) -> None:
        """
        Without translation.add the Japanese column is unavailable.

        The Czech translation the user may still write to is added anyway.
        """
        self.user.is_superuser = False
        self.user.save()
        self.user.groups.clear()
        role = Role.objects.create(name=f"Append {uuid.uuid4().hex[:6]}")
        role.permissions.add(
            Permission.objects.get(codename="glossary.upload"),
            Permission.objects.get(codename="glossary.edit"),
            Permission.objects.get(codename="glossary.add"),
            Permission.objects.get(codename="unit.edit"),
        )
        group = Group.objects.create(name=role.name, language_selection=SELECTION_ALL)
        group.roles.add(role)
        group.projects.add(self.project)
        self.user.groups.add(group)
        self.user.clear_permissions_cache()
        self.assertFalse(self.user.has_perm("translation.add", self.project))
        self.assertTrue(self.user.has_perm("unit.add", self.cs))

        preview = _append_preview(
            {"values": {"en": "Forest", "cs": "Les", "ja": "Mori"}},
            source_language="en",
        )
        result = loc_kit.append_glossary_terms(self._request(), self.glossary, preview)

        self.assertEqual(result.languages["cs"].added, 1)
        self.assertEqual(result.languages["ja"].unavailable_count, 1)
        self.assertNotEqual(result.languages["ja"].unavailable, "")
        self.assertEqual(result.languages["ja"].added, 0)
        self.assertFalse(
            self.glossary.translation_set.filter(language__code="ja").exists()
        )
        self.assertEqual(result.added_terms, 1)

    def test_missing_notes_do_not_block_new_terms(self) -> None:
        preview = _append_preview(
            {"values": {"en": "Staff", "cs": "Hul"}}, source_language="en"
        )

        result = loc_kit.append_glossary_terms(self._request(), self.glossary, preview)

        self.assertEqual(result.added_terms, 1)
        source_unit = self.glossary.source_translation.unit_set.get(source="Staff")
        self.assertEqual(source_unit.explanation, "")

    def test_incoming_notes_never_rewrite_existing_terms(self) -> None:
        seeded = self._seed("Characters", "Hero", "Hrdina")
        preview = _append_preview(
            {
                "section": "Characters",
                "values": {"en": "Hero", "cs": "Hrdina"},
                "notes": {
                    "source": "A brand new description.",
                    "target": {"cs": "Brand new note."},
                },
            },
            source_language="en",
        )

        result = loc_kit.append_glossary_terms(self._request(), self.glossary, preview)

        seeded.refresh_from_db()
        self.assertEqual(seeded.explanation, "")
        self.assertEqual(seeded.source_unit.explanation, "Old meaning.")
        self.assertEqual(result.languages["cs"].existing, 1)
        self.assertEqual(result.languages["cs"].added, 0)

    def test_same_source_under_different_contexts_collides(self) -> None:
        self._seed("Characters", "Hero", "Hrdina")
        # The glossary already knows the source under another section ...
        preview = _append_preview(
            {"section": "Weapons", "values": {"en": "Hero", "cs": "Sampion"}},
            source_language="en",
        )
        with self.assertRaises(loc_kit.GlossaryAppendCollisionError) as ctx:
            loc_kit.append_glossary_terms(self._request(), self.glossary, preview)
        self.assertEqual(ctx.exception.conflicts[0][0], "Hero")
        self.assertFalse(
            self.cs.unit_set.filter(source="Hero").exclude(target="Hrdina").exists()
        )

        # ... and two incoming rows may not introduce the same source either.
        preview = _append_preview(
            {"section": "One", "values": {"en": "Sword", "cs": "Mec"}},
            {"section": "Two", "values": {"en": "Sword", "cs": "Cepel"}},
            source_language="en",
        )
        with self.assertRaises(loc_kit.GlossaryAppendCollisionError):
            loc_kit.append_glossary_terms(self._request(), self.glossary, preview)
        self.assertFalse(self.cs.unit_set.filter(source="Sword").exists())

    def test_counters_mix_added_blank_existing_and_absent(self) -> None:
        self._seed("One", "Old", "Stary")
        self.glossary.translation_set.get(language__code="de")
        preview = _append_preview(
            {"section": "Two", "values": {"en": "New", "cs": "Novy", "pl": ""}},
            {"section": "One", "values": {"en": "Old", "cs": "Ancient"}},
            source_language="en",
        )

        result = loc_kit.append_glossary_terms(self._request(), self.glossary, preview)

        cs, pl, de = (
            result.languages["cs"],
            result.languages["pl"],
            result.languages["de"],
        )
        self.assertEqual((cs.added, cs.existing, cs.blank), (1, 1, 0))
        self.assertEqual((pl.added, pl.existing, pl.blank), (0, 1, 1))
        self.assertTrue(de.absent)

    def test_content_error_rolls_back_every_new_unit(self) -> None:
        real_add_unit = Translation.add_unit
        calls = []

        def failing_add_unit(translation, *args, **kwargs):
            calls.append(translation.language.code)
            if len(calls) == 2:
                msg = "simulated storage failure"
                raise ValueError(msg)
            return real_add_unit(translation, *args, **kwargs)

        preview = _append_preview(
            {"values": {"en": "First", "cs": "Prvni"}},
            {"values": {"en": "Second", "cs": "Druhy"}},
            source_language="en",
        )
        with (
            patch.object(Translation, "add_unit", failing_add_unit),
            self.assertRaises(ValueError),
        ):
            loc_kit.append_glossary_terms(self._request(), self.glossary, preview)

        self.assertFalse(
            self.cs.unit_set.filter(source__in=["First", "Second"]).exists()
        )
        self.assertFalse(
            self.glossary.source_translation.unit_set.filter(
                source__in=["First", "Second"]
            ).exists()
        )

    def test_lock_timeout_propagates_and_reapply_is_idempotent(self) -> None:
        preview = _append_preview(
            {"values": {"en": "Key", "cs": "Klic"}}, source_language="en"
        )

        with (
            patch.object(
                Component,
                "locked_for_update",
                side_effect=WeblateLockTimeoutError("locked", lock=None),
            ),
            self.assertRaises(WeblateLockTimeoutError),
        ):
            loc_kit.append_glossary_terms(self._request(), self.glossary, preview)
        self.assertFalse(self.cs.unit_set.filter(source="Key").exists())

        first = loc_kit.append_glossary_terms(self._request(), self.glossary, preview)
        self.assertEqual(first.added_terms, 1)
        # Replaying the same table adds nothing a second time.
        second = loc_kit.append_glossary_terms(self._request(), self.glossary, preview)
        self.assertEqual(second.added_terms, 0)
        self.assertEqual(second.languages["cs"].existing, 1)
        self.assertEqual(self.cs.unit_set.filter(source="Key").count(), 1)

    def test_failed_content_compensates_created_languages(self) -> None:
        preview = _append_preview(
            {"values": {"en": "Throne", "ja": "Za"}}, source_language="en"
        )

        def failing_add_unit(translation, *args, **kwargs):
            msg = "simulated content failure"
            raise ValueError(msg)

        with (
            patch.object(Translation, "add_unit", failing_add_unit),
            self.assertRaises(ValueError),
        ):
            loc_kit.append_glossary_terms(self._request(), self.glossary, preview)

        self.assertFalse(
            self.glossary.translation_set.filter(language__code="ja").exists()
        )

    def test_failed_compensation_is_reported_and_still_raises(self) -> None:
        preview = _append_preview(
            {"values": {"en": "Prison", "ja": "Vezeni"}}, source_language="en"
        )

        def failing_add_unit(translation, *args, **kwargs):
            msg = "simulated content failure"
            raise ValueError(msg)

        with (
            patch.object(Translation, "add_unit", failing_add_unit),
            patch.object(Translation, "remove", side_effect=OSError("fs down")),
            patch("weblate.utils.errors.report_error") as report,
            self.assertRaises(ValueError),
        ):
            loc_kit.append_glossary_terms(self._request(), self.glossary, preview)

        self.assertTrue(
            any(
                call.args and "compensate" in str(call.args[0])
                for call in report.call_args_list
            ),
            report.call_args_list,
        )
        # The language could not be removed either: it stays, visibly broken,
        # for the operator instead of being hidden behind a success message.
        self.assertTrue(
            self.glossary.translation_set.filter(language__code="ja").exists()
        )


# --------------------------------------------------------------------------- #
# Applying an append table from the preview view
# --------------------------------------------------------------------------- #


class LocKitGlossaryUpdateApplyTest(ViewTestCase):
    """The preview view applies the append service to an existing glossary."""

    CREATE_GLOSSARIES: bool = True

    def setUp(self) -> None:
        super().setUp()
        self.user.is_superuser = True
        self.user.save()
        self.glossary = self.project.glossaries[0]
        self.cs = self.glossary.translation_set.get(language__code="cs")

    def _stage(self, csv_body: str) -> LocKitImportDraft:
        response = self.client.post(
            reverse(
                "loc-kit-glossary-update",
                kwargs={"path": self.glossary.get_url_path()},
            ),
            {
                "table": SimpleUploadedFile(
                    "Terms.csv", csv_body.encode(), content_type="text/csv"
                )
            },
        )
        draft = LocKitImportDraft.objects.get()
        self.assertRedirects(
            response,
            reverse("loc-kit-glossary-preview", kwargs={"token": draft.token}),
        )
        return draft

    def _apply(self, draft, follow=True):
        return self.client.post(
            reverse("loc-kit-glossary-preview", kwargs={"token": draft.token}),
            {"action": "apply"},
            follow=follow,
        )

    def test_apply_imports_new_term_and_consumes_draft(self) -> None:
        draft = self._stage("en,cs\nEnglish,Czech\nShield,Stit\n")

        response = self._apply(draft)

        self.assertContains(response, "1 added")
        self.assertContains(response, "cs")
        self.assertEqual(self.cs.unit_set.get(source="Shield").target, "Stit")
        self.assertFalse(LocKitImportDraft.objects.exists())

    def test_apply_keeps_existing_empty_target_blank(self) -> None:
        context = _context_key("", "Shield")
        self.glossary.source_translation.add_unit(
            None, context, "Shield", author=self.user
        )
        draft = self._stage("en,cs\nEnglish,Czech\nShield,Stit\n")

        self._apply(draft)

        unit = self.cs.unit_set.get(context=context, source="Shield")
        self.assertEqual(unit.target, "")

    def test_apply_reports_blank_targets_as_success(self) -> None:
        # A fully blank column never reaches the service: inference drops
        # it. A partially filled one does, and its blank cells are skipped.
        draft = self._stage(
            "en,cs,pl\nEnglish,Czech,Polish\nBow,Luk,\nSword,Mec,Miecz\n"
        )

        response = self._apply(draft)

        self.assertContains(response, "2 new glossary terms added.")
        self.assertContains(response, "1 blank")
        self.assertEqual(self.cs.unit_set.get(source="Bow").target, "Luk")
        pl = self.glossary.translation_set.get(language__code="pl")
        self.assertEqual(pl.unit_set.get(source="Sword").target, "Miecz")
        # Terminology sync mirrors the new source term into pl with an
        # empty target; the blank cell itself wrote nothing.
        self.assertEqual(pl.unit_set.get(source="Bow").target, "")
        self.assertFalse(LocKitImportDraft.objects.exists())

    def test_apply_creates_missing_language_and_reports_it(self) -> None:
        draft = self._stage("en,ja\nEnglish,Japanese\nMage,Mahou\n")

        response = self._apply(draft)

        self.assertContains(response, "ja")
        self.assertContains(response, "1 added")
        ja = self.glossary.translation_set.get(language__code="ja")
        self.assertEqual(ja.unit_set.get(source="Mage").target, "Mahou")

    def test_apply_without_translation_add_keeps_english_and_flags_japanese(
        self,
    ) -> None:
        draft = self._stage("en,cs,ja\nEnglish,Czech,Japanese\nForest,Les,Mori\n")
        self.user.is_superuser = False
        self.user.save()
        self.user.groups.clear()
        role = Role.objects.create(name=f"Apply {uuid.uuid4().hex[:6]}")
        role.permissions.add(
            Permission.objects.get(codename="glossary.upload"),
            Permission.objects.get(codename="glossary.edit"),
            Permission.objects.get(codename="glossary.add"),
            Permission.objects.get(codename="unit.edit"),
        )
        group = Group.objects.create(name=role.name, language_selection=SELECTION_ALL)
        group.roles.add(role)
        group.projects.add(self.project)
        self.user.groups.add(group)
        self.user.clear_permissions_cache()
        response = self._apply(draft)

        self.assertContains(response, "1 added")
        self.assertContains(response, "Add language for translation")
        self.assertEqual(self.cs.unit_set.get(source="Forest").target, "Les")
        self.assertFalse(
            self.glossary.translation_set.filter(language__code="ja").exists()
        )
        self.assertFalse(LocKitImportDraft.objects.exists())

    def test_collision_returns_to_preview_and_keeps_draft(self) -> None:
        self.glossary.source_translation.add_unit(
            None, _context_key("Characters", "Hero"), "Hero", author=self.user
        )
        draft = self._stage("en,cs\nEnglish,Czech\nHero,Sampion\n")

        response = self._apply(draft)

        self.assertContains(response, "different section")
        # The seeded source unit keeps its empty cs pair; the incoming
        # value must never land in it.
        self.assertFalse(
            self.cs.unit_set.filter(source="Hero").exclude(target="").exists()
        )
        self.assertTrue(LocKitImportDraft.objects.filter(token=draft.token).exists())

    def test_apply_stores_notes_on_new_term(self) -> None:
        draft = self._stage(
            "en,cs,note\nEnglish,Czech,Note\nStaff,Hul,A wooden staff.\n"
        )

        self._apply(draft)

        source_unit = self.glossary.source_translation.unit_set.get(source="Staff")
        self.assertEqual(source_unit.explanation, "A wooden staff.")

    def test_apply_lock_timeout_keeps_draft_and_says_retry(self) -> None:
        draft = self._stage("en,cs\nEnglish,Czech\nKey,Klic\n")

        with patch.object(
            Component,
            "locked_for_update",
            side_effect=WeblateLockTimeoutError("locked", lock=None),
        ):
            response = self._apply(draft)

        self.assertContains(response, "retry")
        self.assertTrue(LocKitImportDraft.objects.filter(token=draft.token).exists())
        self.assertFalse(self.cs.unit_set.filter(source="Key").exists())

    def test_update_preview_promises_descriptions_carry_over(self) -> None:
        draft = self._stage(
            "en,cs,note\nEnglish,Czech,Note\nStaff,Hul,A wooden staff.\n"
        )

        response = self.client.get(
            reverse("loc-kit-glossary-preview", kwargs={"token": draft.token})
        )

        self.assertContains(response, "Add new terms")
        self.assertContains(response, "will not change")
        self.assertContains(response, "Descriptions from this table will be added")
        self.assertNotContains(response, "Create glossary component")


class LocKitGlossaryAppendTerminologySyncTest(ViewTestCase):
    """Appended terms stay usable in later components through terminology sync."""

    CREATE_GLOSSARIES: bool = True

    def setUp(self) -> None:
        super().setUp()
        self.user.is_superuser = True
        self.user.save()
        self.factory = RequestFactory()
        self.glossary = self.project.glossaries[0]
        self.cs = self.glossary.translation_set.get(language__code="cs")

    def _request(self):
        request = self.factory.post("/")
        request.user = self.user
        return request

    def test_sync_mirrors_new_term_without_touching_targets(self) -> None:
        # An existing translated term must survive the sync untouched.
        context = _context_key("Characters", "Hero")
        self.glossary.source_translation.add_unit(
            None, context, "Hero", author=self.user
        )
        hero = self.cs.unit_set.get(context=context, source="Hero")
        hero.target = "Hrdina"
        hero.state = STATE_TRANSLATED
        hero.save(update_fields=["target", "state"], same_content=True)

        # The new term lands in English and Polish only.
        preview = _append_preview(
            {"values": {"en": "Blade", "pl": "Ostrze"}}, source_language="en"
        )
        loc_kit.append_glossary_terms(self._request(), self.glossary, preview)

        source_unit = self.glossary.source_translation.unit_set.get(source="Blade")
        self.assertIn("terminology", source_unit.extra_flags)

        # A later ordinary component gains a language the glossary lacks;
        # the sync must create the glossary language and mirror the term.
        # (The stock test component ships with new_lang="contact", which
        # blocks add_new_language on purpose; opt this component in.)
        self.component.new_base = "po/hello.pot"
        self.component.new_lang = "add"
        self.component.save(update_fields=["new_base", "new_lang"])
        self.component.add_new_language(Language.objects.get(code="fr"), None)

        sync_terminology(self.glossary.pk, self.glossary)

        # Structural pairs appear in every glossary language, empty ...
        self.assertEqual(self.cs.unit_set.get(source="Blade").target, "")
        de = self.glossary.translation_set.get(language__code="de")
        self.assertTrue(de.unit_set.filter(source="Blade").exists())
        fr = self.glossary.translation_set.get(language__code="fr")
        self.assertEqual(fr.unit_set.get(source="Blade").target, "")
        pl = self.glossary.translation_set.get(language__code="pl")
        self.assertEqual(pl.unit_set.get(source="Blade").target, "Ostrze")
        # ... and the existing translated target is unchanged.
        hero.refresh_from_db()
        self.assertEqual(hero.target, "Hrdina")

    def test_japanese_target_and_explanations_survive_sync(self) -> None:
        preview = _append_preview(
            {
                "values": {"en": "Mage", "ja": "Mahou"},
                "notes": {"source": "Casts spells.", "target": {"ja": "Spellcaster."}},
            },
            source_language="en",
        )
        loc_kit.append_glossary_terms(self._request(), self.glossary, preview)

        sync_terminology(self.glossary.pk, self.glossary)

        ja = self.glossary.translation_set.get(language__code="ja")
        ja_unit = ja.unit_set.get(source="Mage")
        self.assertEqual(ja_unit.target, "Mahou")
        self.assertEqual(ja_unit.explanation, "Spellcaster.")
        self.assertEqual(ja_unit.source_unit.explanation, "Casts spells.")


# --------------------------------------------------------------------------- #
# DB-only Explanation import for ordinary string components
# --------------------------------------------------------------------------- #


class KitExplanationApplyServiceTest(ViewTestCase):
    """String-kit explanation import changes only the source explanation."""

    CREATE_GLOSSARIES: bool = True

    def setUp(self) -> None:
        super().setUp()
        self.user.is_superuser = True
        self.user.save()
        self.component.source_translation.add_unit(
            None,
            "greeting",
            "Привет",
            author=self.user,
        )

    def test_matching_key_sets_source_explanation_without_rewriting_content(
        self,
    ) -> None:
        source_unit = self.component.source_translation.unit_set.get(context="greeting")
        source_unit.extra_flags = "read-only"
        source_unit.save(update_fields=["extra_flags"], same_content=True)
        source_unit.refresh_from_db()
        prior_state = source_unit.state
        prior_flags = source_unit.extra_flags

        result = loc_kit.apply_kit_explanations(
            user=self.user,
            component=self.component,
            units=(
                StringUnit(
                    key="greeting",
                    values={"en": "Привет"},
                    comments=(),
                    references=(),
                    row=2,
                    explanation="A friendly opening.",
                ),
            ),
            overwrite=False,
        )

        source_unit.refresh_from_db()
        self.assertEqual(result.set_count, 1)
        self.assertEqual(result.missing_key_count, 0)
        self.assertEqual(source_unit.explanation, "A friendly opening.")
        self.assertEqual(source_unit.source, "Привет")
        self.assertEqual(source_unit.state, prior_state)
        self.assertEqual(source_unit.extra_flags, prior_flags)

    def test_rerun_with_same_value_is_unchanged_and_creates_no_change(self) -> None:
        source_unit = self.component.source_translation.unit_set.get(context="greeting")
        unit = StringUnit(
            key="greeting",
            values={"en": "Привет"},
            comments=(),
            references=(),
            row=2,
            explanation="A friendly opening.",
        )
        loc_kit.apply_kit_explanations(
            user=self.user, component=self.component, units=(unit,), overwrite=False
        )
        change_count_after_first = source_unit.change_set.count()

        result = loc_kit.apply_kit_explanations(
            user=self.user, component=self.component, units=(unit,), overwrite=False
        )

        self.assertEqual(result.unchanged_count, 1)
        self.assertEqual(result.set_count, 0)
        self.assertEqual(source_unit.change_set.count(), change_count_after_first)

    def test_blank_and_missing_key_are_classified_separately(self) -> None:
        result = loc_kit.apply_kit_explanations(
            user=self.user,
            component=self.component,
            units=(
                StringUnit(
                    key="greeting",
                    values={},
                    comments=(),
                    references=(),
                    row=2,
                    explanation="   ",
                ),
                StringUnit(
                    key="no_such_key",
                    values={},
                    comments=(),
                    references=(),
                    row=3,
                    explanation="Orphan explanation.",
                ),
            ),
            overwrite=False,
        )

        self.assertEqual(result.blank_count, 1)
        self.assertEqual(result.missing_key_count, 1)
        self.assertEqual(result.set_count, 0)

    def test_nonempty_existing_value_needs_explicit_overwrite(self) -> None:
        source_unit = self.component.source_translation.unit_set.get(context="greeting")
        source_unit.update_explanation("Original meaning.", self.user)
        unit = StringUnit(
            key="greeting",
            values={"en": "Привет"},
            comments=(),
            references=(),
            row=2,
            explanation="Replacement meaning.",
        )

        skipped = loc_kit.apply_kit_explanations(
            user=self.user, component=self.component, units=(unit,), overwrite=False
        )
        source_unit.refresh_from_db()
        self.assertEqual(skipped.would_overwrite_count, 1)
        self.assertEqual(source_unit.explanation, "Original meaning.")

        applied = loc_kit.apply_kit_explanations(
            user=self.user, component=self.component, units=(unit,), overwrite=True
        )
        source_unit.refresh_from_db()
        self.assertEqual(applied.set_count, 1)
        self.assertEqual(source_unit.explanation, "Replacement meaning.")

    def test_note_equal_to_explanation_is_still_set_and_flagged(self) -> None:
        source_unit = self.component.source_translation.unit_set.get(context="greeting")
        source_unit.note = "Shared context."
        source_unit.save(update_fields=["note"], same_content=True)
        unit = StringUnit(
            key="greeting",
            values={"en": "Привет"},
            comments=(),
            references=(),
            row=2,
            explanation="Shared context.",
        )

        result = loc_kit.apply_kit_explanations(
            user=self.user, component=self.component, units=(unit,), overwrite=False
        )

        source_unit.refresh_from_db()
        self.assertEqual(result.already_in_note_count, 1)
        self.assertEqual(result.set_count, 1)
        self.assertEqual(source_unit.explanation, "Shared context.")

    def test_actor_without_source_edit_is_denied_and_mutates_nothing(self) -> None:
        source_unit = self.component.source_translation.unit_set.get(context="greeting")
        unit = StringUnit(
            key="greeting",
            values={},
            comments=(),
            references=(),
            row=2,
            explanation="Should never land.",
        )

        with self.assertRaises(PermissionDenied):
            loc_kit.apply_kit_explanations(
                user=self.anotheruser,
                component=self.component,
                units=(unit,),
                overwrite=False,
            )

        source_unit.refresh_from_db()
        self.assertEqual(source_unit.explanation, "")

    def test_glossary_component_is_rejected(self) -> None:
        glossary = self.project.glossaries[0]
        unit = StringUnit(
            key="greeting",
            values={},
            comments=(),
            references=(),
            row=2,
            explanation="Never applies.",
        )

        with self.assertRaises(ValidationError):
            loc_kit.apply_kit_explanations(
                user=self.user, component=glossary, units=(unit,), overwrite=False
            )

    def test_locked_component_is_rejected(self) -> None:
        self.component.locked = True
        self.component.save(update_fields=["locked"])
        unit = StringUnit(
            key="greeting",
            values={},
            comments=(),
            references=(),
            row=2,
            explanation="Never applies.",
        )

        with self.assertRaises(ValidationError):
            loc_kit.apply_kit_explanations(
                user=self.user, component=self.component, units=(unit,), overwrite=False
            )

    def test_classify_preview_matches_apply_without_mutating(self) -> None:
        source_unit = self.component.source_translation.unit_set.get(context="greeting")
        unit = StringUnit(
            key="greeting",
            values={"en": "Привет"},
            comments=(),
            references=(),
            row=2,
            explanation="Previewed only.",
        )

        preview = loc_kit.classify_kit_explanations(
            component=self.component, units=(unit,), overwrite=False
        )

        source_unit.refresh_from_db()
        self.assertEqual(preview.set_count, 1)
        self.assertEqual(source_unit.explanation, "")

        applied = loc_kit.apply_kit_explanations(
            user=self.user, component=self.component, units=(unit,), overwrite=False
        )
        self.assertEqual(applied.set_count, preview.set_count)

    def test_changed_source_skips_explanation_and_is_reported(self) -> None:
        source_unit = self.component.source_translation.unit_set.get(context="greeting")
        unit = StringUnit(
            key="greeting",
            values={"en": "A completely different greeting"},
            comments=(),
            references=(),
            row=2,
            explanation="Should be skipped.",
        )

        result = loc_kit.apply_kit_explanations(
            user=self.user, component=self.component, units=(unit,), overwrite=False
        )

        source_unit.refresh_from_db()
        self.assertEqual(result.source_changed_count, 1)
        self.assertEqual(result.set_count, 0)
        self.assertEqual(source_unit.explanation, "")

    def test_blank_source_cell_for_existing_key_still_counts_as_changed(self) -> None:
        """An explicit blank source cell disagrees exactly like a different one."""
        unit = StringUnit(
            key="greeting",
            values={"en": ""},
            comments=(),
            references=(),
            row=2,
            explanation="Should be skipped too.",
        )

        result = loc_kit.apply_kit_explanations(
            user=self.user, component=self.component, units=(unit,), overwrite=False
        )

        self.assertEqual(result.source_changed_count, 1)
        self.assertEqual(result.set_count, 0)

    def test_caller_without_source_evidence_is_never_gated(self) -> None:
        """
        ``values={}`` means "no opinion", not "blank" - unlike an explicit
        blank cell. ``Component.apply_loc_kit_explanations`` (the one-shot
        apply right after a kit-derived component's translations first
        load) only ever tracks key -> explanation and never populates
        ``values``; it must keep applying explanations to freshly created
        units whose real source it never even sees here.
        """
        unit = StringUnit(
            key="greeting",
            values={},
            comments=(),
            references=(),
            row=2,
            explanation="Still applies.",
        )

        result = loc_kit.apply_kit_explanations(
            user=self.user, component=self.component, units=(unit,), overwrite=False
        )

        self.assertEqual(result.source_changed_count, 0)
        self.assertEqual(result.set_count, 1)

    def test_source_changed_between_preview_and_apply_is_still_caught(self) -> None:
        """
        ``classify_kit_explanations`` (preview) must never be trusted stale:
        the real gate lives in ``apply_kit_explanations``'s own locked
        re-classification, so a source edited after preview is still caught.
        """
        unit = StringUnit(
            key="greeting",
            values={"en": "Привет"},
            comments=(),
            references=(),
            row=2,
            explanation="Should be skipped after the source changes.",
        )

        preview = loc_kit.classify_kit_explanations(
            component=self.component, units=(unit,), overwrite=False
        )
        self.assertEqual(preview.set_count, 1)
        self.assertEqual(preview.source_changed_count, 0)

        # Someone edits the source between preview and confirm.
        source_unit = self.component.source_translation.unit_set.get(context="greeting")
        source_unit.source = "Совершенно другой источник"
        source_unit.save(update_fields=["source"], same_content=True)

        applied = loc_kit.apply_kit_explanations(
            user=self.user, component=self.component, units=(unit,), overwrite=False
        )

        source_unit.refresh_from_db()
        self.assertEqual(applied.source_changed_count, 1)
        self.assertEqual(applied.set_count, 0)
        self.assertEqual(source_unit.explanation, "")


# --------------------------------------------------------------------------- #
# Updating an existing string component from a loc-kit table
# --------------------------------------------------------------------------- #


class LocKitStringsUpdateServiceTest(ViewTestCase):
    """apply_loc_kit_string_update adds new keys and sets Explanations."""

    CREATE_GLOSSARIES: bool = True

    def setUp(self) -> None:
        super().setUp()
        self.user.is_superuser = True
        self.user.save()
        self.component = self.create_json_mono(project=self.project, name="Strings")
        self.component.new_lang = "add"
        self.component.edit_template = True
        self.component.save(update_fields=["new_lang", "edit_template"])
        self.existing_unit = self.component.source_translation.add_unit(
            None, "welcome_message", "Welcome!", author=self.user
        )

    def _row(self, **overrides) -> StringUnit:
        base = {
            "key": "welcome_message",
            "values": {"en": "Welcome!"},
            "comments": (),
            "references": (),
            "row": 2,
            "explanation": "",
            "flags": "",
        }
        base.update(overrides)
        return StringUnit(**base)

    def test_new_key_added_to_existing_and_newly_created_language(self) -> None:
        new_row = self._row(
            key="new_key",
            values={"en": "Hello", "cs": "Ahoj", "de": "Hallo"},
        )

        result = loc_kit.apply_loc_kit_string_update(
            user=self.user,
            component=self.component,
            units=(new_row,),
            overwrite_explanations=False,
        )

        self.assertEqual(result.strings.added, 1)
        self.assertIn("de", result.strings.created_languages)
        cs_unit = self.component.translation_set.get(language__code="cs").unit_set.get(
            context="new_key"
        )
        de_unit = self.component.translation_set.get(language__code="de").unit_set.get(
            context="new_key"
        )
        self.assertEqual(cs_unit.target, "Ahoj")
        self.assertEqual(de_unit.target, "Hallo")

    def test_existing_key_source_target_and_flags_never_change(self) -> None:
        self.existing_unit.extra_flags = "max-length:10"
        self.existing_unit.save(update_fields=["extra_flags"], same_content=True)
        cs_translation = self.component.translation_set.get(language__code="cs")
        cs_before = cs_translation.unit_set.get(context="welcome_message").target

        table_row = self._row(
            values={"en": "DIFFERENT TEXT", "cs": "ZMENA"}, flags="read-only"
        )
        loc_kit.apply_loc_kit_string_update(
            user=self.user,
            component=self.component,
            units=(table_row,),
            overwrite_explanations=False,
        )

        self.existing_unit.refresh_from_db()
        self.assertEqual(self.existing_unit.source, "Welcome!")
        self.assertEqual(self.existing_unit.extra_flags, "max-length:10")
        cs_after = cs_translation.unit_set.get(context="welcome_message").target
        self.assertEqual(cs_after, cs_before)

    def test_changed_source_is_reported_separately_from_new_and_existing_counts(
        self,
    ) -> None:
        """
        ``find_changed_sources`` is the discrepancy table the plan requires:
        key/old/new, distinct from what ``append_translation_strings``
        itself decides (it never touches an existing key regardless).
        """
        matching_row = self._row()  # values={"en": "Welcome!"}, matches exactly
        changed_row = self._row(
            key="welcome_message", values={"en": "Welcome, traveler!"}
        )
        new_row = self._row(key="new_key", values={"en": "Hi"})

        changed = loc_kit.find_changed_sources(
            component=self.component, units=(matching_row, new_row)
        )
        self.assertEqual(changed, ())

        changed = loc_kit.find_changed_sources(
            component=self.component, units=(changed_row,)
        )
        self.assertEqual(len(changed), 1)
        self.assertEqual(changed[0].key, "welcome_message")
        self.assertEqual(changed[0].old_source, "Welcome!")
        self.assertEqual(changed[0].new_source, "Welcome, traveler!")

        # Reporting the discrepancy never mutates the stored source.
        self.existing_unit.refresh_from_db()
        self.assertEqual(self.existing_unit.source, "Welcome!")

    def test_existing_and_new_key_get_explanation(self) -> None:
        table_rows = (
            self._row(explanation="Shown at startup."),
            self._row(key="new_key", values={"en": "Hi"}, explanation="A greeting."),
        )

        result = loc_kit.apply_loc_kit_string_update(
            user=self.user,
            component=self.component,
            units=table_rows,
            overwrite_explanations=False,
        )

        self.existing_unit.refresh_from_db()
        new_source = self.component.source_translation.unit_set.get(context="new_key")
        self.assertEqual(self.existing_unit.explanation, "Shown at startup.")
        self.assertEqual(new_source.explanation, "A greeting.")
        self.assertEqual(result.explanations.set_count, 2)

    def test_overwrite_requires_explicit_flag_and_rerun_is_idempotent(self) -> None:
        self.existing_unit.update_explanation("Original.", self.user)
        table_row = self._row(explanation="Replacement.")

        skipped = loc_kit.apply_loc_kit_string_update(
            user=self.user,
            component=self.component,
            units=(table_row,),
            overwrite_explanations=False,
        )
        self.existing_unit.refresh_from_db()
        self.assertEqual(skipped.explanations.would_overwrite_count, 1)
        self.assertEqual(self.existing_unit.explanation, "Original.")

        applied = loc_kit.apply_loc_kit_string_update(
            user=self.user,
            component=self.component,
            units=(table_row,),
            overwrite_explanations=True,
        )
        self.existing_unit.refresh_from_db()
        self.assertEqual(applied.explanations.set_count, 1)
        self.assertEqual(self.existing_unit.explanation, "Replacement.")

        rerun = loc_kit.apply_loc_kit_string_update(
            user=self.user,
            component=self.component,
            units=(table_row,),
            overwrite_explanations=True,
        )
        self.assertEqual(rerun.explanations.unchanged_count, 1)
        self.assertEqual(rerun.explanations.set_count, 0)

    def test_overwrite_skips_an_explanation_edited_after_the_preview(self) -> None:
        """
        A confirmed overwrite never discards an edit made after the preview.

        The operator confirms "overwrite" against what the preview showed.
        If a reviewer rewrites that Explanation while the apply is still
        queued or mid-portion, applying the confirmed value would silently
        destroy the newer human edit, so such a row is skipped and counted.
        """
        self.existing_unit.update_explanation("Previewed.", self.user)
        baseline = {self.existing_unit.context: "Previewed."}
        table_row = self._row(explanation="From the table.")

        self.existing_unit.update_explanation("Reviewer's newer text.", self.user)

        result = loc_kit.apply_loc_kit_string_update(
            user=self.user,
            component=self.component,
            units=(table_row,),
            overwrite_explanations=True,
            explanation_baseline=baseline,
        )
        self.existing_unit.refresh_from_db()
        self.assertEqual(self.existing_unit.explanation, "Reviewer's newer text.")
        self.assertEqual(result.explanations.set_count, 0)
        self.assertEqual(result.explanations.baseline_changed_count, 1)

    def _grant_only(self, user, codenames: list[str], name: str) -> None:
        role = Role.objects.create(name=name)
        role.permissions.add(*Permission.objects.filter(codename__in=codenames))
        group = Group.objects.create(name=name, language_selection=SELECTION_ALL)
        group.roles.add(role)
        group.components.add(self.component)
        user.groups.add(group)
        user.clear_permissions_cache()

    def test_without_source_edit_strings_still_add_explanations_unavailable(
        self,
    ) -> None:
        editor = create_another_user(suffix="-add-only")
        self._grant_only(editor, ["upload.perform", "unit.add"], "Add-only")
        new_row = self._row(
            key="new_key", values={"en": "Hi"}, explanation="Should not apply."
        )

        result = loc_kit.apply_loc_kit_string_update(
            user=editor,
            component=self.component,
            units=(new_row,),
            overwrite_explanations=False,
        )

        self.assertEqual(result.strings.added, 1)
        self.assertEqual(result.explanations.unavailable_count, 1)
        new_source = self.component.source_translation.unit_set.get(context="new_key")
        self.assertEqual(new_source.explanation, "")

    def test_without_upload_permission_explanations_still_apply_strings_unavailable(
        self,
    ) -> None:
        editor = create_another_user(suffix="-explain-only")
        self._grant_only(editor, ["source.edit"], "Explain-only")
        rows = (
            self._row(explanation="Existing key note."),
            self._row(key="new_key", values={"en": "Hi"}),
        )

        result = loc_kit.apply_loc_kit_string_update(
            user=editor,
            component=self.component,
            units=rows,
            overwrite_explanations=False,
        )

        self.assertEqual(result.strings.added, 0)
        self.assertEqual(result.strings.existing, 1)
        self.existing_unit.refresh_from_db()
        self.assertEqual(self.existing_unit.explanation, "Existing key note.")
        self.assertFalse(
            self.component.source_translation.unit_set.filter(
                context="new_key"
            ).exists()
        )

    def test_missing_language_without_translation_add_is_unavailable(self) -> None:
        """
        A missing language is created only with ``translation.add``.

        This isolates that one branch by denying exactly that permission,
        leaving the real ``upload.perform`` / ``unit.add`` grant that
        authorizes the new key itself untouched.
        """
        editor = create_another_user(suffix="-no-add-lang")
        self._grant_only(editor, ["upload.perform", "unit.add"], "No-add-language")
        new_row = self._row(key="new_key", values={"en": "Hi", "de": "Hallo"})
        real_has_perm = editor.has_perm

        def deny_translation_add(perm, obj=None):
            if perm == "translation.add":
                return False
            return real_has_perm(perm, obj)

        with patch.object(editor, "has_perm", side_effect=deny_translation_add):
            result = loc_kit.apply_loc_kit_string_update(
                user=editor,
                component=self.component,
                units=(new_row,),
                overwrite_explanations=False,
            )

        self.assertEqual(result.strings.added, 1)
        self.assertIn("de", result.strings.unavailable_languages)
        self.assertFalse(
            self.component.translation_set.filter(language__code="de").exists()
        )

    def test_glossary_component_is_rejected(self) -> None:
        glossary = self.project.glossaries[0]
        with self.assertRaises(ValidationError):
            loc_kit.apply_loc_kit_string_update(
                user=self.user,
                component=glossary,
                units=(self._row(),),
                overwrite_explanations=False,
            )

    def test_bilingual_component_is_rejected(self) -> None:
        bilingual = self.create_po(project=self.project, name="Bilingual")
        with self.assertRaises(ValidationError):
            loc_kit.apply_loc_kit_string_update(
                user=self.user,
                component=bilingual,
                units=(self._row(),),
                overwrite_explanations=False,
            )

    def test_read_only_flag_lands_after_targets_on_a_new_key(self) -> None:
        """
        Flags apply only after every target is written.

        ``read-only`` on the source unit flips its targets into
        STATE_READONLY (``update_state``); writing targets after the flag
        would overwrite that state back to STATE_TRANSLATED, silently
        un-readonlying the row the table asked to freeze.
        """
        new_row = self._row(
            key="frozen_key",
            values={"en": "Locked", "cs": "Zamceno"},
            flags="read-only",
        )

        result = loc_kit.append_translation_strings(
            user=self.user, component=self.component, units=(new_row,)
        )

        self.assertEqual(result.added, 1)
        source_unit = self.component.source_translation.unit_set.get(
            context="frozen_key"
        )
        self.assertIn("read-only", source_unit.extra_flags)
        cs_unit = self.component.translation_set.get(language__code="cs").unit_set.get(
            context="frozen_key"
        )
        self.assertEqual(cs_unit.target, "Zamceno")
        self.assertEqual(cs_unit.state, STATE_READONLY)

    def test_blank_source_language_cell_keeps_english_value_under_read_only(
        self,
    ) -> None:
        """
        Blank ru + non-blank en + read-only keeps the English value frozen.

        The plan's fallback contract for a row with an empty source-language
        cell: the English value is stored, every language receives read-only,
        and the source-language unit stays untranslated rather than being
        fabricated from the key.
        """
        new_row = self._row(
            key="promo",
            values={"en": "Summer Sale", "ru": ""},
            flags="read-only",
        )

        result = loc_kit.append_translation_strings(
            user=self.user, component=self.component, units=(new_row,)
        )

        self.assertEqual(result.added, 1)
        en_unit = self.component.translation_set.get(language__code="en").unit_set.get(
            context="promo"
        )
        self.assertEqual(en_unit.target, "Summer Sale")
        source_unit = self.component.source_translation.unit_set.get(context="promo")
        self.assertEqual(source_unit.state, STATE_READONLY)
        self.assertIn("read-only", source_unit.extra_flags)

    def test_judge_stale_counter_counts_verdicts_the_apply_would_outdate(self) -> None:
        """
        The preview counts current-context verdicts an apply would stale.

        A verdict is current when its context_hash matches the unit's live
        context. Setting a fresh explanation changes that context, so the
        target units judged against the old explanation must be counted
        (the stale verdicts no longer match the live context).
        """
        # ruff: ignore[import-outside-top-level]
        from weblate.trans.models.judge import (
            JudgeVerdict,
            compute_context_hash,
            compute_target_hash,
            compute_target_storage_hash,
        )

        cs_unit = self.component.translation_set.get(language__code="cs").unit_set.get(
            context="welcome_message"
        )
        context_hash = compute_context_hash(
            source=cs_unit.source,
            note=self.existing_unit.note,
            explanation=self.existing_unit.explanation,
            glossary_terms=[],
        )
        JudgeVerdict.objects.create(
            unit=cs_unit,
            max_severity=JudgeVerdict.Severity.NONE,
            judge_model="vendor/model-a",
            seat=1,
            target_hash=compute_target_hash(cs_unit.get_target_plurals()),
            target_storage_hash=compute_target_storage_hash(cs_unit.target),
            context_hash=context_hash,
        )
        JudgeVerdict.objects.create(
            unit=cs_unit,
            max_severity=JudgeVerdict.Severity.NONE,
            judge_model="vendor/model-b",
            seat=2,
            target_hash=compute_target_hash(cs_unit.get_target_plurals()),
            target_storage_hash=compute_target_storage_hash(cs_unit.target),
            context_hash=context_hash,
        )
        row = self._row(explanation="Fresh note.")

        before = loc_kit.count_judge_stale_after_explanations(
            component=self.component, units=(row,), overwrite=True
        )
        self.assertEqual(before, 1)

        loc_kit.apply_kit_explanations(
            user=self.user,
            component=self.component,
            units=(row,),
            overwrite=True,
        )
        after = loc_kit.count_judge_stale_after_explanations(
            component=self.component, units=(row,), overwrite=True
        )
        self.assertEqual(after, 0)

    def test_new_key_add_is_locked_and_serialized(self) -> None:
        """
        ``append_translation_strings`` runs under ``component.locked_for_update``.

        A lock timeout - the same signal a genuinely concurrent confirm of the
        same or an overlapping draft would raise - must abort before any unit
        is written, exactly like the sibling ``apply_kit_explanations`` path.
        Before this lock was added, a lock timeout here was silently ignored
        and the new key was still written.
        """
        new_row = self._row(key="new_key", values={"en": "Hi"})

        with (
            patch.object(
                Component,
                "locked_for_update",
                side_effect=WeblateLockTimeoutError("locked", lock=None),
            ),
            self.assertRaises(WeblateLockTimeoutError),
        ):
            loc_kit.apply_loc_kit_string_update(
                user=self.user,
                component=self.component,
                units=(new_row,),
                overwrite_explanations=False,
            )

        self.assertFalse(
            self.component.source_translation.unit_set.filter(
                context="new_key"
            ).exists()
        )

    def test_synchronous_limit_rejects_new_key_before_writing(self) -> None:
        """
        The synchronous flow refuses more parsed mutation cells than its limit.

        The guard belongs in the service, not only in its HTTP caller: no
        source unit may exist after a direct or retried call is rejected.
        """
        new_row = self._row(key="new_key", values={"en": "Hello", "cs": "Ahoj"})

        with (
            patch.object(
                loc_kit, "LOC_KIT_STRING_UPDATE_MAX_MUTATIONS", 1, create=True
            ),
            self.assertRaisesRegex(ValidationError, "synchronous update limit"),
        ):
            loc_kit.apply_loc_kit_string_update(
                user=self.user,
                component=self.component,
                units=(new_row,),
                overwrite_explanations=False,
            )

        self.assertFalse(
            self.component.source_translation.unit_set.filter(
                context="new_key"
            ).exists()
        )


class LocKitStringsUpdatePendingCommitTest(ViewTestCase):
    """
    New units actually reach the PO file, with draft-owned finalizing.

    A raw ``Unit.save()`` on a target creates no ``PendingUnitChange``, so a
    translation written that way is silently never committed to the file.
    This locks in the fix: writes go through ``Unit.translate``, get
    flushed via ``store_update_changes``, and this draft's returned
    ``pending_change_ids`` commit through ``commit_pending_subset`` alone.
    """

    def setUp(self) -> None:
        super().setUp()
        self.user.is_superuser = True
        self.user.save()
        self.component = self.create_po_mono(project=self.project, name="Strings")

    def _row(self, **overrides) -> StringUnit:
        base = {
            "key": "frozen_key",
            "values": {"en": "Hello", "cs": "Ahoj"},
            "comments": ("Character: Sample",),
            "references": ("42",),
            "row": 2,
            "explanation": "LLM-only context.",
            "flags": "",
        }
        base.update(overrides)
        return StringUnit(**base)

    def test_target_and_source_metadata_survive_subset_commit_and_reparse(
        self,
    ) -> None:
        result = loc_kit.apply_loc_kit_string_update(
            user=self.user,
            component=self.component,
            units=(self._row(),),
            overwrite_explanations=False,
            pending_owner="draft-token-1",
        )
        self.assertEqual(result.strings.added, 1)
        self.assertEqual(result.explanations.set_count, 1)
        self.assertTrue(result.pending_change_ids)
        self.assertEqual(
            set(
                PendingUnitChange.objects.filter(
                    pk__in=result.pending_change_ids
                ).values_list("metadata__loc_kit_draft_id", flat=True)
            ),
            {"draft-token-1"},
        )
        pending_order = list(
            PendingUnitChange.objects.filter(pk__in=result.pending_change_ids)
            .select_related("unit__translation")
            .order_by("timestamp")
            .values_list("unit__translation__language__code", flat=True)
        )
        self.assertEqual(pending_order, ["en", "cs"])
        commits_before = self.component.repository.count_outgoing()

        committed = self.component.commit_pending_subset(
            "loc-kit test", self.user, set(result.pending_change_ids)
        )
        self.assertEqual(
            self.component.repository.count_outgoing() - commits_before,
            1,
            "finalizing must write source and target in one repository commit",
        )
        self.assertTrue(committed)
        self.assertFalse(
            PendingUnitChange.objects.filter(pk__in=result.pending_change_ids).exists()
        )

        en_translation = self.component.source_translation
        en_translation.drop_store_cache()
        en_content = Path(en_translation.get_filename()).read_text(encoding="utf-8")
        self.assertIn("Character: Sample", en_content)
        self.assertIn("#: 42", en_content)
        self.assertNotIn("LLM-only context.", en_content)

        cs_translation = self.component.translation_set.get(language__code="cs")
        cs_translation.drop_store_cache()
        cs_content = Path(cs_translation.get_filename()).read_text(encoding="utf-8")
        self.assertIn("Ahoj", cs_content)

        source_unit = en_translation.unit_set.get(context="frozen_key")
        self.assertEqual(source_unit.explanation, "LLM-only context.")

    def test_new_unit_extra_flags_reach_the_backing_po_file_on_reparse(self) -> None:
        """
        New loc-kit row flags reach the backing PO file, not the DB alone.

        The runtime pending writer snapshots them into the source pending
        change's metadata and carries them into the file, so a fresh parse
        of the committed file reports the same flag.
        """
        row = self._row(key="frozen_flag_key", flags="read-only")

        result = loc_kit.apply_loc_kit_string_update(
            user=self.user,
            component=self.component,
            units=(row,),
            overwrite_explanations=False,
            pending_owner="draft-token-flags",
        )
        self.assertEqual(result.strings.added, 1)
        source_pending = PendingUnitChange.objects.get(
            pk__in=result.pending_change_ids,
            unit__translation=self.component.source_translation,
        )
        self.assertEqual(source_pending.metadata.get("extra_flags"), "read-only")

        committed = self.component.commit_pending_subset(
            "loc-kit flags test", self.user, set(result.pending_change_ids)
        )
        self.assertTrue(committed)

        source_translation = self.component.source_translation
        source_translation.drop_store_cache()
        content = Path(source_translation.get_filename()).read_text(encoding="utf-8")
        self.assertIn('#, read-only\nmsgid "frozen_flag_key"', content)

        # A true reparse through Weblate's own PoMonoFormat, not only the
        # raw bytes: proves the file round-trips as a recognized flag, not
        # just a matching substring.
        reparsed_pounit, _add = source_translation.store.find_unit(
            "frozen_flag_key", "Hello"
        )
        self.assertIn("read-only", reparsed_pounit.flags)

        source_unit = source_translation.unit_set.get(context="frozen_flag_key")
        self.assertEqual(source_unit.extra_flags, "read-only")

    def test_existing_key_extra_flags_are_never_snapshotted_or_changed(self) -> None:
        """
        An existing key's flags stay exactly as they are.

        The snapshot only ever applies to a brand-new key's own pending
        change.
        """
        existing = self.component.source_translation.add_unit(
            None, "existing_flag_key", "Existing", author=self.user
        )
        existing.extra_flags = "max-length:10"
        existing.save(update_fields=["extra_flags"], same_content=True)

        row = self._row(
            key="existing_flag_key",
            values={"en": "Existing", "cs": "Existujici"},
            flags="read-only",
        )
        result = loc_kit.apply_loc_kit_string_update(
            user=self.user,
            component=self.component,
            units=(row,),
            overwrite_explanations=False,
            pending_owner="draft-token-existing",
        )
        self.assertEqual(result.strings.existing, 1)
        self.assertFalse(
            any(
                pending.metadata.get("extra_flags")
                for pending in PendingUnitChange.objects.filter(
                    unit__context="existing_flag_key"
                )
            )
        )
        existing.refresh_from_db()
        self.assertEqual(existing.extra_flags, "max-length:10")

    def test_retried_flush_does_not_duplicate_note_or_location(self) -> None:
        result = loc_kit.append_translation_strings(
            user=self.user,
            component=self.component,
            units=(self._row(),),
            pending_owner="draft-token-1",
        )
        source_translation = self.component.source_translation
        pending_changes = list(
            PendingUnitChange.objects.filter(pk__in=result.pending_change_ids).filter(
                unit__translation=source_translation
            )
        )
        self.assertEqual(len(pending_changes), 1)

        # First flush: writes the file to disk (inside update_units) without
        # yet running the git commit that would delete the pending row -
        # the exact window a worker crash between subset-commit and
        # terminal state leaves behind.
        store = source_translation.store
        store.ensure_index()
        source_translation.update_units(
            pending_changes, store, self.user.get_author_name()
        )
        source_translation.drop_store_cache()
        once = Path(source_translation.get_filename()).read_text(encoding="utf-8")
        self.assertEqual(once.count("Character: Sample"), 1)

        # Retry: fresh parse from the already-written disk file, same
        # still-pending PendingUnitChange row.
        retried_store = source_translation.store
        retried_store.ensure_index()
        source_translation.update_units(
            pending_changes, retried_store, self.user.get_author_name()
        )
        source_translation.drop_store_cache()
        twice = Path(source_translation.get_filename()).read_text(encoding="utf-8")
        self.assertEqual(twice.count("Character: Sample"), 1)
        self.assertEqual(twice.count("#: 42"), 1)

    def test_subset_commit_processes_source_first_even_with_reversed_timestamps(
        self,
    ) -> None:
        """
        A target's ``new_unit`` template lookup only succeeds once the
        source/template translation's own file write has happened. This
        must hold by explicit ordering, not incidentally because the
        source's pending row usually sorts first by creation timestamp.
        """
        result = loc_kit.append_translation_strings(
            user=self.user,
            component=self.component,
            units=(self._row(),),
            pending_owner="draft-token-2",
        )
        source_translation = self.component.source_translation
        cs_translation = self.component.translation_set.get(language__code="cs")
        source_change = PendingUnitChange.objects.get(
            pk__in=result.pending_change_ids, unit__translation=source_translation
        )
        target_change = PendingUnitChange.objects.get(
            pk__in=result.pending_change_ids, unit__translation=cs_translation
        )
        # Force the opposite of natural creation order.
        PendingUnitChange.objects.filter(pk=target_change.pk).update(
            timestamp=source_change.timestamp - timedelta(seconds=5)
        )

        committed = self.component.commit_pending_subset(
            "loc-kit test", self.user, set(result.pending_change_ids)
        )

        self.assertTrue(committed)
        cs_translation.drop_store_cache()
        cs_content = Path(cs_translation.get_filename()).read_text(encoding="utf-8")
        self.assertIn("Ahoj", cs_content)

    def test_no_diff_replay_clears_pending_rows_without_a_second_commit(
        self,
    ) -> None:
        """
        A no-diff replay clears pending rows without a second commit.

        A pending change describing content already on disk - exactly the
        shape left behind by a crash between a successful VCS commit and
        the DB deleting its pending rows - is recognized as an idempotent
        already-applied replay: the row is cleared and ``True`` is
        returned, but no second repository commit is issued. This must
        hold even when the retry happens minutes later: the replay reuses
        the original pending change's own stable ``timestamp`` for the PO
        header, so a moved wall clock never manufactures a phantom
        revision-date diff.
        """
        row = self._row()
        result = loc_kit.append_translation_strings(
            user=self.user,
            component=self.component,
            units=(row,),
            pending_owner="draft-first",
        )
        # ``commit_pending_subset`` applies ONE header timestamp - the max
        # across every pending change in the batch, source and target
        # alike - to every translation it writes. The replay below must
        # reproduce that exact max, not just the source row's own value.
        original_timestamp = max(
            PendingUnitChange.objects.filter(
                pk__in=result.pending_change_ids
            ).values_list("timestamp", flat=True)
        )
        committed = self.component.commit_pending_subset(
            "first commit", self.user, set(result.pending_change_ids)
        )
        self.assertTrue(committed)
        commits_after_first = self.component.repository.count_outgoing()

        source_unit = self.component.source_translation.unit_set.get(context=row.key)
        # The row a real crash leaves behind is the SAME row, never
        # touched: its ``timestamp`` is whatever the first attempt already
        # used, not a freshly minted one.
        replay_change = PendingUnitChange.objects.create(
            unit=source_unit,
            author=self.user,
            target=source_unit.target,
            explanation=source_unit.explanation,
            source_unit_explanation=source_unit.explanation,
            state=source_unit.state,
            add_unit=False,
            metadata={"loc_kit_draft_id": "draft-first"},
            timestamp=original_timestamp,
        )

        # Advance the wall clock well past the first commit: the replay's
        # no-diff outcome must not depend on when the retry actually runs.
        moved_clock = timezone.now() + timedelta(minutes=10)
        with patch(
            "weblate.trans.models.translation.timezone.now", return_value=moved_clock
        ):
            recovered = self.component.commit_pending_subset(
                "replay commit", self.user, {replay_change.pk}
            )

        self.assertTrue(recovered)
        self.assertEqual(
            self.component.repository.count_outgoing(),
            commits_after_first,
            "a no-diff replay must not create a second commit",
        )
        self.assertFalse(PendingUnitChange.objects.filter(pk=replay_change.pk).exists())


class LocKitStringsUpdateViewTest(ViewTestCase):
    """The start -> preview -> confirm HTTP flow for an existing component."""

    CREATE_GLOSSARIES: bool = True

    def setUp(self) -> None:
        super().setUp()
        self.user.is_superuser = True
        self.user.save()
        self.component = self.create_json_mono(project=self.project, name="Strings")
        self.component.new_lang = "add"
        self.component.edit_template = True
        self.component.save(update_fields=["new_lang", "edit_template"])
        self.component.source_translation.add_unit(
            None, "welcome_message", "Welcome!", author=self.user
        )

    def _upload(self, body: str):
        return SimpleUploadedFile("update.csv", body.encode(), content_type="text/csv")

    def test_full_flow_adds_string_and_sets_explanation(self) -> None:
        kit = (
            "key,en,cs,Explanation\n"
            "welcome_message,Welcome!,Vitejte,Shown at startup.\n"
            "goodbye_message,Goodbye!,Sbohem,Shown at exit.\n"
        )
        start = self.client.post(
            reverse(
                "loc-kit-strings-update", kwargs={"path": self.component.get_url_path()}
            ),
            {"table": self._upload(kit)},
        )
        self.assertEqual(start.status_code, 302)
        preview_url = start["Location"]
        draft = LocKitImportDraft.objects.get()
        prepare_loc_kit_string_update.apply(
            kwargs={"draft_id": draft.pk},
            task_id=str(draft.prepare_task_id),
        )
        draft.refresh_from_db()
        self.assertEqual(draft.state, LocKitImportDraft.State.PREVIEW_READY)

        preview = self.client.get(preview_url)
        self.assertContains(preview, "1")  # new_count rendered somewhere on the page

        confirm = self.client.post(preview_url, {"action": "confirm"}, follow=True)
        self.assertEqual(confirm.status_code, 200)
        draft.refresh_from_db()
        self.assertEqual(draft.state, LocKitImportDraft.State.APPLYING)
        apply_loc_kit_string_update_draft.apply(
            kwargs={"draft_id": draft.pk},
            task_id=str(draft.apply_task_id),
        )
        draft.refresh_from_db()
        self.assertEqual(draft.state, LocKitImportDraft.State.COMPLETED)

        source_translation = self.component.source_translation
        welcome = source_translation.unit_set.get(context="welcome_message")
        goodbye = source_translation.unit_set.get(context="goodbye_message")
        self.assertEqual(welcome.explanation, "Shown at startup.")
        self.assertEqual(goodbye.explanation, "Shown at exit.")
        cs_goodbye = self.component.translation_set.get(
            language__code="cs"
        ).unit_set.get(context="goodbye_message")
        self.assertEqual(cs_goodbye.target, "Sbohem")
        self.assertTrue(LocKitImportDraft.objects.filter(pk=draft.pk).exists())

    def test_preview_reports_and_confirm_preserves_a_changed_source(self) -> None:
        """
        An existing key whose source the table disagrees with is reported
        in its own preview section, its Explanation is skipped, and confirm
        never touches the stored source - the plan's discrepancy table.
        """
        kit = (
            "key,en,cs,Explanation\n"
            "welcome_message,Welcome back!,Vitejte,Shown at startup.\n"
        )
        start = self.client.post(
            reverse(
                "loc-kit-strings-update", kwargs={"path": self.component.get_url_path()}
            ),
            {"table": self._upload(kit)},
        )
        preview_url = start["Location"]
        draft = LocKitImportDraft.objects.get()
        prepare_loc_kit_string_update.apply(
            kwargs={"draft_id": draft.pk},
            task_id=str(draft.prepare_task_id),
        )
        draft.refresh_from_db()
        self.assertEqual(draft.state, LocKitImportDraft.State.PREVIEW_READY)
        summary = json.loads(draft.preview_json)
        self.assertEqual(summary["changed_source_count"], 1)
        self.assertEqual(summary["changed_sources"][0]["key"], "welcome_message")
        self.assertEqual(summary["changed_sources"][0]["old_source"], "Welcome!")
        self.assertEqual(summary["changed_sources"][0]["new_source"], "Welcome back!")
        self.assertEqual(summary["explanations"]["source_changed_count"], 1)
        self.assertEqual(summary["explanations"]["set_count"], 0)

        preview = self.client.get(preview_url)
        self.assertContains(preview, "Welcome back!")
        self.assertContains(
            preview, "Existing keys whose source the table disagrees with"
        )

        confirm = self.client.post(preview_url, {"action": "confirm"}, follow=True)
        self.assertEqual(confirm.status_code, 200)
        draft.refresh_from_db()
        apply_loc_kit_string_update_draft.apply(
            kwargs={"draft_id": draft.pk},
            task_id=str(draft.apply_task_id),
        )
        draft.refresh_from_db()
        self.assertEqual(draft.state, LocKitImportDraft.State.COMPLETED)
        self.assertEqual(draft.progress["explanations_source_changed"], 1)

        welcome = self.component.source_translation.unit_set.get(
            context="welcome_message"
        )
        self.assertEqual(welcome.source, "Welcome!")
        self.assertEqual(welcome.explanation, "")

    @override_settings(CELERY_TASK_ALWAYS_EAGER=False)
    def test_confirm_reserves_background_application_without_unit_writes(self) -> None:
        start = self.client.post(
            reverse(
                "loc-kit-strings-update",
                kwargs={"path": self.component.get_url_path()},
            ),
            {"table": self._upload("key,en\nnew_key,Hello\n")},
        )
        preview_url = start["Location"]
        draft = LocKitImportDraft.objects.get()
        prepare_loc_kit_string_update.apply(
            kwargs={"draft_id": draft.pk},
            task_id=str(draft.prepare_task_id),
        )
        before = self.component.source_translation.unit_set.count()

        with (
            patch(
                "weblate.trans.tasks.apply_loc_kit_string_update_draft.apply_async"
            ) as dispatch,
            self.captureOnCommitCallbacks(execute=True),
        ):
            response = self.client.post(preview_url, {"action": "confirm"})

        self.assertEqual(response.status_code, 302)
        draft.refresh_from_db()
        self.assertEqual(draft.state, LocKitImportDraft.State.APPLYING)
        self.assertIsNotNone(draft.apply_task_id)
        dispatch.assert_called_once_with(
            kwargs={"draft_id": draft.pk},
            task_id=str(draft.apply_task_id),
            priority=0,
        )
        self.assertEqual(self.component.source_translation.unit_set.count(), before)

    @override_settings(CELERY_TASK_ALWAYS_EAGER=False)
    def test_unpublished_intent_is_reclaimed_by_the_drain(self) -> None:
        """
        An unpublished intent is reclaimed by the periodic drain.

        A confirm whose ``on_commit`` callback never ran (process loss
        before the broker publication) still leaves a durable dispatch
        intent; the drain publishes exactly the reserved UUID, and no unit
        is written before a worker executes.
        """
        start = self.client.post(
            reverse(
                "loc-kit-strings-update",
                kwargs={"path": self.component.get_url_path()},
            ),
            {"table": self._upload("key,en\nnew_key,Hello\n")},
        )
        preview_url = start["Location"]
        draft = LocKitImportDraft.objects.get()
        prepare_loc_kit_string_update.apply(
            kwargs={"draft_id": draft.pk},
            task_id=str(draft.prepare_task_id),
        )
        before = self.component.source_translation.unit_set.count()

        with (
            patch(
                "weblate.trans.tasks.apply_loc_kit_string_update_draft.apply_async"
            ) as dispatch,
            # The callback is captured and dropped: the process died between
            # the reserving commit and the fast-path publication.
            self.captureOnCommitCallbacks(execute=False),
        ):
            response = self.client.post(preview_url, {"action": "confirm"})

        self.assertEqual(response.status_code, 302)
        draft.refresh_from_db()
        self.assertEqual(draft.state, LocKitImportDraft.State.APPLYING)
        reserved_task_id = draft.apply_task_id
        self.assertEqual(draft.dispatch_task_id, reserved_task_id)
        self.assertEqual(draft.dispatch_phase, LocKitImportDraft.DispatchPhase.APPLY)
        self.assertIsNone(draft.dispatch_published_at)
        dispatch.assert_not_called()
        self.assertEqual(self.component.source_translation.unit_set.count(), before)

        with patch(
            "weblate.trans.tasks.apply_loc_kit_string_update_draft.apply_async"
        ) as reclaimed:
            drain_loc_kit_dispatches()

        reclaimed.assert_called_once_with(
            kwargs={"draft_id": draft.pk},
            task_id=str(reserved_task_id),
            priority=0,
        )
        draft.refresh_from_db()
        self.assertIsNotNone(draft.dispatch_published_at)
        self.assertEqual(self.component.source_translation.unit_set.count(), before)

    @override_settings(CELERY_TASK_ALWAYS_EAGER=False)
    def test_drain_republishes_same_uuid_after_broker_success_before_record(
        self,
    ) -> None:
        """
        The drain republishes the same UUID after a lost publish record.

        Broker accepted the message but the process died before
        ``dispatch_published_at`` was recorded: the drain republishes the
        SAME UUID (a duplicate delivery the portion contract makes safe),
        never a fresh generation.
        """
        start = self.client.post(
            reverse(
                "loc-kit-strings-update",
                kwargs={"path": self.component.get_url_path()},
            ),
            {"table": self._upload("key,en\nnew_key,Hello\n")},
        )
        preview_url = start["Location"]
        draft = LocKitImportDraft.objects.get()
        prepare_loc_kit_string_update.apply(
            kwargs={"draft_id": draft.pk},
            task_id=str(draft.prepare_task_id),
        )
        with self.captureOnCommitCallbacks(execute=False):
            self.client.post(preview_url, {"action": "confirm"})
        draft.refresh_from_db()
        reserved_task_id = draft.apply_task_id
        self.assertIsNone(draft.dispatch_published_at)

        with (
            patch(
                "weblate.trans.tasks.apply_loc_kit_string_update_draft.apply_async"
            ) as dispatch,
            # The DB mark dies after the broker call returned.
            patch(
                "weblate.trans.tasks._record_loc_kit_dispatch_published",
                side_effect=KeyboardInterrupt,
            ),
            self.assertRaises(KeyboardInterrupt),
        ):
            _publish_loc_kit_dispatch(draft_id=draft.pk)

        dispatch.assert_called_once_with(
            kwargs={"draft_id": draft.pk},
            task_id=str(reserved_task_id),
            priority=0,
        )
        draft.refresh_from_db()
        self.assertIsNone(draft.dispatch_published_at)

        with patch(
            "weblate.trans.tasks.apply_loc_kit_string_update_draft.apply_async"
        ) as republished:
            drain_loc_kit_dispatches()

        republished.assert_called_once_with(
            kwargs={"draft_id": draft.pk},
            task_id=str(reserved_task_id),
            priority=0,
        )
        draft.refresh_from_db()
        self.assertIsNotNone(draft.dispatch_published_at)

    @override_settings(CELERY_TASK_ALWAYS_EAGER=False)
    def test_dispatch_broker_failure_exhausts_to_retryable_failed(self) -> None:
        """
        Persistent broker failures are bounded.

        The cap flips the draft to a retryable FAILED with a sanitized
        message, never touching the payload or the cursor.
        """
        start = self.client.post(
            reverse(
                "loc-kit-strings-update",
                kwargs={"path": self.component.get_url_path()},
            ),
            {"table": self._upload("key,en\nnew_key,Hello\n")},
        )
        preview_url = start["Location"]
        draft = LocKitImportDraft.objects.get()
        prepare_loc_kit_string_update.apply(
            kwargs={"draft_id": draft.pk},
            task_id=str(draft.prepare_task_id),
        )
        with (
            patch(
                "weblate.trans.tasks.apply_loc_kit_string_update_draft.apply_async",
                side_effect=RuntimeError("broker credentials leaked here"),
            ),
            self.captureOnCommitCallbacks(execute=True),
        ):
            self.client.post(preview_url, {"action": "confirm"})

        draft.refresh_from_db()
        self.assertEqual(draft.state, LocKitImportDraft.State.APPLYING)
        self.assertEqual(draft.dispatch_attempts, 1)
        self.assertIsNone(draft.dispatch_published_at)

        with patch(
            "weblate.trans.tasks.apply_loc_kit_string_update_draft.apply_async",
            side_effect=RuntimeError("broker credentials leaked here"),
        ):
            for _ in range(LOC_KIT_DISPATCH_MAX_ATTEMPTS - 1):
                _publish_loc_kit_dispatch(draft_id=draft.pk)

        draft.refresh_from_db()
        self.assertEqual(draft.state, LocKitImportDraft.State.FAILED)
        self.assertEqual(draft.error_code, "dispatch-failed")
        self.assertEqual(draft.retry_phase, "apply")
        self.assertEqual(draft.dispatch_attempts, LOC_KIT_DISPATCH_MAX_ATTEMPTS)
        # Broker exception text must never reach the owner-visible fields.
        self.assertNotIn("credentials", draft.error_details)
        self.assertNotIn("credentials", draft.dispatch_error)
        # Cursor and private payload are untouched by dispatch bookkeeping.
        self.assertEqual(draft.next_row, 0)
        self.assertTrue(draft.prepared_payload.name)

    def test_cancel_deletes_the_draft_without_changing_anything(self) -> None:
        kit = "key,en\nnew_key,Hello\n"
        start = self.client.post(
            reverse(
                "loc-kit-strings-update", kwargs={"path": self.component.get_url_path()}
            ),
            {"table": self._upload(kit)},
        )
        preview_url = start["Location"]

        self.client.post(preview_url, {"action": "cancel"})

        self.assertFalse(LocKitImportDraft.objects.exists())
        self.assertFalse(
            self.component.source_translation.unit_set.filter(
                context="new_key"
            ).exists()
        )

    def test_multi_sheet_workbook_is_rejected(self) -> None:
        workbook = Workbook()
        workbook.active.title = "One"
        workbook.active.append(["key", "en"])
        workbook.active.append(["a", "A"])
        workbook.create_sheet("Two").append(["key", "en"])
        buffer = io.BytesIO()
        workbook.save(buffer)

        response = self.client.post(
            reverse(
                "loc-kit-strings-update", kwargs={"path": self.component.get_url_path()}
            ),
            {
                "table": SimpleUploadedFile(
                    "kit.xlsx",
                    buffer.getvalue(),
                    content_type=(
                        "application/vnd.openxmlformats-officedocument."
                        "spreadsheetml.sheet"
                    ),
                )
            },
        )

        self.assertEqual(response.status_code, 302)
        draft = LocKitImportDraft.objects.get()
        prepare_loc_kit_string_update.apply(
            kwargs={"draft_id": draft.pk},
            task_id=str(draft.prepare_task_id),
        )
        draft.refresh_from_db()
        self.assertEqual(draft.state, LocKitImportDraft.State.FAILED)
        self.assertIn("exactly one worksheet", draft.error_details)

    def test_glossary_component_has_no_start_route(self) -> None:
        glossary = self.project.glossaries[0]
        response = self.client.post(
            reverse("loc-kit-strings-update", kwargs={"path": glossary.get_url_path()}),
            {"table": self._upload("key,en\na,A\n")},
        )
        self.assertEqual(response.status_code, 404)

    @override_settings(
        TRANSLATION_UPLOAD_MAX_SIZE=200, COMPONENT_ZIP_UPLOAD_MAX_SIZE=1_000_000
    )
    def test_oversized_table_under_conflicting_limits_is_rejected_synchronously(
        self,
    ) -> None:
        """
        A table over ``TRANSLATION_UPLOAD_MAX_SIZE`` must fail synchronously.

        This must hold even when the much larger
        ``COMPONENT_ZIP_UPLOAD_MAX_SIZE`` would have let it through: no
        draft row, no storage object and no background dispatch for a table
        that can only ever fail asynchronously.
        """
        kit = "key,en\n" + "".join(
            f"padding_key_{number},{'x' * 20}\n" for number in range(20)
        )
        self.assertGreater(len(kit.encode()), 200)
        self.assertLess(len(kit.encode()), 1_000_000)

        with patch(
            "weblate.trans.tasks.prepare_loc_kit_string_update.apply_async"
        ) as dispatch:
            response = self.client.post(
                reverse(
                    "loc-kit-strings-update",
                    kwargs={"path": self.component.get_url_path()},
                ),
                {"table": self._upload(kit)},
            )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "too big")
        self.assertFalse(LocKitImportDraft.objects.exists())
        dispatch.assert_not_called()

    def test_apply_lock_timeout_fails_retryably_and_shows_retry(self) -> None:
        kit = "key,en\nnew_key,Hello\n"
        start = self.client.post(
            reverse(
                "loc-kit-strings-update", kwargs={"path": self.component.get_url_path()}
            ),
            {"table": self._upload(kit)},
        )
        preview_url = start["Location"]
        draft = LocKitImportDraft.objects.get()
        prepare_loc_kit_string_update.apply(
            kwargs={"draft_id": draft.pk},
            task_id=str(draft.prepare_task_id),
        )
        self.client.post(preview_url, {"action": "confirm"})
        draft.refresh_from_db()
        self.assertEqual(draft.state, LocKitImportDraft.State.APPLYING)

        with patch.object(
            Component,
            "locked_for_update",
            side_effect=WeblateLockTimeoutError("locked", lock=None),
        ):
            apply_loc_kit_string_update_draft.apply(
                kwargs={"draft_id": draft.pk},
                task_id=str(draft.apply_task_id),
                retries=8,
            )

        draft.refresh_from_db()
        self.assertEqual(draft.state, LocKitImportDraft.State.FAILED)
        self.assertEqual(draft.retry_phase, "apply")
        self.assertFalse(
            self.component.source_translation.unit_set.filter(
                context="new_key"
            ).exists()
        )

        response = self.client.get(preview_url)
        self.assertContains(response, "retry")

    @override_settings(CELERY_TASK_ALWAYS_EAGER=False)
    def test_failed_apply_retry_reserves_a_new_task_and_completes(self) -> None:
        kit = "key,en\nnew_key,Hello\n"
        start = self.client.post(
            reverse(
                "loc-kit-strings-update", kwargs={"path": self.component.get_url_path()}
            ),
            {"table": self._upload(kit)},
        )
        preview_url = start["Location"]
        draft = LocKitImportDraft.objects.get()
        prepare_loc_kit_string_update.apply(
            kwargs={"draft_id": draft.pk},
            task_id=str(draft.prepare_task_id),
        )
        self.client.post(preview_url, {"action": "confirm"})
        draft.refresh_from_db()
        with patch.object(
            Component,
            "locked_for_update",
            side_effect=WeblateLockTimeoutError("locked", lock=None),
        ):
            apply_loc_kit_string_update_draft.apply(
                kwargs={"draft_id": draft.pk},
                task_id=str(draft.apply_task_id),
                retries=8,
            )
        draft.refresh_from_db()
        old_task_id = draft.apply_task_id

        with (
            patch(
                "weblate.trans.tasks.apply_loc_kit_string_update_draft.apply_async"
            ) as dispatch,
            self.captureOnCommitCallbacks(execute=True),
        ):
            response = self.client.post(preview_url, {"action": "retry"})

        self.assertEqual(response.status_code, 302)
        draft.refresh_from_db()
        self.assertEqual(draft.state, LocKitImportDraft.State.APPLYING)
        self.assertNotEqual(draft.apply_task_id, old_task_id)
        dispatch.assert_called_once_with(
            kwargs={"draft_id": draft.pk},
            task_id=str(draft.apply_task_id),
            priority=0,
        )

        apply_loc_kit_string_update_draft.apply(
            kwargs={"draft_id": draft.pk},
            task_id=str(draft.apply_task_id),
        )
        draft.refresh_from_db()
        self.assertEqual(draft.state, LocKitImportDraft.State.COMPLETED)
        self.assertTrue(
            self.component.source_translation.unit_set.filter(
                context="new_key"
            ).exists()
        )

    @override_settings(CELERY_TASK_ALWAYS_EAGER=False, CELERY_VISIBILITY_TIMEOUT=4)
    def test_apply_delivery_budget_reserves_one_fenced_continuation(self) -> None:
        """
        A delivery stops after a committed portion, hands ownership to one
        fresh UUID, and does not finish the following portion itself.
        """
        rows = "key,en\n" + "".join(f"chain{i},Value {i}\n" for i in range(26))
        start = self.client.post(
            reverse(
                "loc-kit-strings-update",
                kwargs={"path": self.component.get_url_path()},
            ),
            {"table": self._upload(rows)},
        )
        draft = LocKitImportDraft.objects.get()
        prepare_loc_kit_string_update.apply(
            kwargs={"draft_id": draft.pk}, task_id=str(draft.prepare_task_id)
        )
        self.client.post(start["Location"], {"action": "confirm"})
        draft.refresh_from_db()
        original_task_id = draft.apply_task_id

        with (
            patch(
                "weblate.trans.tasks.time.monotonic",
                side_effect=[0.0, 3.0],
            ),
            patch(
                "weblate.trans.tasks.apply_loc_kit_string_update_draft.apply_async"
            ) as dispatch,
            self.captureOnCommitCallbacks(execute=True),
        ):
            apply_loc_kit_string_update_draft.apply(
                kwargs={"draft_id": draft.pk}, task_id=str(original_task_id)
            )

        draft.refresh_from_db()
        self.assertEqual(draft.state, LocKitImportDraft.State.APPLYING)
        self.assertEqual(draft.next_row, 25)
        self.assertNotEqual(draft.apply_task_id, original_task_id)
        self.assertEqual(
            self.component.source_translation.unit_set.filter(
                context__startswith="chain"
            ).count(),
            25,
        )
        dispatch.assert_called_once_with(
            kwargs={"draft_id": draft.pk},
            task_id=str(draft.apply_task_id),
            priority=0,
        )

    def test_cumulative_apply_cap_commits_completed_portions_before_stopping(
        self,
    ) -> None:
        """
        The cap never leaves this draft's already-written units pending only
        in the database: its committed portion is flushed before reporting
        the partial, terminal result.
        """
        rows = "key,en\n" + "".join(f"cap{i},Value {i}\n" for i in range(26))
        start = self.client.post(
            reverse(
                "loc-kit-strings-update",
                kwargs={"path": self.component.get_url_path()},
            ),
            {"table": self._upload(rows)},
        )
        draft = LocKitImportDraft.objects.get()
        prepare_loc_kit_string_update.apply(
            kwargs={"draft_id": draft.pk}, task_id=str(draft.prepare_task_id)
        )
        self.client.post(start["Location"], {"action": "confirm"})
        draft.refresh_from_db()

        real_apply = loc_kit.apply_loc_kit_portion
        calls = {"count": 0}

        def interrupt_before_second_portion(*args, **kwargs):
            calls["count"] += 1
            if calls["count"] == 2:
                raise KeyboardInterrupt
            return real_apply(*args, **kwargs)

        with (
            patch.object(
                loc_kit, "LOC_KIT_STRING_UPDATE_PORTION_SIZE", 25, create=True
            ),
            patch.object(
                loc_kit,
                "apply_loc_kit_portion",
                side_effect=interrupt_before_second_portion,
            ),
            self.assertRaises(KeyboardInterrupt),
        ):
            apply_loc_kit_string_update_draft.apply(
                kwargs={"draft_id": draft.pk}, task_id=str(draft.apply_task_id)
            )

        draft.refresh_from_db()
        draft.apply_started_at = (
            timezone.now()
            - LOC_KIT_STRING_UPDATE_APPLY_TIME_LIMIT
            - timedelta(seconds=1)
        )
        draft.save(update_fields=["apply_started_at"])

        apply_loc_kit_string_update_draft.apply(
            kwargs={"draft_id": draft.pk}, task_id=str(draft.apply_task_id)
        )

        draft.refresh_from_db()
        self.assertEqual(draft.state, LocKitImportDraft.State.FAILED)
        self.assertEqual(draft.retry_phase, "")
        self.assertEqual(draft.progress["processed_rows"], 25)
        self.assertEqual(draft.progress["added"], 25)
        self.assertEqual(
            self.component.source_translation.unit_set.filter(
                context__startswith="cap"
            ).count(),
            25,
        )
        self.assertFalse(
            PendingUnitChange.objects.filter(
                metadata__loc_kit_draft_id=str(draft.pk)
            ).exists()
        )

    def test_apply_retries_a_transient_component_lock_timeout(self) -> None:
        """A transient lock timeout retries the same durable portion."""
        start = self.client.post(
            reverse(
                "loc-kit-strings-update",
                kwargs={"path": self.component.get_url_path()},
            ),
            {"table": self._upload("key,en\nretry_lock,Value\n")},
        )
        draft = LocKitImportDraft.objects.get()
        prepare_loc_kit_string_update.apply(
            kwargs={"draft_id": draft.pk}, task_id=str(draft.prepare_task_id)
        )
        self.client.post(start["Location"], {"action": "confirm"})
        draft.refresh_from_db()

        real_apply = loc_kit.apply_loc_kit_portion
        calls = {"count": 0}

        def fail_once(*args, **kwargs):
            calls["count"] += 1
            if calls["count"] == 1:
                msg = "locked"
                raise WeblateLockTimeoutError(msg, lock=None)
            return real_apply(*args, **kwargs)

        with patch.object(loc_kit, "apply_loc_kit_portion", side_effect=fail_once):
            apply_loc_kit_string_update_draft.apply(
                kwargs={"draft_id": draft.pk},
                task_id=str(draft.apply_task_id),
                throw=False,
            )
            apply_loc_kit_string_update_draft.apply(
                kwargs={"draft_id": draft.pk},
                task_id=str(draft.apply_task_id),
                retries=1,
            )

        draft.refresh_from_db()
        self.assertEqual(draft.state, LocKitImportDraft.State.COMPLETED)
        self.assertEqual(draft.progress["added"], 1)
        self.assertTrue(
            self.component.source_translation.unit_set.filter(
                context="retry_lock"
            ).exists()
        )

    @override_settings(CELERY_TASK_ALWAYS_EAGER=False)
    def test_stale_applying_draft_offers_a_fenced_retry(self) -> None:
        """A stalled APPLYING delivery can be replaced without a new upload."""
        start = self.client.post(
            reverse(
                "loc-kit-strings-update",
                kwargs={"path": self.component.get_url_path()},
            ),
            {"table": self._upload("key,en\nstalled,Value\n")},
        )
        preview_url = start["Location"]
        draft = LocKitImportDraft.objects.get()
        prepare_loc_kit_string_update.apply(
            kwargs={"draft_id": draft.pk}, task_id=str(draft.prepare_task_id)
        )
        self.client.post(preview_url, {"action": "confirm"})
        draft.refresh_from_db()
        original_task_id = draft.apply_task_id
        draft.last_activity_at = timezone.now() - timedelta(minutes=31)
        draft.save(update_fields=["last_activity_at"])

        response = self.client.get(preview_url)
        self.assertContains(response, 'name="action" value="retry"', html=False)

        with (
            patch(
                "weblate.trans.tasks.apply_loc_kit_string_update_draft.apply_async"
            ) as dispatch,
            self.captureOnCommitCallbacks(execute=True),
        ):
            response = self.client.post(preview_url, {"action": "retry"})

        self.assertEqual(response.status_code, 302)
        draft.refresh_from_db()
        replacement_task_id = draft.apply_task_id
        self.assertEqual(draft.state, LocKitImportDraft.State.APPLYING)
        self.assertNotEqual(replacement_task_id, original_task_id)
        dispatch.assert_called_once_with(
            kwargs={"draft_id": draft.pk},
            task_id=str(replacement_task_id),
            priority=0,
        )

        apply_loc_kit_string_update_draft.apply(
            kwargs={"draft_id": draft.pk}, task_id=str(original_task_id)
        )
        draft.refresh_from_db()
        self.assertEqual(draft.apply_task_id, replacement_task_id)
        self.assertEqual(draft.state, LocKitImportDraft.State.APPLYING)

    def test_completed_result_is_discoverable_from_the_component_page(self) -> None:
        """
        The plan's core UX requirement: closing the preview tab must not
        strand the result. A fresh request for the component page -
        without ever revisiting the preview URL - carries a link back to
        it, bound to the same owner and session the draft was created
        under (matching every other draft access check in this flow).
        """
        kit = "key,en\nnew_key,Hello\n"
        start = self.client.post(
            reverse(
                "loc-kit-strings-update", kwargs={"path": self.component.get_url_path()}
            ),
            {"table": self._upload(kit)},
        )
        preview_url = start["Location"]
        draft = LocKitImportDraft.objects.get()
        prepare_loc_kit_string_update.apply(
            kwargs={"draft_id": draft.pk}, task_id=str(draft.prepare_task_id)
        )
        self.client.post(preview_url, {"action": "confirm"})
        draft.refresh_from_db()
        apply_loc_kit_string_update_draft.apply(
            kwargs={"draft_id": draft.pk}, task_id=str(draft.apply_task_id)
        )
        draft.refresh_from_db()
        self.assertEqual(draft.state, LocKitImportDraft.State.COMPLETED)

        component_page = self.client.get(
            reverse("show", kwargs={"path": self.component.get_url_path()})
        )

        self.assertContains(component_page, preview_url)

    def test_completed_result_link_is_absent_for_a_different_session(self) -> None:
        """The discovery link never renders a token the session check would 404."""
        kit = "key,en\nnew_key,Hello\n"
        start = self.client.post(
            reverse(
                "loc-kit-strings-update", kwargs={"path": self.component.get_url_path()}
            ),
            {"table": self._upload(kit)},
        )
        preview_url = start["Location"]
        draft = LocKitImportDraft.objects.get()
        prepare_loc_kit_string_update.apply(
            kwargs={"draft_id": draft.pk}, task_id=str(draft.prepare_task_id)
        )
        self.client.post(preview_url, {"action": "confirm"})
        draft.refresh_from_db()
        apply_loc_kit_string_update_draft.apply(
            kwargs={"draft_id": draft.pk}, task_id=str(draft.apply_task_id)
        )

        self.client.session.flush()
        self.client.force_login(self.user)
        component_page = self.client.get(
            reverse("show", kwargs={"path": self.component.get_url_path()})
        )

        self.assertNotContains(component_page, preview_url)

    def test_apply_resumes_from_cursor_and_merges_totals_across_deliveries(
        self,
    ) -> None:
        """
        A redelivered task resumes from the durable cursor, and the
        COMPLETED summary reflects every portion applied across every
        delivery - not only the last one, which resets its local counters
        to zero on every fresh task invocation.
        """
        rows = "key,en\n" + "".join(f"k{i},V{i}\n" for i in range(23))
        start = self.client.post(
            reverse(
                "loc-kit-strings-update",
                kwargs={"path": self.component.get_url_path()},
            ),
            {"table": self._upload(rows)},
        )
        draft = LocKitImportDraft.objects.get()
        prepare_loc_kit_string_update.apply(
            kwargs={"draft_id": draft.pk}, task_id=str(draft.prepare_task_id)
        )
        self.client.post(start["Location"], {"action": "confirm"})
        draft.refresh_from_db()

        real_apply = loc_kit.apply_loc_kit_portion
        calls = {"n": 0}

        def crash_after_first_portion(*args, **kwargs):
            calls["n"] += 1
            if calls["n"] == 2:
                # Simulate a hard worker crash mid-table: no exception
                # handler in the task ever runs, unlike a normal Python
                # exception, so this must never leave a FAILED marking
                # behind - only the redelivery resumes it.
                raise KeyboardInterrupt
            return real_apply(*args, **kwargs)

        with (
            patch.object(
                loc_kit, "LOC_KIT_STRING_UPDATE_PORTION_SIZE", 10, create=True
            ),
            patch.object(
                loc_kit,
                "apply_loc_kit_portion",
                side_effect=crash_after_first_portion,
            ),
            self.assertRaises(KeyboardInterrupt),
        ):
            apply_loc_kit_string_update_draft.apply(
                kwargs={"draft_id": draft.pk}, task_id=str(draft.apply_task_id)
            )

        draft.refresh_from_db()
        self.assertEqual(draft.state, LocKitImportDraft.State.APPLYING)
        self.assertEqual(draft.next_row, 10)
        self.assertEqual(draft.progress.get("added"), 10)

        with patch.object(
            loc_kit, "LOC_KIT_STRING_UPDATE_PORTION_SIZE", 10, create=True
        ):
            apply_loc_kit_string_update_draft.apply(
                kwargs={"draft_id": draft.pk}, task_id=str(draft.apply_task_id)
            )

        draft.refresh_from_db()
        self.assertEqual(draft.state, LocKitImportDraft.State.COMPLETED)
        self.assertEqual(draft.progress["added"], 23)
        self.assertEqual(
            self.component.source_translation.unit_set.filter(
                context__startswith="k"
            ).count(),
            23,
        )

    @override_settings(CELERY_TASK_ALWAYS_EAGER=False)
    def test_retry_view_resumes_from_cursor_and_merges_totals(self) -> None:
        """
        The same cursor/totals durability as a raw redelivery, but through
        the actual failed-state retry button: a new ``apply_task_id`` must
        still resume from the durable ``next_row`` and add the prior
        delivery's counters to its own, not start counting from zero.
        """
        rows = "key,en\n" + "".join(f"r{i},V{i}\n" for i in range(23))
        start = self.client.post(
            reverse(
                "loc-kit-strings-update",
                kwargs={"path": self.component.get_url_path()},
            ),
            {"table": self._upload(rows)},
        )
        preview_url = start["Location"]
        draft = LocKitImportDraft.objects.get()
        prepare_loc_kit_string_update.apply(
            kwargs={"draft_id": draft.pk}, task_id=str(draft.prepare_task_id)
        )
        self.client.post(preview_url, {"action": "confirm"})
        draft.refresh_from_db()

        real_apply = loc_kit.apply_loc_kit_portion
        calls = {"n": 0}

        def fail_after_first_portion(*args, **kwargs):
            calls["n"] += 1
            if calls["n"] == 2:
                msg = "Locked components cannot be updated."
                raise ValidationError(msg)
            return real_apply(*args, **kwargs)

        with (
            patch.object(
                loc_kit, "LOC_KIT_STRING_UPDATE_PORTION_SIZE", 10, create=True
            ),
            patch.object(
                loc_kit,
                "apply_loc_kit_portion",
                side_effect=fail_after_first_portion,
            ),
        ):
            apply_loc_kit_string_update_draft.apply(
                kwargs={"draft_id": draft.pk}, task_id=str(draft.apply_task_id)
            )

        draft.refresh_from_db()
        self.assertEqual(draft.state, LocKitImportDraft.State.FAILED)
        self.assertEqual(draft.next_row, 10)
        self.assertEqual(draft.progress.get("added"), 10)
        old_task_id = draft.apply_task_id

        with (
            patch(
                "weblate.trans.tasks.apply_loc_kit_string_update_draft.apply_async"
            ) as dispatch,
            self.captureOnCommitCallbacks(execute=True),
        ):
            response = self.client.post(preview_url, {"action": "retry"})
        self.assertEqual(response.status_code, 302)
        draft.refresh_from_db()
        self.assertNotEqual(draft.apply_task_id, old_task_id)
        dispatch.assert_called_once_with(
            kwargs={"draft_id": draft.pk},
            task_id=str(draft.apply_task_id),
            priority=0,
        )
        # The retry view itself must not have touched the durable cursor.
        self.assertEqual(draft.next_row, 10)
        self.assertEqual(draft.progress.get("added"), 10)

        with patch.object(
            loc_kit, "LOC_KIT_STRING_UPDATE_PORTION_SIZE", 10, create=True
        ):
            apply_loc_kit_string_update_draft.apply(
                kwargs={"draft_id": draft.pk}, task_id=str(draft.apply_task_id)
            )

        draft.refresh_from_db()
        self.assertEqual(draft.state, LocKitImportDraft.State.COMPLETED)
        self.assertEqual(draft.progress["added"], 23)
        self.assertEqual(draft.progress["existing"], 0)
        self.assertEqual(
            self.component.source_translation.unit_set.filter(
                context__startswith="r"
            ).count(),
            23,
        )

    @override_settings(CELERY_TASK_ALWAYS_EAGER=False)
    def test_retry_view_resumes_with_real_portion_size_two_portions_done(
        self,
    ) -> None:
        """
        Same as ``test_retry_view_resumes_from_cursor_and_merges_totals``
        but with the real, unpatched ``LOC_KIT_STRING_UPDATE_PORTION_SIZE``
        and two successful portions before the failure - the exact shape
        observed in a live E2E run (75 rows, portions of 25, failure on
        the third portion after two succeed).
        """
        row_count = 75
        rows = "key,en\n" + "".join(f"p{i},V{i}\n" for i in range(row_count))
        start = self.client.post(
            reverse(
                "loc-kit-strings-update",
                kwargs={"path": self.component.get_url_path()},
            ),
            {"table": self._upload(rows)},
        )
        preview_url = start["Location"]
        draft = LocKitImportDraft.objects.get()
        prepare_loc_kit_string_update.apply(
            kwargs={"draft_id": draft.pk}, task_id=str(draft.prepare_task_id)
        )
        self.client.post(preview_url, {"action": "confirm"})
        draft.refresh_from_db()

        real_apply = loc_kit.apply_loc_kit_portion
        calls = {"n": 0}

        def fail_on_third_portion(*args, **kwargs):
            calls["n"] += 1
            if calls["n"] == 3:
                msg = "Locked components cannot be updated."
                raise ValidationError(msg)
            return real_apply(*args, **kwargs)

        with patch.object(
            loc_kit,
            "apply_loc_kit_portion",
            side_effect=fail_on_third_portion,
        ):
            apply_loc_kit_string_update_draft.apply(
                kwargs={"draft_id": draft.pk}, task_id=str(draft.apply_task_id)
            )

        draft.refresh_from_db()
        self.assertEqual(draft.state, LocKitImportDraft.State.FAILED)
        self.assertEqual(draft.next_row, 50)
        self.assertEqual(draft.progress.get("added"), 50)

        with (
            patch(
                "weblate.trans.tasks.apply_loc_kit_string_update_draft.apply_async"
            ) as dispatch,
            self.captureOnCommitCallbacks(execute=True),
        ):
            self.client.post(preview_url, {"action": "retry"})
        draft.refresh_from_db()
        dispatch.assert_called_once_with(
            kwargs={"draft_id": draft.pk},
            task_id=str(draft.apply_task_id),
            priority=0,
        )
        # Nothing about the retry dispatch itself may touch the cursor.
        self.assertEqual(draft.next_row, 50)
        self.assertEqual(draft.progress.get("added"), 50)

        apply_loc_kit_string_update_draft.apply(
            kwargs={"draft_id": draft.pk}, task_id=str(draft.apply_task_id)
        )

        draft.refresh_from_db()
        self.assertEqual(draft.state, LocKitImportDraft.State.COMPLETED)
        self.assertEqual(draft.progress["added"], row_count)
        self.assertEqual(draft.progress["existing"], 0)
        self.assertEqual(
            self.component.source_translation.unit_set.filter(
                context__startswith="p"
            ).count(),
            row_count,
        )

    @override_settings(CELERY_TASK_ALWAYS_EAGER=False)
    def test_explanation_failure_rolls_back_the_whole_portion(self) -> None:
        """
        An Explanation failure must roll back the whole atomic portion.

        Both mutations run inside one atomic portion (``apply_loc_kit_portion``).
        A failure in the Explanation step - after the append would previously
        have already committed - must roll back the append too: no new key,
        no owned pending row and no cursor advance survive without their
        paired outcome. A retry then applies the whole portion fresh,
        reported as ``added``, never as a duplicate-avoiding ``existing``.
        The completed totals add up to every row exactly once; no unit is
        created twice.
        """
        row_count = 75
        rows = "key,en\n" + "".join(f"m{i},V{i}\n" for i in range(row_count))
        start = self.client.post(
            reverse(
                "loc-kit-strings-update",
                kwargs={"path": self.component.get_url_path()},
            ),
            {"table": self._upload(rows)},
        )
        preview_url = start["Location"]
        draft = LocKitImportDraft.objects.get()
        prepare_loc_kit_string_update.apply(
            kwargs={"draft_id": draft.pk}, task_id=str(draft.prepare_task_id)
        )
        self.client.post(preview_url, {"action": "confirm"})
        draft.refresh_from_db()

        # ruff: ignore[private-member-access]
        real_explanations = loc_kit._apply_kit_explanations_locked
        calls = {"n": 0}

        def fail_explanations_on_third_portion(*args, **kwargs):
            calls["n"] += 1
            if calls["n"] == 3:
                msg = "Locked components cannot be updated."
                raise ValidationError(msg)
            return real_explanations(*args, **kwargs)

        with patch.object(
            loc_kit,
            "_apply_kit_explanations_locked",
            side_effect=fail_explanations_on_third_portion,
        ):
            apply_loc_kit_string_update_draft.apply(
                kwargs={"draft_id": draft.pk}, task_id=str(draft.apply_task_id)
            )

        draft.refresh_from_db()
        self.assertEqual(draft.state, LocKitImportDraft.State.FAILED)
        # The third portion's append rolled back together with the failed
        # Explanation call: the cursor and counters never advanced past the
        # second portion's boundary, and no "m50".."m74" unit exists.
        self.assertEqual(draft.next_row, 50)
        self.assertEqual(draft.progress.get("added"), 50)
        self.assertEqual(
            self.component.source_translation.unit_set.filter(
                context__startswith="m"
            ).count(),
            50,
        )
        # The first two portions' owned pending rows are still legitimately
        # awaiting finalization (the flush only runs once the whole table
        # is applied); only the rolled-back third portion left none.
        owned_contexts = set(
            PendingUnitChange.objects.filter(
                metadata__loc_kit_draft_id=str(draft.pk)
            ).values_list("unit__context", flat=True)
        )
        self.assertTrue(owned_contexts)
        self.assertFalse({f"m{i}" for i in range(50, 75)} & owned_contexts)

        with (
            patch(
                "weblate.trans.tasks.apply_loc_kit_string_update_draft.apply_async"
            ) as dispatch,
            self.captureOnCommitCallbacks(execute=True),
        ):
            self.client.post(preview_url, {"action": "retry"})
        draft.refresh_from_db()
        dispatch.assert_called_once_with(
            kwargs={"draft_id": draft.pk},
            task_id=str(draft.apply_task_id),
            priority=0,
        )

        apply_loc_kit_string_update_draft.apply(
            kwargs={"draft_id": draft.pk}, task_id=str(draft.apply_task_id)
        )

        draft.refresh_from_db()
        self.assertEqual(draft.state, LocKitImportDraft.State.COMPLETED)
        # Every row landed exactly once, all reported as added: the rolled
        # back third portion was never partially applied, so the retry adds
        # it fresh instead of replaying it as already-existing.
        self.assertEqual(draft.progress["added"], row_count)
        self.assertEqual(draft.progress["existing"], 0)
        units = self.component.source_translation.unit_set.filter(
            context__startswith="m"
        )
        self.assertEqual(units.count(), row_count)
        self.assertEqual(len(set(units.values_list("context", flat=True))), row_count)
        self.assertFalse(
            PendingUnitChange.objects.filter(
                metadata__loc_kit_draft_id=str(draft.pk)
            ).exists()
        )

    @override_settings(CELERY_TASK_ALWAYS_EAGER=False)
    def test_start_queues_large_table_preparation_without_unit_writes(self) -> None:
        kit = "key,en,cs\n" + "".join(
            f"new_key_{number},Hello {number},Ahoj {number}\n"
            for number in range(2_501)
        )
        before = self.component.source_translation.unit_set.count()

        with (
            patch(
                "weblate.trans.tasks.prepare_loc_kit_string_update.apply_async"
            ) as dispatch,
            self.captureOnCommitCallbacks(execute=True),
        ):
            response = self.client.post(
                reverse(
                    "loc-kit-strings-update",
                    kwargs={"path": self.component.get_url_path()},
                ),
                {"table": self._upload(kit)},
            )

        self.assertEqual(response.status_code, 302)
        draft = LocKitImportDraft.objects.get()
        self.assertEqual(draft.state, LocKitImportDraft.State.PREPARING)
        dispatch.assert_called_once_with(
            kwargs={"draft_id": draft.pk},
            task_id=str(draft.prepare_task_id),
            priority=0,
        )
        self.assertEqual(self.component.source_translation.unit_set.count(), before)

    @override_settings(CELERY_TASK_ALWAYS_EAGER=False)
    def test_preparation_writes_private_packet_without_units(self) -> None:
        kit = "key,en,cs,Explanation\nnew_key,Hello,Ahoj,Shown at startup.\n"
        start = self.client.post(
            reverse(
                "loc-kit-strings-update",
                kwargs={"path": self.component.get_url_path()},
            ),
            {"table": self._upload(kit)},
        )
        draft = LocKitImportDraft.objects.get()
        before = self.component.source_translation.unit_set.count()

        result = prepare_loc_kit_string_update.apply(
            kwargs={"draft_id": draft.pk},
            task_id=str(draft.prepare_task_id),
        )

        self.assertTrue(result.successful())
        draft.refresh_from_db()
        self.assertEqual(start.status_code, 302)
        self.assertEqual(draft.state, LocKitImportDraft.State.PREVIEW_READY)
        self.assertTrue(draft.prepared_payload.name)
        self.assertFalse(draft.uploaded.name)
        self.assertEqual(self.component.source_translation.unit_set.count(), before)


class LocKitDispatchDrainConcurrencyTest(TransactionTestCase):
    """
    The periodic drain skips a dispatch intent whose row is locked.

    The drain claims with ``SELECT ... FOR UPDATE SKIP LOCKED`` so a
    concurrent claim (the on-commit fast path or a worker-side reservation)
    never aborts the sweep; the lock owner publishes. A plain TestCase's
    outer transaction would hide the second connection, so this needs real
    commits.
    """

    def setUp(self) -> None:
        super().setUp()
        self.project = Project.objects.create(
            name="Probe", slug="probe", web="https://example.com/"
        )
        self.user = create_another_user(suffix="-dispatch")
        self.user.is_superuser = True
        self.user.save(update_fields=["is_superuser"])

    def _draft(self) -> LocKitImportDraft:
        task_id = uuid.uuid4()
        return LocKitImportDraft.objects.create(
            owner=self.user,
            session_key="dispatch-test",
            project=self.project,
            slug="component",
            name="Component",
            source_filename="update.csv",
            state=LocKitImportDraft.State.APPLYING,
            apply_task_id=task_id,
            dispatch_task_id=task_id,
            dispatch_phase=LocKitImportDraft.DispatchPhase.APPLY,
            dispatch_requested_at=timezone.now(),
        )

    def test_drain_skips_a_locked_intent_and_publishes_after_release(self) -> None:
        draft = self._draft()
        connection.close()  # this connection must not straddle the threads

        barrier = threading.Barrier(2)
        release = threading.Event()

        def holder() -> None:
            try:
                with transaction.atomic():
                    LocKitImportDraft.objects.select_for_update().get(pk=draft.pk)
                    barrier.wait(timeout=10)
                    release.wait(timeout=10)
            finally:
                connection.close()

        thread = threading.Thread(target=holder)
        thread.start()
        try:
            barrier.wait(timeout=10)
            with patch(
                "weblate.trans.tasks.apply_loc_kit_string_update_draft.apply_async"
            ) as dispatch:
                drain_loc_kit_dispatches()
            dispatch.assert_not_called()
        finally:
            release.set()
            thread.join(timeout=15)
            connection.close()

        draft.refresh_from_db()
        self.assertIsNone(draft.dispatch_published_at)

        with patch(
            "weblate.trans.tasks.apply_loc_kit_string_update_draft.apply_async"
        ) as dispatch:
            drain_loc_kit_dispatches()
        dispatch.assert_called_once_with(
            kwargs={"draft_id": draft.pk},
            task_id=str(draft.apply_task_id),
            priority=0,
        )
        draft.refresh_from_db()
        self.assertIsNotNone(draft.dispatch_published_at)


class LocKitAtomicPortionConcurrencyTest(RepoTestMixin, TransactionTestCase):
    """
    Real concurrent deliveries racing ``apply_loc_kit_portion``'s fencing.

    A plain TestCase runs inside one outer transaction (savepoints, not
    separate commits), so a second thread's connection could never see a
    committed change from another - this needs real commits, hence
    TransactionTestCase, mirroring ``JudgeResolutionRealConcurrencyTest``.
    Each worker fetches its own fresh ``Component`` instance, exactly like
    two independent Celery deliveries would.
    """

    def setUp(self) -> None:
        self.clone_test_repos()
        super().setUp()
        self.component = self.create_po_mono(name="ConcurrentStrings")
        self.component.new_lang = "add"
        self.component.edit_template = True
        self.component.save(update_fields=["new_lang", "edit_template"])
        self.user = create_test_user()
        self.user.is_superuser = True
        self.user.save(update_fields=["is_superuser"])

    @staticmethod
    def _row(key: str) -> StringUnit:
        return StringUnit(
            key=key,
            values={"en": f"Value {key}"},
            comments=(),
            references=(),
            row=2,
            explanation="",
            flags="",
        )

    def _draft(self, task_id) -> LocKitImportDraft:
        return LocKitImportDraft.objects.create(
            owner=self.user,
            session_key="concurrency-test",
            kind=LocKitImportDraft.Kind.STRING,
            project=self.component.project,
            slug=self.component.slug,
            name=self.component.name,
            source_filename="update.csv",
            target_component=self.component,
            state=LocKitImportDraft.State.APPLYING,
            apply_task_id=task_id,
            apply_started_at=timezone.now(),
        )

    def test_stale_uuid_creates_zero_units_after_retry_installs_new_uuid(
        self,
    ) -> None:
        """
        A stale, paused task must create nothing once retry replaces its UUID.

        The retry replaces the fencing UUID before the stale task ever
        reaches the draft lock.
        """
        old_task_id = uuid.uuid4()
        new_task_id = uuid.uuid4()
        draft = self._draft(old_task_id)
        connection.close()  # this connection must not straddle the threads

        retried = threading.Event()
        outcome: dict[str, object] = {}

        def stale_worker() -> None:
            try:
                retried.wait(timeout=10)
                component = Component.objects.get(pk=self.component.pk)
                outcome["result"] = loc_kit.apply_loc_kit_portion(
                    user=self.user,
                    component=component,
                    units=(self._row("stale_key"),),
                    overwrite_explanations=False,
                    pending_owner=str(draft.pk),
                    explanation_baseline={},
                    draft_id=draft.pk,
                    task_id=old_task_id,
                    expected_cursor=0,
                    total_rows=1,
                )
            finally:
                connection.close()

        thread = threading.Thread(target=stale_worker)
        thread.start()
        # A real committed retry installs a fresh UUID while the stale
        # worker's delivery is paused before it ever reaches the draft lock.
        LocKitImportDraft.objects.filter(pk=draft.pk).update(apply_task_id=new_task_id)
        retried.set()
        thread.join(timeout=15)

        self.assertIsNone(outcome.get("result"))
        self.assertFalse(
            self.component.source_translation.unit_set.filter(
                context="stale_key"
            ).exists()
        )

    def test_duplicate_same_uuid_delivery_commits_at_most_one_portion(
        self,
    ) -> None:
        """
        Two workers delivering the same UUID commit at most one portion.

        The second worker observes the advanced cursor inside the lock and
        returns without mutating anything.
        """
        task_id = uuid.uuid4()
        draft = self._draft(task_id)
        connection.close()

        barrier = threading.Barrier(2, timeout=10)
        outcomes: list[object] = []
        lock = threading.Lock()

        def worker() -> None:
            try:
                barrier.wait(timeout=10)
                component = Component.objects.get(pk=self.component.pk)
                result = loc_kit.apply_loc_kit_portion(
                    user=self.user,
                    component=component,
                    units=(self._row("dup_key"),),
                    overwrite_explanations=False,
                    pending_owner=str(draft.pk),
                    explanation_baseline={},
                    draft_id=draft.pk,
                    task_id=task_id,
                    expected_cursor=0,
                    total_rows=1,
                )
            finally:
                connection.close()
            with lock:
                outcomes.append(result)

        threads = [threading.Thread(target=worker) for _ in range(2)]
        for thread in threads:
            thread.start()
        for thread in threads:
            thread.join(timeout=15)

        self.assertEqual(len(outcomes), 2, "both workers must finish and report")
        successes = [outcome for outcome in outcomes if outcome is not None]
        self.assertEqual(
            len(successes), 1, f"exactly one portion must commit, got: {outcomes}"
        )
        self.assertEqual(
            self.component.source_translation.unit_set.filter(
                context="dup_key"
            ).count(),
            1,
        )
        draft.refresh_from_db()
        self.assertEqual(draft.next_row, 1)
        self.assertEqual(draft.progress.get("added"), 1)

    def test_explanation_error_rolls_back_an_already_added_key(self) -> None:
        """
        An Explanation failure must roll back an already-added key too.

        An exception raised while applying the Explanation, after the
        append would previously have already committed as its own separate
        transaction, must roll back the append too under real commit
        semantics: no key, no owned pending row and no cursor advance
        survives without its paired outcome. A second, unpatched call then
        applies the whole portion cleanly, reported as added.
        """
        task_id = uuid.uuid4()
        draft = self._draft(task_id)
        row = self._row("rollback_key")
        row = replace(row, explanation="Some context.")

        with (
            patch.object(
                loc_kit,
                "_apply_kit_explanations_locked",
                side_effect=ValidationError("boom"),
            ),
            self.assertRaises(ValidationError),
        ):
            loc_kit.apply_loc_kit_portion(
                user=self.user,
                component=self.component,
                units=(row,),
                overwrite_explanations=False,
                pending_owner=str(draft.pk),
                explanation_baseline={},
                draft_id=draft.pk,
                task_id=task_id,
                expected_cursor=0,
                total_rows=1,
            )

        self.assertFalse(
            self.component.source_translation.unit_set.filter(
                context="rollback_key"
            ).exists()
        )
        self.assertFalse(
            PendingUnitChange.objects.filter(
                metadata__loc_kit_draft_id=str(draft.pk)
            ).exists()
        )
        draft.refresh_from_db()
        self.assertEqual(draft.next_row, 0)
        self.assertEqual(draft.progress, {})

        result = loc_kit.apply_loc_kit_portion(
            user=self.user,
            component=self.component,
            units=(row,),
            overwrite_explanations=False,
            pending_owner=str(draft.pk),
            explanation_baseline={},
            draft_id=draft.pk,
            task_id=task_id,
            expected_cursor=0,
            total_rows=1,
        )
        self.assertIsNotNone(result)
        self.assertEqual(result.strings.added, 1)
        self.assertEqual(result.explanations.set_count, 1)
        self.assertTrue(
            self.component.source_translation.unit_set.filter(
                context="rollback_key"
            ).exists()
        )
        draft.refresh_from_db()
        self.assertEqual(draft.next_row, 1)
        self.assertEqual(draft.progress.get("added"), 1)


class LocKitFinalizerRecoveryTest(ViewTestCase):
    """
    ``_flush_loc_kit_pending_changes`` locking, ownership and recovery.

    Exercises the finalizer directly against a real po-mono repository so
    the repository lock, metadata-derived ownership and the VCS-commit
    recovery path are proven against real commits, not only through the
    full task/view machinery.
    """

    def setUp(self) -> None:
        super().setUp()
        self.user.is_superuser = True
        self.user.save()
        self.component = self.create_po_mono(project=self.project, name="Strings")

    def _row(self, key: str, **overrides) -> StringUnit:
        base = {
            "key": key,
            "values": {"en": f"Value {key}"},
            "comments": (),
            "references": (),
            "row": 2,
            "explanation": "",
            "flags": "",
        }
        base.update(overrides)
        return StringUnit(**base)

    def _draft(self, task_id) -> LocKitImportDraft:
        return LocKitImportDraft.objects.create(
            owner=self.user,
            session_key="finalizer-test",
            kind=LocKitImportDraft.Kind.STRING,
            project=self.component.project,
            slug=self.component.slug,
            name=self.component.name,
            source_filename="update.csv",
            target_component=self.component,
            state=LocKitImportDraft.State.APPLYING,
            apply_task_id=task_id,
            apply_started_at=timezone.now(),
        )

    def test_finalizer_holds_repository_lock_before_committing(self) -> None:
        """A lock timeout must leave every pending row and the draft untouched."""
        task_id = uuid.uuid4()
        draft = self._draft(task_id)
        result = loc_kit.append_translation_strings(
            user=self.user,
            component=self.component,
            units=(self._row("locked_key"),),
            pending_owner=str(draft.pk),
        )

        with (
            patch.object(
                Component,
                "locked_for_update",
                side_effect=WeblateLockTimeoutError("locked", lock=None),
            ),
            self.assertRaises(WeblateLockTimeoutError),
        ):
            _flush_loc_kit_pending_changes(
                draft_id=draft.pk,
                task_id=task_id,
                component=self.component,
                owner=self.user,
                retry_phase="finalize",
            )

        self.assertTrue(
            PendingUnitChange.objects.filter(pk__in=result.pending_change_ids).exists()
        )
        draft.refresh_from_db()
        self.assertEqual(draft.state, LocKitImportDraft.State.APPLYING)

    def test_finalizer_excludes_a_foreign_pending_row(self) -> None:
        """
        A foreign pending change must survive finalization untouched.

        A pending change never tagged with this draft's metadata - a manual
        translator edit, or another import entirely - must survive
        untouched: only this draft's owned, metadata-tagged rows are
        committed.
        """
        task_id = uuid.uuid4()
        draft = self._draft(task_id)
        owned = loc_kit.append_translation_strings(
            user=self.user,
            component=self.component,
            units=(self._row("owned_key"),),
            pending_owner=str(draft.pk),
        )
        foreign_unit = self.component.source_translation.add_unit(
            None, "foreign_key", "Foreign", author=self.user
        )
        foreign_change = PendingUnitChange.objects.get(unit=foreign_unit, add_unit=True)
        self.assertEqual(foreign_change.metadata, {})

        outcome = _flush_loc_kit_pending_changes(
            draft_id=draft.pk,
            task_id=task_id,
            component=self.component,
            owner=self.user,
            retry_phase="finalize",
        )

        self.assertTrue(outcome)
        self.assertFalse(
            PendingUnitChange.objects.filter(pk__in=owned.pending_change_ids).exists()
        )
        self.assertTrue(PendingUnitChange.objects.filter(pk=foreign_change.pk).exists())
        self.component.source_translation.drop_store_cache()
        content = Path(self.component.source_translation.get_filename()).read_text(
            encoding="utf-8"
        )
        self.assertIn("owned_key", content)
        self.assertNotIn("foreign_key", content)

    def test_finalizer_stops_before_commit_when_permission_is_revoked(self) -> None:
        """Permission loss must prevent any commit, not only fail afterward."""
        task_id = uuid.uuid4()
        draft = self._draft(task_id)
        result = loc_kit.append_translation_strings(
            user=self.user,
            component=self.component,
            units=(self._row("perm_key"),),
            pending_owner=str(draft.pk),
        )
        limited = create_another_user(suffix="-finalize")

        with patch.object(Component, "commit_pending_subset") as commit_mock:
            outcome = _flush_loc_kit_pending_changes(
                draft_id=draft.pk,
                task_id=task_id,
                component=self.component,
                owner=limited,
                retry_phase="finalize",
            )

        self.assertFalse(outcome)
        commit_mock.assert_not_called()
        self.assertTrue(
            PendingUnitChange.objects.filter(pk__in=result.pending_change_ids).exists()
        )
        draft.refresh_from_db()
        self.assertEqual(draft.state, LocKitImportDraft.State.FAILED)
        self.assertEqual(draft.error_code, "finalize-forbidden")
        self.assertEqual(draft.retry_phase, "finalize")

    def test_finalizer_recovers_from_a_crash_after_vcs_commit(self) -> None:
        """
        Crash recovery after a VCS commit must never create a second commit.

        A pending row describing content already on disk - exactly what such
        a crash leaves behind - is cleared without a second commit, and the
        finalizer reports success so the draft can still complete. This
        must hold even when the retry happens minutes later: the replay
        reuses the original batch's own max ``timestamp`` for the PO
        header, so a moved wall clock never manufactures a phantom
        revision-date diff.
        """
        task_id = uuid.uuid4()
        draft = self._draft(task_id)
        result = loc_kit.append_translation_strings(
            user=self.user,
            component=self.component,
            units=(self._row("crash_key"),),
            pending_owner=str(draft.pk),
        )
        # ``commit_pending_subset`` applies ONE header timestamp - the max
        # across every pending change in the batch - to every translation
        # it writes. The replay below must reproduce that exact max.
        original_timestamp = max(
            PendingUnitChange.objects.filter(
                pk__in=result.pending_change_ids
            ).values_list("timestamp", flat=True)
        )
        first_outcome = _flush_loc_kit_pending_changes(
            draft_id=draft.pk,
            task_id=task_id,
            component=self.component,
            owner=self.user,
            retry_phase="finalize",
        )
        self.assertTrue(first_outcome)
        commits_after_first = self.component.repository.count_outgoing()

        source_unit = self.component.source_translation.unit_set.get(
            context="crash_key"
        )
        # The row a real crash leaves behind is the SAME row, never
        # touched: its ``timestamp`` is whatever the first attempt already
        # used, not a freshly minted one.
        replay_change = PendingUnitChange.objects.create(
            unit=source_unit,
            author=self.user,
            target=source_unit.target,
            explanation=source_unit.explanation,
            source_unit_explanation=source_unit.explanation,
            state=source_unit.state,
            add_unit=False,
            metadata={"loc_kit_draft_id": str(draft.pk)},
            timestamp=original_timestamp,
        )

        moved_clock = timezone.now() + timedelta(minutes=10)
        with patch(
            "weblate.trans.models.translation.timezone.now", return_value=moved_clock
        ):
            recovered = _flush_loc_kit_pending_changes(
                draft_id=draft.pk,
                task_id=task_id,
                component=self.component,
                owner=self.user,
                retry_phase="finalize",
            )

        self.assertTrue(recovered)
        self.assertEqual(
            self.component.repository.count_outgoing(),
            commits_after_first,
            "recovery must not create a second commit",
        )
        self.assertFalse(PendingUnitChange.objects.filter(pk=replay_change.pk).exists())
        draft.refresh_from_db()
        self.assertEqual(draft.state, LocKitImportDraft.State.APPLYING)
